from pathlib import Path
from math import cos, sin, pi

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
PROJECT = ROOT / "projects" / "slide09_whitepaper_static_v2"
ASSET_DIR = PROJECT / "assets"
OUT_DIR = PROJECT / "exports"
VALIDATION_DIR = PROJECT / "validation"
OUT = OUT_DIR / "slide09_whitepaper_static_v2.pptx"

W, H = 16, 9
FONT = "Microsoft YaHei"

PAPER = RGBColor(248, 250, 250)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(0, 72, 84)
TEAL_2 = RGBColor(0, 96, 112)
ICE = RGBColor(226, 243, 247)
ICE_LINE = RGBColor(183, 216, 224)
BLUE = RGBColor(75, 172, 198)
DARK = RGBColor(28, 44, 49)
MUTED = RGBColor(88, 108, 114)
LINE = RGBColor(211, 225, 229)
ORANGE = RGBColor(226, 126, 39)
GRAY = RGBColor(243, 246, 246)


FACTORS = [
    ("F01", "温度", "Temperature", "℃", "-15 / -10 / -5 / 0 ℃"),
    ("F02", "压强", "Pressure", "P", "0.05 / 0.10 / 0.20 MPa"),
    ("F03", "金属材质", "Metal Material", "▰", "Al / Steel / Ti"),
    ("F04", "表面粗糙度", "Surface Roughness", "≈", "Ra 0.2 / 1.0 / 3.0 μm"),
    ("F05", "接触面积", "Contact Area", "□", "1 / 4 / 9 cm²"),
    ("F06", "冰的成分", "Ice Composition", "⌬", "淡水冰 / 海冰 / 盐冰"),
]


def make_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)

    # A restrained raster texture used only as a local visual asset; all text and diagrams stay editable.
    bg = Image.new("RGB", (2048, 1152), (248, 250, 250))
    d = ImageDraw.Draw(bg, "RGBA")
    for x in range(-120, 2200, 180):
        d.line([(x, 130), (x + 95, 1060)], fill=(191, 212, 218, 30), width=3)
    for y in range(170, 1060, 145):
        d.line([(110, y), (1980, y)], fill=(191, 212, 218, 24), width=3)
    for r, a in [(520, 18), (410, 22), (280, 18)]:
        d.ellipse((1080 - r, 80 - r // 3, 1080 + r, 80 + r), outline=(75, 172, 198, a), width=4)
    bg = bg.filter(ImageFilter.GaussianBlur(0.35))
    texture = ASSET_DIR / "whitepaper_subtle_grid_texture.png"
    bg.save(texture, quality=95)

    scene = Image.new("RGB", (1200, 680), (246, 249, 250))
    d = ImageDraw.Draw(scene, "RGBA")
    for y in range(680):
        shade = int(250 - y * 0.018)
        d.line([(0, y), (1200, y)], fill=(shade, shade + 1, shade + 1, 255))
    d.rounded_rectangle((95, 410, 1070, 470), radius=18, fill=(214, 231, 235, 210), outline=(150, 180, 188, 210), width=3)
    d.rectangle((155, 390, 1010, 415), fill=(235, 246, 249, 170))
    d.rounded_rectangle((430, 230, 740, 405), radius=10, fill=(174, 186, 188, 255), outline=(120, 137, 140, 255), width=3)
    for i in range(18):
        x = 440 + i * 17
        d.line([(x, 240), (x + 80, 398)], fill=(236, 241, 242, 32), width=2)
    d.line((585, 105, 585, 230), fill=(0, 72, 84, 220), width=7)
    d.polygon([(565, 120), (605, 120), (585, 92)], fill=(0, 72, 84, 220))
    d.ellipse((455, 390, 715, 438), fill=(75, 172, 198, 26))
    for x in [230, 340, 860, 970]:
        d.line((x, 430, x + 65, 392), fill=(255, 255, 255, 135), width=2)
    d.rounded_rectangle((70, 70, 1130, 610), radius=46, outline=(211, 225, 229, 180), width=4)
    scene = scene.filter(ImageFilter.UnsharpMask(radius=1.2, percent=130, threshold=3))
    scene_path = ASSET_DIR / "metal_ice_contact_inset.png"
    scene.save(scene_path, quality=96)
    return texture, scene_path


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=False, transparency=0, name=None):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    if name:
        shp.name = name
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.fill.transparency = transparency
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    return shp


def line(slide, x1, y1, x2, y2, color=LINE, width=1.0, name=None):
    shp = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    if name:
        shp.name = name
    shp.line.color.rgb = color
    shp.line.width = Pt(width)
    return shp


def text(slide, x, y, w, h, value, size, color=DARK, bold=False, align=PP_ALIGN.LEFT, name=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for idx, raw in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = raw
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def tab(slide, x, y, w, h, fill=TEAL):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = fill
    return shp


def add_header(slide):
    tab(slide, -0.10, 0.08, 1.35, 0.66)
    text(slide, 0.18, 0.20, 0.58, 0.32, "09", 22, WHITE, True, name="slide09_title_index")
    text(slide, 1.42, 0.22, 3.1, 0.36, "结果分析", 23.5, TEAL, True, name="slide09_title_main")
    line(slide, 0.72, 0.86, 1.28, 0.86, TEAL, 3)
    line(slide, 1.40, 0.86, 2.62, 0.86, TEAL_2, 1.4)
    text(
        slide, 1.42, 0.98, 7.2, 0.28,
        "控制变量法下的摩擦系数变化规律  /  Controlled Variable Analysis of Metal-Ice Friction Coefficient",
        11.6, MUTED, name="slide09_title_sub"
    )
    tab(slide, 15.12, 8.50, 0.80, 0.36)
    text(slide, 15.36, 8.58, 0.27, 0.12, "09", 10, WHITE, True, PP_ALIGN.CENTER)
    text(slide, 0.92, 8.48, 8.4, 0.16, "注：当前数据为版式示例，后续可替换为真实实验数据。", 7.5, MUTED, name="slide09_data_source_note")


def add_model_framework(slide):
    text(slide, 0.92, 1.42, 2.25, 0.24, "六因素矩阵 / 模型框架", 13.2, TEAL, True)
    text(slide, 4.68, 1.44, 1.65, 0.16, "Response variable: μ", 7.8, BLUE, True, PP_ALIGN.RIGHT)
    rect(slide, 0.88, 1.82, 6.45, 5.72, WHITE, LINE, True, name="slide09_factor_matrix")

    cx, cy = 4.08, 4.68
    rx, ry = 2.18, 1.85
    for scale, alpha_w in [(1.00, 1.0), (0.76, 0.7)]:
        for i in range(72):
            a1 = 2 * pi * i / 72
            a2 = 2 * pi * (i + 0.52) / 72
            line(slide, cx + rx * scale * cos(a1), cy + ry * scale * sin(a1), cx + rx * scale * cos(a2), cy + ry * scale * sin(a2), ICE_LINE, alpha_w)

    angles = [-120, -60, 0, 60, 120, 180]
    for (fid, cn, en, glyph, levels), deg in zip(FACTORS, angles):
        a = pi * deg / 180
        nx, ny = cx + rx * cos(a), cy + ry * sin(a)
        line(slide, cx, cy, nx, ny, ICE_LINE, 1.1)
    rect(slide, cx - 0.74, cy - 0.42, 1.48, 0.84, ICE, ICE_LINE, True, name="model_response_core")
    text(slide, cx - 0.46, cy - 0.29, 0.92, 0.22, "μ", 24, TEAL, True, PP_ALIGN.CENTER)
    text(slide, cx - 0.52, cy + 0.08, 1.04, 0.12, "friction coefficient", 6.6, MUTED, False, PP_ALIGN.CENTER)
    for (fid, cn, en, glyph, levels), deg in zip(FACTORS, angles):
        a = pi * deg / 180
        nx, ny = cx + rx * cos(a), cy + ry * sin(a)
        rect(slide, nx - 0.86, ny - 0.38, 1.72, 0.76, WHITE, ICE_LINE, True, name=f"factor_node_{en.lower().replace(' ', '_')}")
        rect(slide, nx - 0.76, ny - 0.23, 0.34, 0.34, ICE, ICE_LINE, True)
        text(slide, nx - 0.70, ny - 0.14, 0.22, 0.10, glyph, 8.0, TEAL, True, PP_ALIGN.CENTER)
        text(slide, nx - 0.30, ny - 0.27, 0.40, 0.12, fid, 7.0, BLUE, True)
        text(slide, nx - 0.30, ny - 0.09, 0.94, 0.15, cn, 9.5, TEAL, True)
        text(slide, nx - 0.30, ny + 0.12, 1.02, 0.10, en, 5.8, MUTED)
        text(slide, nx - 0.58, ny + 0.40, 1.16, 0.11, levels, 5.8, MUTED, False, PP_ALIGN.CENTER)

    text(
        slide, 1.22, 7.18, 5.70, 0.18,
        "矩阵总览用于定位六个控制变量；下一阶段确认视觉后再叠加交互分析态。",
        8.2, MUTED, name="slide09_matrix_note"
    )


def add_right_overview(slide, scene_path):
    rect(slide, 7.86, 1.82, 7.06, 5.72, WHITE, LINE, True, name="slide09_experiment_model_overview")
    text(slide, 8.26, 2.20, 3.15, 0.25, "实验手段选择 + 模型构建", 13.2, TEAL, True)
    pic = slide.shapes.add_picture(str(scene_path), Inches(8.26), Inches(2.72), width=Inches(2.86), height=Inches(1.62))
    pic.name = "slide09_metal_ice_inset_image"
    text(slide, 8.42, 4.50, 2.54, 0.16, "金属-冰接触实验场景（局部视觉资产）", 7.8, MUTED, False, PP_ALIGN.CENTER)
    text(
        slide, 11.65, 2.82, 2.62, 1.00,
        "以控制变量法组织实验水平，在同一测量框架下改变单一因素，提取 μ 的变化趋势并回到物理机理解释。",
        10.5, DARK, name="overview_method_text"
    )

    labels = ["控制变量法", "水平设计", "μ 数据采集", "趋势建模", "结果分析"]
    for i, label in enumerate(labels):
        x = 8.34 + i * 1.25
        fill = TEAL if i == 4 else ICE
        fg = WHITE if i == 4 else TEAL
        rect(slide, x, 5.34, 1.08, 0.56, fill, ICE_LINE, True, name=f"overview_step_{i+1}")
        text(slide, x + 0.08, 5.53, 0.92, 0.12, label, 7.6, fg, True, PP_ALIGN.CENTER)
        if i < 4:
            line(slide, x + 1.10, 5.62, x + 1.22, 5.62, MUTED, 0.8)
            text(slide, x + 1.14, 5.51, 0.14, 0.12, "→", 7.2, MUTED, True, PP_ALIGN.CENTER)

    rect(slide, 8.34, 6.54, 5.95, 0.62, GRAY, LINE, True)
    text(slide, 8.72, 6.73, 5.12, 0.18, "静态阶段只呈现分析入口：模型总览、实验方法与建模流程。", 9.3, MUTED)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    VALIDATION_DIR.mkdir(parents=True, exist_ok=True)
    texture, scene = make_assets()
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(texture), Inches(0), Inches(0), width=Inches(W), height=Inches(H)).name = "slide09_background_texture"
    add_header(slide)
    add_model_framework(slide)
    add_right_overview(slide, scene)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
