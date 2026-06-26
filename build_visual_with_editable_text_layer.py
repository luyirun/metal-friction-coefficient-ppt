from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(r"F:\Users\user\Documents\测金属摩擦系数")
SRC_DIR = ROOT / "outputs" / "single_pages"
ASSET_DIR = ROOT / "outputs" / "visual_editable_layer_assets"
OUT = ROOT / "outputs" / "metal_ice_99fid_with_editable_text_layer.pptx"

SLIDE_W, SLIDE_H = 16, 9
COLS, ROWS = 8, 6
FONT = "Microsoft YaHei"


TEXT_ITEMS = {
    1: [
        ("S01_主标题", 0.95, 1.95, 6.50, 1.65, "金属与冰\n摩擦系数的测量", 40, True),
        ("S01_副标题", 0.95, 4.05, 4.40, 0.30, "物理实验创新竞赛答辩", 18, True),
        ("S01_团队", 1.65, 5.80, 3.20, 0.28, "团队：祝冰间隙者", 17, False),
        ("S01_日期", 1.65, 6.78, 3.20, 0.28, "日期：2024.05.20", 17, False),
    ],
    2: [
        ("S02_标题", 1.30, 0.25, 3.20, 0.36, "选题背景", 24, True),
        ("S02_右侧标题", 9.50, 1.50, 4.20, 0.35, "北极航运与冰阻力问题", 20, True),
        ("S02_关键数字_1", 9.55, 2.85, 2.20, 0.40, "25%-55%", 24, True),
        ("S02_关键数字_2", 9.55, 4.18, 2.20, 0.40, "13%-15%", 24, True),
        ("S02_研究意义", 1.05, 5.90, 7.20, 1.25, "研究意义\n揭示船舶破冰与航行效率的关键因素之一\n为船舶材料选型、表面处理与结构设计提供依据\n推动极地装备与工程安全的进步", 12, False),
    ],
    3: [
        ("S03_标题", 1.30, 0.25, 3.20, 0.36, "题目解读", 24, True),
        ("S03_机理分析", 1.65, 1.55, 3.00, 4.90, "机理分析\n冰的黏附性变弱\n表面微观水膜\n水膜润滑作用\n温度与压力影响", 13, False),
        ("S03_装置搭建", 6.05, 1.55, 3.00, 4.90, "装置搭建\n斜面摩擦实验平台\n力学测量与采集\n低温环境控制\n稳定与校准", 13, False),
        ("S03_结果与不确定度", 10.25, 1.55, 3.50, 4.90, "结果与不确定度\n摩擦系数计算\n不确定度评估\n数据性分析\n结果讨论与验证", 13, False),
    ],
    4: [
        ("S04_标题", 1.30, 0.25, 5.20, 0.36, "方法选择 —— 动态斜面法", 24, True),
        ("S04_力学模型", 1.20, 1.30, 2.20, 0.30, "力学模型", 15, True),
        ("S04_理论公式", 9.05, 1.30, 4.80, 3.30, "理论公式\na = g(sinθ - μcosθ)\nμ = (sinθ - a/g) / cosθ", 16, False),
        ("S04_符号说明", 9.20, 5.20, 4.20, 1.30, "a —— 沿斜面方向的加速度\ng —— 重力加速度\nθ —— 斜面倾角\nμ —— 动摩擦系数", 12, False),
    ],
    5: [
        ("S05_标题", 1.30, 0.25, 3.20, 0.36, "装置设计", 24, True),
        ("S05_左模块", 1.15, 1.12, 2.60, 0.28, "斜面冰面模块", 14, True),
        ("S05_右模块", 8.85, 1.12, 3.60, 0.28, "垂直冰面模块（压力加载）", 14, True),
        ("S05_底部标签", 1.30, 7.10, 13.20, 0.35, "模块化设计    快速更换    低温环境模拟    高精度测量", 11, False),
    ],
    6: [
        ("S06_标题", 1.30, 0.25, 3.20, 0.36, "变量设计", 24, True),
        ("S06_副标题", 1.30, 0.95, 4.00, 0.28, "六大影响变量与水平设置", 14, False),
        ("S06_变量表", 1.05, 3.75, 13.80, 1.55, "温度：-5 ℃ / -10 ℃ / -15 ℃\n压强：0.05 MPa / 0.10 MPa / 0.20 MPa\n材质：铝合金 / 不锈钢 / 钛合金\n粗糙度：Ra 0.2 μm / Ra 1.0 μm / Ra 3.0 μm\n面积：1 cm² / 4 cm² / 9 cm²\n成分：淡水冰 / 海冰 / 人工掺盐冰", 10, False),
    ],
    7: [
        ("S07_标题", 1.30, 0.25, 3.20, 0.36, "问题改进", 24, True),
        ("S07_问题与方案", 2.30, 1.45, 10.50, 5.80, "冰面制备不平整：表面粗糙和沉积过大，影响测量重复性 → 定制制冰机 + 控温体系，实现光滑透明平整冰面\n低温环境稳定性差：温度波动导致数据漂移，影响实验结果 → 双层保温 + PID温控，温度波动 ≤ ±0.2 ℃\n传感器可靠性不足：低温漂移与噪声干扰大，影响信号准确性 → 低温型传感器 + 屏蔽布线，多点校准与滤波处理", 11, False),
    ],
    8: [
        ("S08_标题", 1.30, 0.25, 3.20, 0.36, "数据处理", 24, True),
        ("S08_流程", 1.25, 1.70, 13.50, 5.30, "1 数据采集：力信号、位移信号、温度/压力、采样频率 100 Hz\n2 加速度计算：位移二次微分、滤波处理、得到 a(t)\n3 摩擦系数计算：μ = (sinθ - a/g) / cosθ，代入角度与加速度计算 μ\n4 不确定度评估：A类不确定度、B类不确定度、合成标准不确定度、扩展不确定度 U", 11, False),
    ],
    9: [
        ("S09_标题", 1.30, 0.25, 3.20, 0.36, "结果分析", 24, True),
        ("S09_图表标题", 2.85, 1.25, 4.00, 0.30, "摩擦系数对比趋势（示例）", 14, True),
        ("S09_数据解读", 9.45, 1.55, 4.00, 2.70, "数据解读\n温度升高，μ整体下降\n粗糙度增大，μ上升\n压强增大，先升后降\n材质差异显著，钛合金综合表现最优", 12, False),
        ("S09_结论", 9.75, 6.05, 3.10, 0.45, "结果与文献趋势一致，验证方法的有效性", 11, True),
    ],
    10: [
        ("S10_标题", 1.30, 0.25, 3.20, 0.36, "总结展望", 24, True),
        ("S10_创新点", 1.65, 1.25, 4.60, 4.90, "创新点\n模块化多工况装置设计\n低温稳定控制与高精度测量\n多因素系统研究与量化分析\n不确定度评估与可靠性提升", 12, False),
        ("S10_应用前景", 10.80, 1.25, 3.60, 4.90, "工程应用前景\n极地船舶设计优化：降低航行阻力\n表面工程与减阻开发：提升抗冰性能\n极地装备与结构安全：提供理论依据", 12, False),
        ("S10_展望", 2.00, 7.35, 11.20, 0.30, "未来将拓展更多材料与工况，建立摩擦数据库，服务极地工程与安全航行。", 10, False),
    ],
}


def make_run_transparent(run):
    run.font.color.rgb = RGBColor(255, 255, 255)
    rpr = run._r.get_or_add_rPr()
    solid = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill")
    if solid is None:
        solid = OxmlElement("a:solidFill")
        rpr.append(solid)
    srgb = solid.find("{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
    if srgb is None:
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", "FFFFFF")
        solid.append(srgb)
    alpha = OxmlElement("a:alpha")
    alpha.set("val", "0")
    srgb.append(alpha)


def add_invisible_text(slide, name, x, y, w, h, text, size, bold):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold if idx == 0 else False
        make_run_transparent(run)
    box.fill.background()
    box.line.fill.background()
    return box


def add_sliced_page(slide, page_no):
    src = SRC_DIR / f"page_{page_no:02d}_industrial_whitepaper.png"
    im = Image.open(src).convert("RGB")
    w, h = im.size
    tw, th = w // COLS, h // ROWS
    for r in range(ROWS):
        for c in range(COLS):
            left = c * tw
            top = r * th
            right = w if c == COLS - 1 else (c + 1) * tw
            bottom = h if r == ROWS - 1 else (r + 1) * th
            tile = im.crop((left, top, right, bottom))
            path = ASSET_DIR / f"p{page_no:02d}_{r:02d}_{c:02d}.png"
            tile.save(path, optimize=True)
            slide.shapes.add_picture(
                str(path),
                Inches(SLIDE_W * left / w),
                Inches(SLIDE_H * top / h),
                width=Inches(SLIDE_W * (right - left) / w),
                height=Inches(SLIDE_H * (bottom - top) / h),
            )


def build():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    for page_no in range(1, 11):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_sliced_page(slide, page_no)
        for item in TEXT_ITEMS[page_no]:
            add_invisible_text(slide, *item)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
