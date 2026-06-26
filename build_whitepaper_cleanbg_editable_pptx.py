from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "outputs"
SRC_DIR = OUT_DIR / "single_pages"
BG_DIR = OUT_DIR / "whitepaper_clean_backgrounds"
OUT = OUT_DIR / "metal_ice_whitepaper_cleanbg_editable.pptx"

W, H = 16, 9
PW, PH = 2048, 1152
FONT = "Microsoft YaHei"

TEAL = RGBColor(0, 72, 84)
DARK = RGBColor(28, 44, 49)
MUTED = RGBColor(78, 98, 104)
ORANGE = RGBColor(226, 126, 39)
GREEN = RGBColor(61, 137, 78)
WHITE = RGBColor(255, 255, 255)

PAPER_HEX = "#f7fafa"
WHITE_HEX = "#ffffff"
TEAL_HEX = "#004854"


def px_box(box):
    x, y, w, h = box
    return (
        round(x / W * PW),
        round(y / H * PH),
        round((x + w) / W * PW),
        round((y + h) / H * PH),
    )


def clean_page(page, masks):
    BG_DIR.mkdir(parents=True, exist_ok=True)
    im = Image.open(SRC_DIR / f"page_{page:02d}_industrial_whitepaper.png").convert("RGB")
    draw = ImageDraw.Draw(im)
    for item in masks:
        if len(item) == 4:
            box, color = item, WHITE_HEX
        else:
            box, color = item[:4], item[4]
        draw.rectangle(px_box(box), fill=color)
    out = BG_DIR / f"page_{page:02d}_clean.png"
    im.save(out, quality=95)
    return out


def build_backgrounds():
    # Masks are in slide inches. They erase baked text while leaving the source visual system intact.
    return {
        1: clean_page(1, [
            (0, 0, 1.58, 0.78, PAPER_HEX),
            (0.70, 1.40, 5.65, 1.86, PAPER_HEX),
            (0.70, 3.32, 3.85, 0.40, PAPER_HEX),
            (0.68, 4.16, 3.05, 1.55, PAPER_HEX),
            (0.78, 7.92, 3.85, 0.36, TEAL_HEX),
            (15.10, 8.50, 0.78, 0.36, TEAL_HEX),
        ]),
        2: clean_page(2, [
            (0, 0, 4.40, 0.95, PAPER_HEX),
            (8.86, 1.35, 5.30, 3.24, WHITE_HEX),
            (0.98, 5.34, 6.20, 1.66, WHITE_HEX),
            (15.10, 8.50, 0.78, 0.36, TEAL_HEX),
        ]),
        3: clean_page(3, [
            (0, 0, 4.20, 0.95, PAPER_HEX),
            (1.10, 1.35, 13.60, 5.70, WHITE_HEX),
            (15.10, 8.50, 0.78, 0.36, TEAL_HEX),
        ]),
        4: clean_page(4, [
            (0, 0, 6.50, 0.95, PAPER_HEX),
            (1.05, 1.15, 1.72, 0.34, WHITE_HEX),
            (8.60, 1.15, 5.10, 5.85, WHITE_HEX),
            (15.10, 8.50, 0.78, 0.36, TEAL_HEX),
        ]),
        5: clean_page(5, [
            (0, 0, 3.90, 0.95, PAPER_HEX),
            (0.90, 1.00, 2.70, 0.42, WHITE_HEX),
            (8.20, 1.00, 4.70, 0.42, WHITE_HEX),
            (0.62, 1.62, 1.45, 4.05, WHITE_HEX),
            (12.65, 1.62, 1.72, 4.15, WHITE_HEX),
            (0.82, 7.50, 13.90, 0.42, WHITE_HEX),
            (15.10, 8.50, 0.78, 0.36, TEAL_HEX),
        ]),
        6: clean_page(6, [
            (0, 0, 4.20, 0.95, PAPER_HEX),
            (0.90, 0.95, 5.70, 0.34, WHITE_HEX),
            (0.92, 3.55, 14.25, 2.10, WHITE_HEX),
            (15.10, 8.50, 0.78, 0.36, TEAL_HEX),
        ]),
        7: clean_page(7, [
            (0, 0, 4.05, 0.95, PAPER_HEX),
            (1.34, 1.46, 3.75, 5.05, "#f7fafa"),
            (7.30, 1.46, 3.85, 5.05, "#f7fafa"),
            (15.10, 8.50, 0.78, 0.36, TEAL_HEX),
        ]),
        8: clean_page(8, [
            (0, 0, 4.05, 0.95, PAPER_HEX),
            (1.20, 1.95, 13.60, 4.90, WHITE_HEX),
            (15.10, 8.50, 0.78, 0.36, TEAL_HEX),
        ]),
        9: clean_page(9, [
            (0, 0, 4.05, 0.95, PAPER_HEX),
            (1.82, 1.08, 4.30, 0.34, WHITE_HEX),
            (8.80, 1.30, 4.55, 3.15, WHITE_HEX),
            (8.95, 5.48, 3.85, 0.95, "#f2f8f8"),
            (15.10, 8.50, 0.78, 0.36, TEAL_HEX),
        ]),
        10: clean_page(10, [
            (0, 0, 4.05, 0.95, PAPER_HEX),
            (2.10, 1.05, 1.35, 0.30, WHITE_HEX),
            (10.35, 1.05, 2.20, 0.30, WHITE_HEX),
            (1.85, 1.52, 4.95, 4.20, WHITE_HEX),
            (11.00, 1.52, 3.35, 4.20, WHITE_HEX),
            (1.85, 7.32, 10.80, 0.36, WHITE_HEX),
            (15.10, 8.50, 0.78, 0.36, TEAL_HEX),
        ]),
    }


def txt(slide, x, y, w, h, value, size, color=DARK, bold=False, align=PP_ALIGN.LEFT, name=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for idx, raw in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = raw
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def add_bg(slide, path):
    slide.shapes.add_picture(str(path), Inches(0), Inches(0), width=Inches(W), height=Inches(H))


def bullet(slide, x, y, value, size=10.5, color=DARK):
    txt(slide, x, y, 0.12, 0.15, "•", size + 2, TEAL, True, PP_ALIGN.CENTER)
    txt(slide, x + 0.20, y + 0.01, 4.1, 0.20, value, size, color)


def page_text(slide, page):
    txt(slide, 0.18, 0.17, 0.55, 0.26, f"{page:02d}", 21.5, WHITE, True)
    txt(slide, 15.36, 8.58, 0.27, 0.12, f"{page:02d}", 10, WHITE, True, PP_ALIGN.CENTER)


def title(slide, page, value, subtitle=None):
    page_text(slide, page)
    txt(slide, 1.33, 0.20, 6.4, 0.32, value, 22, TEAL, True, name=f"S{page:02d}_title")
    if subtitle:
        txt(slide, 1.33, 0.94, 6.4, 0.22, subtitle, 12.5, MUTED, name=f"S{page:02d}_subtitle")


def s01(slide, bg):
    add_bg(slide, bg[1])
    page_text(slide, 1)
    txt(slide, 0.78, 1.52, 4.65, 1.25, "金属与冰\n摩擦系数的测量", 31, TEAL, True, name="S01_main")
    txt(slide, 0.78, 3.36, 3.1, 0.26, "物理实验创新竞赛答辩", 15, DARK, True)
    txt(slide, 1.16, 4.34, 2.1, 0.18, "团队：祝冰间隙者", 11.2, DARK)
    txt(slide, 1.16, 5.16, 2.1, 0.18, "日期：2024.05.20", 11.2, DARK)
    txt(slide, 1.02, 8.03, 3.30, 0.15, "精准测量  ·  理解机理  ·  服务工程", 9.3, WHITE, True)


def s02(slide, bg):
    add_bg(slide, bg[2])
    title(slide, 2, "选题背景")
    txt(slide, 9.00, 1.48, 4.4, 0.30, "北极航运与冰阻力问题", 17.5, TEAL, True)
    txt(slide, 9.00, 2.14, 3.8, 0.18, "航行阻力中，冰障碍阻力占比", 11.2, MUTED)
    txt(slide, 9.00, 2.46, 2.1, 0.30, "25%-55%", 21, ORANGE, True)
    txt(slide, 9.00, 3.82, 4.4, 0.18, "冰面相关事故中，摩擦失控问题占比", 11.2, MUTED)
    txt(slide, 9.00, 4.14, 2.1, 0.30, "13%-15%", 21, ORANGE, True)
    txt(slide, 1.35, 5.55, 1.2, 0.20, "研究意义", 12.2, TEAL, True)
    for i, item in enumerate(["揭示船舶破冰与航行效率的关键因素之一", "为船舶材料选型、表面处理与结构设计提供依据", "推动极地装备与工程安全的进步"]):
        bullet(slide, 1.22, 5.90 + i * 0.32, item, 10.2)


def s03(slide, bg):
    add_bg(slide, bg[3])
    title(slide, 3, "题目解读")
    sections = [
        ("机理分析", "冰的黏附性变弱\n表面微观水膜\n水膜润滑作用\n温度与压力影响"),
        ("装置搭建", "斜面摩擦实验平台\n力学测量与采集\n低温环境控制\n稳定与校准"),
        ("结果与不确定度", "摩擦系数计算\n不确定度评估\n数据性分析\n结果讨论与验证"),
    ]
    for i, (h, body) in enumerate(sections):
        x = 1.35 + i * 4.45
        txt(slide, x + 0.55, 1.63, 2.90, 0.22, h, 14.5, TEAL, True, PP_ALIGN.CENTER)
        for j, item in enumerate(body.split("\n")):
            bullet(slide, x + 0.62, 4.16 + j * 0.42, item, 10.8)


def s04(slide, bg):
    add_bg(slide, bg[4])
    title(slide, 4, "方法选择 —— 动态斜面法")
    txt(slide, 1.10, 1.26, 1.4, 0.22, "力学模型", 13.2, TEAL, True)
    txt(slide, 8.78, 1.26, 1.4, 0.22, "理论公式", 13.2, TEAL, True)
    txt(slide, 9.48, 1.98, 3.50, 0.22, "a = g(sinθ - μcosθ)", 16, DARK, False, PP_ALIGN.CENTER)
    txt(slide, 11.10, 2.65, 0.25, 0.18, "↓", 18, TEAL, True, PP_ALIGN.CENTER)
    txt(slide, 9.45, 3.34, 3.65, 0.24, "μ = (sinθ - a/g) / cosθ", 16, DARK, False, PP_ALIGN.CENTER)
    txt(slide, 9.20, 4.34, 3.7, 1.15, "a —— 沿斜面方向的加速度\ng —— 重力加速度\nθ —— 斜面倾角\nμ —— 动摩擦系数", 11.5, DARK)


def s05(slide, bg):
    add_bg(slide, bg[5])
    title(slide, 5, "装置设计")
    txt(slide, 1.00, 1.13, 2.10, 0.20, "斜面冰面模块", 12.8, TEAL, True)
    txt(slide, 8.35, 1.13, 3.30, 0.20, "垂直冰面模块（压力加载）", 12.8, TEAL, True)
    for i, label in enumerate(["导轨与滑块", "力传感器", "位移传感器", "冰面板", "角度调节机构", "支撑框架"]):
        txt(slide, 0.72, 1.86 + i * 0.48, 1.20, 0.16, label, 9.2, MUTED)
    for i, label in enumerate(["压力加载装置", "力传感器", "冰面板", "温控仓体", "底盘"]):
        txt(slide, 13.10, 1.82 + i * 0.70, 1.25, 0.16, label, 9.2, MUTED)
    for i, label in enumerate(["模块化设计", "快速更换", "低温环境模拟", "高精度测量"]):
        txt(slide, 1.10 + i * 3.35, 7.72, 1.35, 0.16, label, 9.6, TEAL, True, PP_ALIGN.CENTER)


def s06(slide, bg):
    add_bg(slide, bg[6])
    title(slide, 6, "变量设计", "六大影响变量与水平设置")
    items = [
        ("温度", "-5 ℃\n-10 ℃\n-15 ℃"),
        ("压强", "0.05 MPa\n0.10 MPa\n0.20 MPa"),
        ("材质", "铝合金\n不锈钢\n钛合金"),
        ("粗糙度", "Ra 0.2 μm\nRa 1.0 μm\nRa 3.0 μm"),
        ("面积", "1 cm²\n4 cm²\n9 cm²"),
        ("成分", "淡水冰\n海冰\n人工掺盐冰"),
    ]
    for i, (h, body) in enumerate(items):
        x = 1.25 + i * 2.40
        txt(slide, x - 0.45, 3.34, 0.90, 0.18, h, 12.0, TEAL, True, PP_ALIGN.CENTER)
        txt(slide, x - 0.62, 4.02, 1.24, 0.80, body, 10.5, DARK, False, PP_ALIGN.CENTER)


def s07(slide, bg):
    add_bg(slide, bg[7])
    title(slide, 7, "问题改进")
    rows = [
        ("冰面制备不平整", "表面粗糙和气泡大\n影响测量重复性", "改进方案", "定制制冰机 + 控温体系\n实现光滑透明平整冰面"),
        ("低温环境稳定性差", "温度波动导致数据漂移\n影响实验结果", "改进方案", "双层保温 + PID温控\n温度波动 ≤ ±0.2 ℃"),
        ("传感器可靠性不足", "低温漂移与噪声干扰大\n影响信号准确性", "改进方案", "低温型传感器 + 屏蔽布线\n多点校准与滤波处理"),
    ]
    for i, (p, desc, st, sol) in enumerate(rows):
        y = 1.66 + i * 1.72
        txt(slide, 1.70, y, 2.42, 0.20, p, 11.4, DARK, True)
        txt(slide, 1.70, y + 0.30, 2.55, 0.38, desc, 9.7, MUTED)
        txt(slide, 7.58, y, 1.16, 0.20, st, 11.4, GREEN, True)
        txt(slide, 7.58, y + 0.30, 2.70, 0.38, sol, 9.7, DARK)


def s08(slide, bg):
    add_bg(slide, bg[8])
    title(slide, 8, "数据处理")
    steps = [
        ("数据采集", "力信号\n位移信号\n温度/压力\n采样频率 100 Hz"),
        ("加速度计算", "位移二次微分\n滤波处理\n得到 a(t)"),
        ("摩擦系数计算", "μ = (sinθ - a/g) / cosθ\n代入角度与加速度\n计算 μ"),
        ("不确定度评估", "A类不确定度\nB类不确定度\n合成标准不确定度\n扩展不确定度 U"),
    ]
    for i, (h, body) in enumerate(steps):
        x = 1.18 + i * 3.45
        txt(slide, x + 0.50, 2.17, 1.46, 0.20, h, 11.4, TEAL, True, PP_ALIGN.CENTER)
        if i == 2:
            txt(slide, x + 0.25, 3.95, 1.95, 0.36, "μ = (sinθ - a/g)\n      / cosθ", 12.5, DARK, False, PP_ALIGN.CENTER)
            txt(slide, x + 0.50, 5.18, 1.50, 0.32, "代入角度与加速度\n计算 μ", 9.0, DARK)
        else:
            for j, item in enumerate(body.split("\n")):
                bullet(slide, x + 0.40, 5.00 + j * 0.30, item, 9.0)


def s09(slide, bg):
    add_bg(slide, bg[9])
    title(slide, 9, "结果分析")
    txt(slide, 2.35, 1.18, 3.25, 0.22, "摩擦系数对比趋势（示例）", 12.2, TEAL, True, PP_ALIGN.CENTER)
    txt(slide, 9.05, 1.46, 1.30, 0.20, "数据解读", 12.6, TEAL, True)
    for i, item in enumerate(["温度升高，μ整体下降", "粗糙度增大，μ上升", "压强增大，先升后降", "材质差异显著，钛合金综合表现最优"]):
        bullet(slide, 9.05, 2.05 + i * 0.48, item, 10.2)
    txt(slide, 9.35, 5.82, 2.70, 0.30, "结果与文献趋势一致，\n验证方法的有效性", 10.2, TEAL, True, PP_ALIGN.CENTER)


def s10(slide, bg):
    add_bg(slide, bg[10])
    title(slide, 10, "总结展望")
    txt(slide, 2.25, 1.12, 1.00, 0.18, "创新点", 11.8, TEAL, True)
    for i, item in enumerate(["模块化多工况装置设计", "低温稳定控制与高精度测量", "多因素系统研究与量化分析", "不确定度评估与可靠性提升"]):
        txt(slide, 2.08, 1.76 + i * 0.86, 3.50, 0.18, item, 10.5, DARK)
    txt(slide, 10.40, 1.12, 1.70, 0.18, "工程应用前景", 11.8, TEAL, True)
    for i, (h, b) in enumerate([("极地船舶设计优化", "降低航行阻力"), ("表面工程与减阻开发", "提升抗冰性能"), ("极地装备与结构安全", "提供理论依据")]):
        y = 1.82 + i * 1.22
        txt(slide, 11.10, y, 2.15, 0.18, h, 10.5, TEAL, True)
        txt(slide, 11.10, y + 0.30, 1.60, 0.16, b, 9.0, MUTED)
    txt(slide, 2.00, 7.40, 9.80, 0.18, "未来将拓展更多材料与工况，建立摩擦数据库，服务极地工程与安全航行。", 9.0, DARK)


def build():
    bg = build_backgrounds()
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    builders = [s01, s02, s03, s04, s05, s06, s07, s08, s09, s10]
    for builder in builders:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        builder(slide, bg)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
