from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
PROJECT = ROOT / "projects" / "slide09_whitepaper"
OUT = PROJECT / "exports" / "slide09_whitepaper_page09_interactive.pptx"

W, H = 16, 9
FONT = "Microsoft YaHei"

PAPER = RGBColor(248, 250, 250)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(0, 72, 84)
TEAL_2 = RGBColor(0, 96, 112)
TEAL_LIGHT = RGBColor(226, 243, 247)
ICE_BLUE = RGBColor(75, 172, 198)
DARK = RGBColor(28, 44, 49)
MUTED = RGBColor(88, 108, 114)
LINE = RGBColor(211, 225, 229)
ORANGE = RGBColor(226, 126, 39)
GREEN = RGBColor(54, 134, 76)
GRAY_FILL = RGBColor(243, 246, 246)


DATA_CONFIG = {
    "temperature": {
        "id": "F01",
        "cn": "温度",
        "en": "Temperature",
        "glyph": "℃",
        "chart_type": "line",
        "x_label": "Temperature / ℃",
        "levels": [
            {"level": "T1", "value": "-15 ℃", "mu": 0.118},
            {"level": "T2", "value": "-10 ℃", "mu": 0.103},
            {"level": "T3", "value": "-5 ℃", "mu": 0.089},
            {"level": "T4", "value": "0 ℃", "mu": 0.076},
        ],
        "conclusion": "温度升高时，冰表面液膜增强，摩擦系数 μ 整体下降。",
    },
    "pressure": {
        "id": "F02",
        "cn": "压强",
        "en": "Pressure",
        "glyph": "P",
        "chart_type": "line",
        "x_label": "Pressure / MPa",
        "levels": [
            {"level": "P1", "value": "0.05 MPa", "mu": 0.082},
            {"level": "P2", "value": "0.10 MPa", "mu": 0.096},
            {"level": "P3", "value": "0.20 MPa", "mu": 0.091},
        ],
        "conclusion": "压强增大改变局部接触状态与压力熔化程度，摩擦系数可能呈现先升后降趋势。",
    },
    "material": {
        "id": "F03",
        "cn": "金属材质",
        "en": "Metal Material",
        "glyph": "▰",
        "chart_type": "bar",
        "x_label": "Material",
        "levels": [
            {"level": "M1", "value": "铝合金", "mu": 0.108},
            {"level": "M2", "value": "不锈钢", "mu": 0.095},
            {"level": "M3", "value": "钛合金", "mu": 0.083},
        ],
        "conclusion": "不同金属材料的导热性、表面能和微观接触状态不同，导致摩擦系数存在显著差异。",
    },
    "roughness": {
        "id": "F04",
        "cn": "表面粗糙度",
        "en": "Surface Roughness",
        "glyph": "≈",
        "chart_type": "line",
        "x_label": "Roughness Ra / μm",
        "levels": [
            {"level": "R1", "value": "Ra 0.2 μm", "mu": 0.074},
            {"level": "R2", "value": "Ra 1.0 μm", "mu": 0.096},
            {"level": "R3", "value": "Ra 3.0 μm", "mu": 0.124},
        ],
        "conclusion": "表面粗糙度增大时，微观嵌入和犁削效应增强，摩擦系数 μ 上升。",
    },
    "area": {
        "id": "F05",
        "cn": "接触面积",
        "en": "Contact Area",
        "glyph": "□",
        "chart_type": "line",
        "x_label": "Contact Area / cm²",
        "levels": [
            {"level": "A1", "value": "1 cm²", "mu": 0.112},
            {"level": "A2", "value": "4 cm²", "mu": 0.096},
            {"level": "A3", "value": "9 cm²", "mu": 0.087},
        ],
        "conclusion": "在其他条件不变时，接触面积变化会改变单位压强和局部液膜状态，使摩擦系数发生变化。",
    },
    "ice": {
        "id": "F06",
        "cn": "冰的成分",
        "en": "Ice Composition",
        "glyph": "⌬",
        "chart_type": "bar",
        "x_label": "Ice Composition",
        "levels": [
            {"level": "I1", "value": "淡水冰", "mu": 0.104},
            {"level": "I2", "value": "海冰", "mu": 0.089},
            {"level": "I3", "value": "人工盐冰", "mu": 0.078},
        ],
        "conclusion": "冰中盐分或杂质会改变冰点与表面液膜状态，从而影响金属-冰界面的摩擦行为。",
    },
}


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=False, transparency=0, name=None):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    if name:
        shp.name = name
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.fill.transparency = transparency
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    return shp


def connector(slide, x1, y1, x2, y2, color=LINE, width=1.0, name=None):
    shp = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    if name:
        shp.name = name
    shp.line.color.rgb = color
    shp.line.width = Pt(width)
    return shp


def text(slide, x, y, w, h, value, size, color=DARK, bold=False, align=PP_ALIGN.LEFT, name=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for idx, raw in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = raw
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def para_tab(slide, x, y, w, h, fill=TEAL, name=None):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shp.name = name
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = fill
    return shp


def page_base(slide):
    rect(slide, 0, 0, W, H, PAPER, name="slide09_background_texture")
    for i in range(12):
        x = 0.9 + i * 1.22
        connector(slide, x, 1.1, x + 0.4, 7.75, RGBColor(232, 240, 242), 0.35, "slide09_background_grid")
    for i in range(7):
        y = 1.25 + i * 0.92
        connector(slide, 0.85, y, 15.15, y, RGBColor(232, 240, 242), 0.35, "slide09_background_grid")

    para_tab(slide, -0.10, 0.08, 1.35, 0.66)
    text(slide, 0.18, 0.20, 0.58, 0.32, "09", 22, WHITE, True, name="slide09_title_index")
    text(slide, 1.42, 0.22, 3.0, 0.36, "结果分析", 23.5, TEAL, True, name="slide09_title_main")
    connector(slide, 0.72, 0.86, 1.28, 0.86, TEAL, 3)
    connector(slide, 1.40, 0.86, 2.62, 0.86, TEAL_2, 1.4)
    text(
        slide,
        1.42,
        0.98,
        6.6,
        0.28,
        "控制变量法下的摩擦系数变化规律  /  Controlled Variable Analysis of Metal-Ice Friction Coefficient",
        11.8,
        MUTED,
        name="slide09_title_sub",
    )
    para_tab(slide, 15.12, 8.50, 0.80, 0.36)
    text(slide, 15.36, 8.58, 0.27, 0.12, "09", 10, WHITE, True, PP_ALIGN.CENTER, name="S09_footer")
    text(
        slide,
        0.92,
        8.48,
        8.4,
        0.16,
        "注：当前数据为版式示例，后续可替换为真实实验数据。",
        7.5,
        MUTED,
        name="slide09_data_source_note",
    )


def add_factor_node(slide, key, x, y):
    item = DATA_CONFIG[key]
    prefix = f"factor_node_{key}"
    rect(slide, x, y, 2.23, 1.30, WHITE, LINE, True, name=f"{prefix}_card")
    rect(slide, x + 0.12, y + 0.12, 0.53, 0.53, TEAL_LIGHT, RGBColor(183, 216, 224), True, name=f"{prefix}_icon_bg")
    text(slide, x + 0.20, y + 0.25, 0.38, 0.18, item["glyph"], 12.2, TEAL, True, PP_ALIGN.CENTER, name=f"{prefix}_icon")
    text(slide, x + 0.78, y + 0.18, 0.52, 0.18, item["id"], 9.5, ICE_BLUE, True, name=f"{prefix}_id")
    text(slide, x + 0.78, y + 0.43, 1.15, 0.20, item["cn"], 11.8, TEAL, True, name=f"{prefix}_cn")
    text(slide, x + 0.78, y + 0.72, 1.32, 0.20, item["en"], 7.8, MUTED, name=f"{prefix}_en")
    levels = " · ".join([row["level"] for row in item["levels"]])
    text(slide, x + 0.18, y + 1.04, 1.78, 0.12, levels, 7.3, MUTED, name=f"{prefix}_levels")


def add_matrix(slide):
    text(slide, 0.92, 1.34, 2.2, 0.24, "六因素模型框架", 13.5, TEAL, True, name="slide09_factor_matrix_title")
    text(slide, 3.50, 1.34, 2.5, 0.18, "点击节点查看变量分析态", 8.8, MUTED, name="slide09_factor_matrix_hint")
    rect(slide, 0.85, 1.72, 6.65, 5.86, WHITE, LINE, True, name="slide09_factor_matrix")
    text(slide, 3.52, 4.20, 1.25, 0.22, "μ", 25, TEAL, True, PP_ALIGN.CENTER, name="slide09_factor_matrix_center_mu")
    rect(slide, 3.30, 4.08, 1.04, 0.58, TEAL_LIGHT, RGBColor(183, 216, 224), True, name="slide09_factor_matrix_center")
    positions = {
        "temperature": (1.18, 2.05),
        "pressure": (4.92, 2.05),
        "material": (1.18, 3.78),
        "roughness": (4.92, 3.78),
        "area": (1.18, 5.51),
        "ice": (4.92, 5.51),
    }
    anchors = [(2.30, 2.70), (5.98, 2.70), (2.30, 4.43), (5.98, 4.43), (2.30, 6.16), (5.98, 6.16)]
    for ax, ay in anchors:
        connector(slide, 3.82, 4.37, ax, ay, RGBColor(196, 218, 224), 1.0, name="slide09_matrix_connector")
    for key, (x, y) in positions.items():
        add_factor_node(slide, key, x, y)


def add_overview(slide):
    rect(slide, 8.00, 1.42, 6.95, 6.16, WHITE, LINE, True, name="slide09_experiment_model_overview_bg")
    text(slide, 8.38, 1.74, 3.0, 0.24, "实验手段选择与模型构建", 13.5, TEAL, True, name="slide09_experiment_model_overview")
    rect(slide, 8.38, 2.28, 2.18, 1.72, GRAY_FILL, LINE, True, name="overview_contact_stage")
    rect(slide, 8.70, 3.18, 1.45, 0.18, RGBColor(215, 224, 226), RGBColor(166, 180, 184), name="overview_ice_surface")
    rect(slide, 9.14, 2.77, 0.62, 0.42, RGBColor(188, 198, 200), RGBColor(136, 150, 154), name="overview_metal_block")
    connector(slide, 9.45, 2.78, 9.45, 2.38, TEAL, 1.2, name="overview_pressure_arrow")
    text(slide, 8.65, 3.60, 1.65, 0.16, "金属-冰接触实验场景", 8.6, MUTED, False, PP_ALIGN.CENTER, name="overview_scene_caption")
    text(
        slide,
        10.90,
        2.30,
        3.20,
        1.35,
        "以控制变量法组织实验水平，将装置设计、变量设计与数据处理串联为可替换的数据分析入口。",
        11.0,
        DARK,
        name="overview_summary",
    )
    flow = ["控制变量法", "实验水平设计", "μ 数据采集", "趋势建模", "结果分析"]
    for idx, label in enumerate(flow):
        x = 8.42 + idx * 1.22
        rect(slide, x, 4.72, 1.08, 0.56, TEAL_LIGHT if idx < 4 else TEAL, RGBColor(183, 216, 224), True, name=f"overview_flow_{idx+1}")
        text(slide, x + 0.09, 4.91, 0.90, 0.13, label, 7.9, WHITE if idx == 4 else TEAL, True, PP_ALIGN.CENTER, name=f"overview_flow_text_{idx+1}")
        if idx < len(flow) - 1:
            text(slide, x + 1.10, 4.88, 0.18, 0.14, "→", 9.5, MUTED, True, PP_ALIGN.CENTER, name=f"overview_flow_arrow_{idx+1}")
    rect(slide, 8.42, 6.08, 5.95, 0.62, RGBColor(247, 250, 250), LINE, True, name="overview_note_box")
    text(
        slide,
        8.74,
        6.27,
        5.2,
        0.18,
        "初始态保留模型总览；节点触发后切换为实验水平表、趋势图与物理结论。",
        9.6,
        MUTED,
        name="overview_note_text",
    )


def add_level_table(slide, key, x, y, w, h):
    item = DATA_CONFIG[key]
    rows = len(item["levels"]) + 1
    table_shape = slide.shapes.add_table(rows, 3, Inches(x), Inches(y), Inches(w), Inches(h))
    table_shape.name = f"table_group_{key}"
    table = table_shape.table
    headers = ["水平", "变量取值", "μ"]
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = TEAL
    for row_idx, row in enumerate(item["levels"], 1):
        table.cell(row_idx, 0).text = row["level"]
        table.cell(row_idx, 1).text = row["value"]
        table.cell(row_idx, 2).text = f'{row["mu"]:.3f}'
        for col in range(3):
            table.cell(row_idx, col).fill.solid()
            table.cell(row_idx, col).fill.fore_color.rgb = WHITE
    for row in table.rows:
        for cell in row.cells:
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.name = FONT
                    r.font.size = Pt(7.4)
                    r.font.color.rgb = WHITE if cell.fill.fore_color.rgb == TEAL else DARK
    return table_shape


def add_native_chart(slide, key, x, y, w, h):
    item = DATA_CONFIG[key]
    data = CategoryChartData()
    data.categories = [row["value"] for row in item["levels"]]
    data.add_series("μ", [row["mu"] for row in item["levels"]])
    chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED if item["chart_type"] == "bar" else XL_CHART_TYPE.LINE_MARKERS
    frame = slide.shapes.add_chart(chart_type, Inches(x), Inches(y), Inches(w), Inches(h), data)
    frame.name = f"chart_group_{key}"
    chart = frame.chart
    chart.has_legend = False
    chart.value_axis.minimum_scale = 0.04
    chart.value_axis.maximum_scale = 0.14
    chart.value_axis.major_unit = 0.02
    chart.value_axis.tick_labels.font.size = Pt(7)
    chart.category_axis.tick_labels.font.size = Pt(7)
    chart.chart_title.has_text_frame = True
    chart.chart_title.text_frame.text = item["x_label"]
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.name = FONT
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format = "0.000"
    plot.data_labels.font.size = Pt(7)
    series = plot.series[0]
    series.format.line.color.rgb = ICE_BLUE
    series.format.line.width = Pt(1.6)
    return frame


def add_analysis_panel(slide, key):
    item = DATA_CONFIG[key]
    prefix = f"analysis_group_{key}"
    rect(slide, 8.00, 1.42, 6.95, 6.16, WHITE, LINE, True, name=f"{prefix}_bg")
    text(slide, 8.35, 1.72, 0.55, 0.18, item["id"], 9.5, ICE_BLUE, True, name=f"{prefix}_id")
    text(slide, 8.90, 1.66, 2.8, 0.26, f'{item["cn"]} / {item["en"]}', 13.3, TEAL, True, name=f"{prefix}_title")
    text(slide, 13.65, 1.72, 0.64, 0.18, "TRIGGER", 6.5, MUTED, False, PP_ALIGN.RIGHT, name=f"{prefix}_trigger_label")
    text(slide, 8.35, 2.23, 1.35, 0.18, "实验水平表", 9.3, TEAL, True, name=f"{prefix}_table_title")
    add_level_table(slide, key, 8.35, 2.52, 2.25, 1.45)
    text(slide, 10.95, 2.23, 1.80, 0.18, "摩擦系数 μ 趋势", 9.3, TEAL, True, name=f"{prefix}_chart_title")
    add_native_chart(slide, key, 10.75, 2.46, 3.62, 2.35)
    rect(slide, 8.35, 4.52, 2.25, 1.05, RGBColor(247, 250, 250), LINE, True, name=f"{prefix}_meta_box")
    text(slide, 8.58, 4.75, 1.82, 0.45, f"变量：{item['cn']}\n因变量：摩擦系数 μ", 8.8, DARK, name=f"{prefix}_meta_text")
    rect(slide, 10.75, 5.10, 3.62, 0.80, TEAL_LIGHT, RGBColor(183, 216, 224), True, name=f"conclusion_{key}_bg")
    text(slide, 11.05, 5.34, 3.02, 0.32, item["conclusion"], 9.4, TEAL, True, PP_ALIGN.CENTER, name=f"conclusion_{key}")


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    page_base(slide)
    add_matrix(slide)
    for key in DATA_CONFIG:
        add_analysis_panel(slide, key)
    add_overview(slide)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
