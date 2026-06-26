from pathlib import Path

from PIL import Image, ImageEnhance
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "outputs"
ASSET_DIR = OUT_DIR / "keynote_pptx_assets_v3"
SAMPLE_ASSET_DIR = OUT_DIR / "keynote_style_samples_assets"
OUT = OUT_DIR / "metal_ice_keynote_style_samples_v1.pptx"

W, H = 13.333333, 7.5

NAVY = RGBColor(8, 59, 102)
BLUE = RGBColor(0, 91, 153)
TEAL = RGBColor(0, 105, 137)
ICE = RGBColor(222, 239, 248)
ICE_2 = RGBColor(238, 247, 252)
PAPER = RGBColor(246, 251, 254)
WHITE = RGBColor(255, 255, 255)
SILVER = RGBColor(169, 184, 198)
MIST = RGBColor(232, 243, 249)
GRID = RGBColor(218, 232, 241)
GRAY = RGBColor(65, 85, 100)
SOFT = RGBColor(116, 145, 162)
ORANGE = RGBColor(242, 107, 42)


def rgb(hex_value: str) -> RGBColor:
    hex_value = hex_value.strip("#")
    return RGBColor(int(hex_value[0:2], 16), int(hex_value[2:4], 16), int(hex_value[4:6], 16))


def ensure_assets() -> None:
    SAMPLE_ASSET_DIR.mkdir(exist_ok=True)


def crop_asset(source_name: str, out_name: str, box: tuple[float, float, float, float], cool: bool = False) -> Path:
    src = Image.open(ASSET_DIR / source_name).convert("RGB")
    sw, sh = src.size
    left, top, right, bottom = box
    crop = src.crop((int(sw * left), int(sh * top), int(sw * right), int(sh * bottom)))
    if cool:
        crop = ImageEnhance.Color(crop).enhance(0.82)
        crop = ImageEnhance.Contrast(crop).enhance(1.06)
        overlay = Image.new("RGB", crop.size, (238, 248, 253))
        crop = Image.blend(crop, overlay, 0.08)
    out = SAMPLE_ASSET_DIR / out_name
    crop.save(out, quality=95)
    return out


def add_text(slide, text, x, y, w, h, size, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei UI"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_multiline(slide, lines, x, y, w, h, size=11.5, color=GRAY, bold=False, gap=1.05):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    for idx, line_text in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line_text
        p.font.name = "Microsoft YaHei UI"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.line_spacing = gap
    return shape


def add_rect(slide, x, y, w, h, fill, line=None, transparency=0, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    return shape


def add_line(slide, x1, y1, x2, y2, color=GRID, width=0.8):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def add_picture(slide, path, x, y, w, h):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def background(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PAPER
    add_rect(slide, 0, 0, W, H, rgb("F5FAFD"))
    add_rect(slide, 0, 0, W, 0.13, rgb("E2F0F7"))
    for y in [1.45, 3.95, 6.35]:
        add_line(slide, 0.42, y, 12.9, y - 0.18, rgb("E7F2F8"), 0.45)


def footer(slide, page):
    add_line(slide, 0.56, 7.03, 12.78, 7.03, rgb("C8DDEA"), 0.7)
    icon = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(0.58), Inches(7.14), Inches(0.24), Inches(0.18))
    icon.fill.solid()
    icon.fill.fore_color.rgb = NAVY
    icon.line.fill.background()
    add_text(slide, "中国大学生机械工程创新创意大赛", 0.9, 7.12, 3.4, 0.2, 7.8, color=SOFT)
    add_text(slide, f"{page:02d} / 10", 11.88, 7.11, 0.8, 0.2, 8.5, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)


def page_mark(slide, page, x=0.58, y=0.36, accent=False):
    if accent:
        add_rect(slide, x - 0.08, y + 0.08, 0.035, 0.46, ORANGE)
    add_text(slide, f"{page:02d}", x, y, 1.1, 0.54, 31, bold=True, color=NAVY)
    add_rect(slide, x + 0.01, y + 0.68, 0.55, 0.045, NAVY)


def callout(slide, title, body, x, y, w, icon_text):
    add_rect(slide, x, y, w, 0.68, WHITE, rgb("D4E4EE"), transparency=10, radius=True)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.14), Inches(y + 0.15), Inches(0.34), Inches(0.34))
    oval.fill.solid()
    oval.fill.fore_color.rgb = ICE
    oval.line.color.rgb = BLUE
    add_text(slide, icon_text, x + 0.18, y + 0.22, 0.25, 0.12, 8.5, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    add_text(slide, title, x + 0.62, y + 0.13, w - 0.8, 0.18, 10.2, bold=True, color=NAVY)
    add_text(slide, body, x + 0.62, y + 0.38, w - 0.8, 0.16, 7.6, color=GRAY)


def slide_01(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    hero = crop_asset("s01_hero_clean.png", "sample_01_hero.png", (0.06, 0.02, 1.0, 0.98), cool=True)
    add_picture(slide, hero, 4.92, 0.34, 8.18, 6.5)
    add_rect(slide, 0, 0, 5.55, 7.03, WHITE, transparency=6)
    add_rect(slide, 4.72, 0.34, 1.45, 6.5, WHITE, transparency=38)
    page_mark(slide, 1)
    add_text(slide, "金属与冰\n摩擦系数的测量", 0.62, 1.3, 4.65, 1.45, 30, bold=True, color=NAVY)
    add_text(slide, "极地环境下摩擦行为的实验研究", 0.64, 3.15, 4.3, 0.26, 13.2, color=GRAY)
    add_rect(slide, 0.64, 3.66, 3.65, 0.035, ORANGE)
    callout(slide, "参赛项目", "极地工程与材料摩擦实验", 0.64, 4.25, 3.95, "A")
    callout(slide, "作品编号", "2025-ICE-001", 0.64, 5.06, 3.95, "ID")
    add_text(slide, "Metal / Ice Friction Measurement", 0.65, 6.48, 3.6, 0.18, 8.3, bold=True, color=SOFT)
    footer(slide, 1)


def slide_05(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    apparatus = crop_asset("s05_apparatus.png", "sample_05_apparatus_no_text.png", (0.0, 0.0, 0.78, 1.0), cool=True)
    add_picture(slide, apparatus, 4.0, 0.92, 6.25, 5.62)
    add_rect(slide, 0.0, 0.0, 4.55, 7.03, WHITE, transparency=7)
    add_rect(slide, 9.68, 0.72, 2.72, 5.88, WHITE, rgb("D7E7F0"), transparency=15, radius=True)
    page_mark(slide, 5)
    add_text(slide, "装置设计", 0.62, 1.2, 3.45, 0.5, 30, bold=True, color=NAVY)
    add_text(slide, "可切换双模式冰面摩擦测试装置", 0.64, 1.85, 3.65, 0.28, 12.5, bold=True, color=TEAL)
    add_multiline(
        slide,
        [
            "通过倾斜冰面与竖直冰面两种结构，分别控制重力分量与法向压力。",
            "装置保留传感、温控与角度调节接口，为后续变量实验提供稳定平台。",
        ],
        0.66,
        2.55,
        3.18,
        0.92,
        size=10.6,
        color=GRAY,
    )
    points = [
        ("倾斜冰面模式", "通过坡角控制重力分量", "01"),
        ("竖直冰面模式", "模拟法向压力变化", "02"),
        ("传感与温控", "同步记录位移、温度与压力", "03"),
    ]
    for idx, (title, body, no) in enumerate(points):
        callout(slide, title, body, 0.66, 3.95 + idx * 0.76, 3.35, no)
    labels = [
        ("金属滑块", 8.3, 3.15, 10.0, 2.05),
        ("透明冰面", 7.0, 3.85, 10.0, 3.32),
        ("角度调节机构", 5.3, 2.1, 10.0, 4.58),
        ("力 / 位移传感器", 5.72, 4.72, 10.0, 5.42),
    ]
    for text, x1, y1, x2, y2 in labels:
        add_line(slide, x1, y1, x2 - 0.12, y2 + 0.08, BLUE, 0.85)
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x1 - 0.035), Inches(y1 - 0.035), Inches(0.07), Inches(0.07))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ORANGE
        dot.line.fill.background()
        add_text(slide, text, x2, y2, 1.6, 0.18, 9.5, bold=True, color=NAVY)
    footer(slide, 5)


def style_chart(chart, y_min, y_max):
    chart.has_legend = False
    chart.category_axis.tick_labels.font.size = Pt(7)
    chart.category_axis.tick_labels.font.color.rgb = GRAY
    chart.value_axis.tick_labels.font.size = Pt(7)
    chart.value_axis.tick_labels.font.color.rgb = GRAY
    chart.value_axis.minimum_scale = y_min
    chart.value_axis.maximum_scale = y_max
    chart.value_axis.major_gridlines.format.line.color.rgb = GRID
    chart.value_axis.major_gridlines.format.line.width = Pt(0.45)
    plot = chart.plots[0]
    plot.has_data_labels = False
    series = chart.series[0]
    series.format.line.color.rgb = BLUE
    series.format.line.width = Pt(1.6)
    series.marker.style = XL_MARKER_STYLE.CIRCLE
    series.marker.size = 5
    series.marker.format.fill.solid()
    series.marker.format.fill.fore_color.rgb = WHITE
    series.marker.format.line.color.rgb = BLUE
    series.marker.format.line.width = Pt(1)


def add_native_chart(slide, title, x, y, w, h, categories, values, x_label, y_label, note, y_min, y_max):
    add_rect(slide, x, y, w, h, WHITE, rgb("D9E9F2"), transparency=8, radius=True)
    add_text(slide, title, x + 0.22, y + 0.18, w - 0.44, 0.2, 10.2, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series("摩擦系数 μ", values)
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(x + 0.22),
        Inches(y + 0.55),
        Inches(w - 0.44),
        Inches(h - 1.12),
        chart_data,
    )
    chart = frame.chart
    style_chart(chart, y_min, y_max)
    add_text(slide, x_label, x + 0.86, y + h - 0.42, w - 1.72, 0.12, 7.2, color=SOFT, align=PP_ALIGN.CENTER)
    add_text(slide, y_label, x + 0.18, y + 0.47, 0.55, 0.12, 7.2, color=SOFT)
    add_text(slide, note, x + 0.24, y + h - 0.19, w - 0.48, 0.13, 7.5, bold=True, color=GRAY, align=PP_ALIGN.CENTER)


def slide_09(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    add_rect(slide, 0.0, 0.0, W, 7.03, WHITE, transparency=20)
    page_mark(slide, 9, accent=True)
    add_text(slide, "结果分析", 1.62, 0.62, 3.0, 0.48, 28, bold=True, color=NAVY)
    add_text(slide, "摩擦系数随温度、压强与粗糙度的变化趋势", 1.64, 1.22, 5.8, 0.25, 12.2, bold=True, color=TEAL)
    add_rect(slide, 8.92, 0.72, 3.3, 0.52, ICE, rgb("CEE4F0"), transparency=8, radius=True)
    add_text(slide, "数据为样张示意值，可在后续替换为真实实验数据", 9.15, 0.9, 2.82, 0.14, 8.5, color=GRAY)
    add_native_chart(
        slide,
        "μ - 温度关系",
        0.72,
        2.08,
        3.72,
        3.15,
        ["-30", "-25", "-20", "-15", "-10", "-5", "0"],
        [0.18, 0.16, 0.13, 0.10, 0.075, 0.055, 0.040],
        "温度 (°C)",
        "μ",
        "温度升高，摩擦系数降低",
        0,
        0.2,
    )
    add_native_chart(
        slide,
        "μ - 压强关系",
        4.82,
        2.08,
        3.72,
        3.15,
        ["0.05", "0.10", "0.20", "0.30", "0.40", "0.50"],
        [0.060, 0.075, 0.095, 0.120, 0.140, 0.155],
        "压强 (MPa)",
        "μ",
        "压强增大，接触状态改变",
        0,
        0.18,
    )
    add_native_chart(
        slide,
        "μ - 粗糙度关系",
        8.92,
        2.08,
        3.72,
        3.15,
        ["0.02", "0.05", "0.10", "0.30", "0.80", "1.50"],
        [0.045, 0.060, 0.080, 0.110, 0.145, 0.175],
        "Ra (μm)",
        "μ",
        "粗糙度增大，摩擦系数上升",
        0,
        0.2,
    )
    add_rect(slide, 0.78, 5.86, 11.78, 0.72, ICE, rgb("CFE2ED"), transparency=8, radius=True)
    add_text(slide, "温度是主要影响因素；压强与粗糙度改变接触状态，材料与成分差异在低温环境下进一步放大。", 1.1, 6.1, 10.92, 0.22, 13.2, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    footer(slide, 9)


def main():
    ensure_assets()
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide_01(prs)
    slide_05(prs)
    slide_09(prs)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
