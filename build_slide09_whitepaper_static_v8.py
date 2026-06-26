from pathlib import Path
from math import cos, sin, pi

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
PROJECT = ROOT / "projects" / "slide09_whitepaper_static_v8"
ASSET_DIR = PROJECT / "assets"
OUT_DIR = PROJECT / "exports"
VALIDATION_DIR = PROJECT / "validation"
OUT = OUT_DIR / "slide09_whitepaper_static_v8.pptx"

W, H = 16, 9
FONT = "Microsoft YaHei"

PAPER = RGBColor(246, 250, 251)
WHITE = RGBColor(255, 255, 255)
NAVY = RGBColor(5, 38, 56)
NAVY_2 = RGBColor(7, 51, 70)
INK = RGBColor(8, 45, 58)
MUTED = RGBColor(92, 115, 125)
ICE = RGBColor(215, 245, 251)
ICE_2 = RGBColor(226, 248, 252)
CYAN = RGBColor(96, 215, 238)
BLUE_LINE = RGBColor(150, 198, 213)
SOFT_LINE = RGBColor(216, 232, 237)
ORANGE = RGBColor(223, 138, 66)
STEEL = RGBColor(128, 151, 160)

FACTORS = [
    {
        "id": "F01",
        "cn": "温度",
        "en": "Temperature",
        "level": "-15 / -10 / -5 / 0 ℃",
        "note": "相态 / 融水",
        "glyph": "T",
    },
    {
        "id": "F02",
        "cn": "压强",
        "en": "Pressure",
        "level": "0.05 / 0.10 / 0.20 MPa",
        "note": "法向载荷",
        "glyph": "P",
    },
    {
        "id": "F03",
        "cn": "金属材质",
        "en": "Metal Material",
        "level": "Al / Steel / Ti",
        "note": "硬度 / 导热",
        "glyph": "M",
    },
    {
        "id": "F04",
        "cn": "表面粗糙度",
        "en": "Surface Roughness",
        "level": "Ra 0.2 / 1.0 / 3.0 μm",
        "note": "微凸体 / 润滑膜",
        "glyph": "Ra",
    },
    {
        "id": "F05",
        "cn": "接触面积",
        "en": "Contact Area",
        "level": "1 / 4 / 9 cm²",
        "note": "载荷分布",
        "glyph": "A",
    },
    {
        "id": "F06",
        "cn": "冰的成分",
        "en": "Ice Composition",
        "level": "淡水冰 / 海冰 / 盐冰",
        "note": "盐度 / 晶粒",
        "glyph": "Ice",
    },
]


def make_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)

    bg = Image.new("RGB", (2048, 1152), (246, 250, 251))
    d = ImageDraw.Draw(bg, "RGBA")
    for x in range(0, 2048, 64):
        d.line([(x, 0), (x, 1152)], fill=(7, 70, 92, 13), width=1)
    for y in range(0, 1152, 64):
        d.line([(0, y), (2048, y)], fill=(7, 70, 92, 10), width=1)
    for x in range(0, 2048, 256):
        d.line([(x, 0), (x, 1152)], fill=(7, 70, 92, 18), width=1)
    for y in range(0, 1152, 256):
        d.line([(0, y), (2048, y)], fill=(7, 70, 92, 14), width=1)
    for box, alpha, width in [
        ((550, -260, 1520, 710), 22, 4),
        ((660, -150, 1410, 600), 16, 3),
        ((120, 980, 460, 1320), 14, 3),
    ]:
        d.ellipse(box, outline=(96, 215, 238, alpha), width=width)
    d.line([(1600, 80), (1990, 80)], fill=(5, 38, 56, 55), width=4)
    d.line([(1600, 80), (1600, 220)], fill=(5, 38, 56, 28), width=2)
    bg = bg.filter(ImageFilter.GaussianBlur(0.25))
    bg_path = ASSET_DIR / "technical_paper_background_v6.png"
    bg.save(bg_path, quality=96)

    sample = ROOT / "output" / "imagegen" / "slide09_factor_matrix_premium_sample_v1.png"
    if sample.exists():
        im = Image.open(sample).convert("RGB")
        # Crop the right top apparatus region from the approved style sample.
        apparatus = im.crop((1540, 80, 1990, 430))
        canvas = Image.new("RGB", (900, 520), (247, 250, 251))
        scale = min(820 / apparatus.width, 430 / apparatus.height)
        apparatus = apparatus.resize((int(apparatus.width * scale), int(apparatus.height * scale)), Image.Resampling.LANCZOS)
        x = (900 - apparatus.width) // 2
        y = (520 - apparatus.height) // 2
        canvas.paste(apparatus, (x, y))
    else:
        canvas = Image.new("RGB", (900, 520), (247, 250, 251))
        d = ImageDraw.Draw(canvas, "RGBA")
        d.rectangle((160, 240, 750, 310), fill=(220, 246, 250, 255), outline=(120, 170, 182, 190), width=3)
        d.rectangle((365, 165, 535, 240), fill=(124, 140, 146, 255), outline=(60, 82, 90, 220), width=3)
        d.line((450, 80, 450, 165), fill=(5, 38, 56, 220), width=5)
    overlay = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    od = ImageDraw.Draw(overlay, "RGBA")
    od.rectangle((0, 0, 899, 519), outline=(150, 198, 213, 110), width=2)
    od.line((28, 32, 135, 32), fill=(5, 38, 56, 115), width=4)
    od.line((28, 32, 28, 112), fill=(5, 38, 56, 55), width=2)
    apparatus_path = ASSET_DIR / "apparatus_evidence_v6.png"
    Image.alpha_composite(canvas.convert("RGBA"), overlay).convert("RGB").save(apparatus_path, quality=97)
    return bg_path, apparatus_path


def set_line(shp, color, width=0.75, transparency=0):
    shp.line.color.rgb = color
    shp.line.width = Pt(width)
    shp.line.transparency = transparency


def shape(slide, shape_type, x, y, w, h, fill=None, line=None, width=0.75, trans=0, name=None):
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shp.name = name
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        shp.fill.transparency = trans
    if line is None:
        shp.line.fill.background()
    else:
        set_line(shp, line, width)
    return shp


def rect(slide, x, y, w, h, fill=None, line=None, width=0.75, trans=0, name=None):
    return shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, w, h, fill, line, width, trans, name)


def oval(slide, x, y, w, h, fill=None, line=BLUE_LINE, width=0.75, trans=0, name=None):
    return shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, y, w, h, fill, line, width, trans, name)


def add_line(slide, x1, y1, x2, y2, color=BLUE_LINE, width=0.75, transparency=0, name=None):
    shp = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    if name:
        shp.name = name
    shp.line.color.rgb = color
    shp.line.width = Pt(width)
    shp.line.transparency = transparency
    return shp


def textbox(slide, x, y, w, h, value, size, color=INK, bold=False, align=PP_ALIGN.LEFT, name=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for idx, raw in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = raw
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def add_header(slide):
    tab = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, Inches(-0.08), Inches(0.0), Inches(1.45), Inches(0.70))
    tab.fill.solid()
    tab.fill.fore_color.rgb = NAVY
    tab.line.fill.background()
    textbox(slide, 0.18, 0.18, 0.56, 0.30, "09", 22, WHITE, True, PP_ALIGN.CENTER, "slide09_page_no")
    textbox(slide, 1.55, 0.20, 1.70, 0.28, "结果分析", 22, INK, True, name="slide09_title_cn")
    textbox(slide, 3.25, 0.255, 4.85, 0.18, "Controlled Variable Framework for Friction Coefficient", 9.2, MUTED, False, name="slide09_title_en")
    add_line(slide, 1.55, 0.68, 2.68, 0.68, NAVY, 2.0)
    add_line(slide, 2.83, 0.68, 6.52, 0.68, BLUE_LINE, 0.75, 25)


def add_hex_icon(slide, cx, cy, r, line_color=BLUE_LINE):
    pts = []
    for i in range(6):
        a = pi / 6 + i * pi / 3
        pts.append((cx + r * cos(a), cy + r * sin(a)))
    for i in range(6):
        x1, y1 = pts[i]
        x2, y2 = pts[(i + 1) % 6]
        add_line(slide, x1, y1, x2, y2, line_color, 0.55, 16)
    for x, y in pts[::2]:
        add_line(slide, cx, cy, x, y, line_color, 0.45, 28)
    oval(slide, cx - 0.025, cy - 0.025, 0.05, 0.05, NAVY_2, None)


def add_factor_node(slide, cx, cy, data, name, label_pos="bottom"):
    # Precision circular module plus editable labels.
    oval(slide, cx - 0.58, cy - 0.58, 1.16, 1.16, WHITE, NAVY_2, 1.05, 0, f"{name}_outer")
    oval(slide, cx - 0.42, cy - 0.42, 0.84, 0.84, ICE_2, BLUE_LINE, 0.55, 12000, f"{name}_inner")
    add_hex_icon(slide, cx, cy - 0.08, 0.19, BLUE_LINE)
    oval(slide, cx - 0.055, cy + 0.38, 0.11, 0.11, ORANGE, WHITE, 0.55, name=f"{name}_measure_dot")
    textbox(slide, cx - 0.31, cy + 0.02, 0.62, 0.11, data["id"], 5.8, CYAN, True, PP_ALIGN.CENTER, f"{name}_id")


def add_factor_label(slide, x, y, w, data, align=PP_ALIGN.LEFT, name="factor_label"):
    rect(slide, x, y, w, 0.58, WHITE, SOFT_LINE, 0.45, 22000, f"{name}_box")
    rect(slide, x, y, 0.035, 0.58, NAVY, None, name=f"{name}_bar")
    textbox(slide, x + 0.12, y + 0.09, 0.30, 0.08, data["id"], 5.4, CYAN, True, align, f"{name}_id")
    textbox(slide, x + 0.44, y + 0.06, 0.72, 0.12, data["cn"], 8.2, INK, True, align, f"{name}_cn")
    textbox(slide, x + 0.44, y + 0.24, 0.92, 0.09, data["en"], 4.8, MUTED, False, align, f"{name}_en")
    textbox(slide, x + 0.44, y + 0.39, w - 0.54, 0.08, data["level"], 4.8, STEEL, False, align, f"{name}_level")


def add_matrix(slide):
    rect(slide, 0.65, 1.10, 10.65, 6.78, WHITE, SOFT_LINE, 0.8, 9000, "factor_matrix_panel")
    rect(slide, 0.65, 1.10, 10.65, 0.07, NAVY, None)
    textbox(slide, 0.93, 1.38, 1.30, 0.14, "PRIMARY MODEL", 6.8, MUTED, True, name="matrix_kicker")
    textbox(slide, 0.93, 1.62, 2.28, 0.22, "六因素控制变量矩阵", 14.5, INK, True, name="matrix_title")
    textbox(slide, 7.22, 1.39, 3.50, 0.36, "变量围绕 μ 响应量组织；实验装置仅提供测量语境。", 7.6, MUTED, False, PP_ALIGN.RIGHT, "matrix_subtitle")

    cx, cy = 5.95, 4.24
    for r, alpha, width in [(2.62, 30, 0.7), (2.06, 42, 0.65), (1.38, 55, 0.65), (0.86, 45, 0.55)]:
        shp = oval(slide, cx - r, cy - r, 2 * r, 2 * r, None, BLUE_LINE, width)
        shp.line.transparency = alpha
    for deg in [0, 30, 60, 90, 120, 150]:
        a = deg * pi / 180
        add_line(slide, cx - 2.66 * cos(a), cy - 2.66 * sin(a), cx + 2.66 * cos(a), cy + 2.66 * sin(a), BLUE_LINE, 0.42, 64)
    for deg in range(0, 360, 15):
        a = deg * pi / 180
        r1 = 2.54
        r2 = 2.62 if deg % 45 else 2.70
        add_line(slide, cx + r1 * cos(a), cy + r1 * sin(a), cx + r2 * cos(a), cy + r2 * sin(a), BLUE_LINE, 0.35, 45)

    core_size = 1.22
    oval(slide, cx - 0.80, cy - 0.80, 1.60, 1.60, ICE_2, BLUE_LINE, 0.65, 45000, "response_outer_halo")
    oval(slide, cx - 0.61, cy - 0.61, core_size, core_size, NAVY, BLUE_LINE, 1.0, 0, "response_core")
    textbox(slide, cx - 0.30, cy - 0.38, 0.60, 0.36, "μ", 34, WHITE, True, PP_ALIGN.CENTER, "response_mu")
    textbox(slide, cx - 0.38, cy + 0.20, 0.76, 0.10, "response", 5.8, ICE, False, PP_ALIGN.CENTER, "response_label")

    positions = [
        (cx, cy - 2.05, "bottom"),
        (cx + 1.95, cy - 1.10, "bottom"),
        (cx + 1.95, cy + 1.10, "top"),
        (cx, cy + 2.05, "top"),
        (cx - 1.95, cy + 1.10, "top"),
        (cx - 1.95, cy - 1.10, "bottom"),
    ]
    for idx, (nx, ny, _label_pos) in enumerate(positions):
        add_line(slide, cx, cy, nx, ny, BLUE_LINE, 0.55, 42, f"connector_{idx+1}")
    for idx, ((nx, ny, label_pos), data) in enumerate(zip(positions, FACTORS), start=1):
        add_factor_node(slide, nx, ny, data, f"factor_{idx:02d}", label_pos)

    label_specs = [
        (4.86, 1.84, 2.15, FACTORS[0], "factor_label_01"),
        (8.42, 3.12, 2.05, FACTORS[1], "factor_label_02"),
        (8.42, 5.15, 2.05, FACTORS[2], "factor_label_03"),
        (4.86, 6.54, 2.15, FACTORS[3], "factor_label_04"),
        (1.32, 5.15, 2.05, FACTORS[4], "factor_label_05"),
        (1.32, 3.12, 2.05, FACTORS[5], "factor_label_06"),
    ]
    for lx, ly, lw, data, lname in label_specs:
        add_factor_label(slide, lx, ly, lw, data, name=lname)

    # Subtle small charts around the panel, drawn as editable lines.
    add_line(slide, 0.92, 6.98, 2.12, 6.98, BLUE_LINE, 0.5, 50)
    add_line(slide, 0.92, 6.98, 0.92, 6.18, BLUE_LINE, 0.5, 50)
    for i in range(5):
        x = 1.02 + i * 0.22
        y = 6.88 - i * i * 0.022
        oval(slide, x - 0.018, y - 0.018, 0.036, 0.036, ORANGE if i == 2 else NAVY_2, None)
        if i:
            px = 1.02 + (i - 1) * 0.22
            py = 6.88 - (i - 1) * (i - 1) * 0.022
            add_line(slide, px, py, x, y, BLUE_LINE, 0.45, 44)

    footer_y = 7.50
    labels = ["控制变量", "实验水平", "μ 数据采集", "趋势拟合", "结果解释", "模型回写"]
    step_w = 1.48
    x0 = 0.92
    add_line(slide, x0, footer_y - 0.12, x0 + step_w * 6, footer_y - 0.12, SOFT_LINE, 0.55)
    for i, label in enumerate(labels):
        x = x0 + i * step_w
        if i:
            add_line(slide, x, footer_y - 0.12, x, footer_y + 0.38, SOFT_LINE, 0.45)
        textbox(slide, x + 0.08, footer_y + 0.00, 0.24, 0.10, f"{i+1:02d}", 5.0, CYAN, True)
        textbox(slide, x + 0.08, footer_y + 0.19, 1.10, 0.12, label, 6.2, INK, True)


def add_right_column(slide, apparatus_path):
    add_line(slide, 11.68, 1.10, 11.68, 7.88, SOFT_LINE, 0.9)
    textbox(slide, 12.02, 1.22, 2.62, 0.24, "实验手段选择 + 模型构建", 13.4, INK, True, name="right_title")
    textbox(slide, 12.02, 1.60, 2.88, 0.38, "作为证据与方法层，解释六因素矩阵如何落到实验方案和建模链路。", 7.5, MUTED, False, name="right_intro")

    pic = slide.shapes.add_picture(str(apparatus_path), Inches(12.02), Inches(2.15), width=Inches(2.72), height=Inches(1.55))
    pic.name = "apparatus_evidence_image"
    textbox(slide, 12.23, 3.80, 2.30, 0.12, "金属-冰接触实验语境（辅助证据层）", 6.7, MUTED, False, PP_ALIGN.CENTER, "apparatus_caption")

    x_axis = 12.32
    add_line(slide, x_axis, 4.32, x_axis, 7.25, BLUE_LINE, 0.65, 32)
    methods = [
        ("01", "控制变量法", "一次只改变一个因素，保持其余条件稳定。"),
        ("02", "实验水平设计", "为每个因素设置离散水平，形成可比较序列。"),
        ("03", "μ 数据采集", "从接触界面测量摩擦响应并记录波动。"),
        ("04", "趋势建模", "拟合单因素趋势并识别主导影响区间。"),
        ("05", "结果分析", "回到物理机制解释并支撑结论表达。"),
    ]
    y = 4.30
    for idx, (num, title, desc) in enumerate(methods):
        oval(slide, x_axis - 0.055, y + 0.02, 0.11, 0.11, WHITE, BLUE_LINE, 0.75)
        oval(slide, x_axis - 0.025, y + 0.05, 0.05, 0.05, CYAN if idx < 4 else ORANGE, None)
        textbox(slide, 12.52, y - 0.02, 0.30, 0.10, num, 5.8, CYAN, True)
        textbox(slide, 12.88, y - 0.04, 1.34, 0.12, title, 7.6, INK, True)
        textbox(slide, 12.88, y + 0.16, 2.04, 0.22, desc, 5.8, MUTED, False)
        y += 0.58

    rect(slide, 12.02, 7.42, 2.88, 0.36, WHITE, SOFT_LINE, 0.55, 9000)
    textbox(slide, 12.20, 7.53, 2.52, 0.10, "静态版仅呈现模型总览；交互热区将在确认后单独添加。", 5.8, MUTED, False, PP_ALIGN.CENTER)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    VALIDATION_DIR.mkdir(parents=True, exist_ok=True)
    bg, apparatus = make_assets()
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(bg), Inches(0), Inches(0), width=Inches(W), height=Inches(H)).name = "technical_background_texture"
    add_header(slide)
    add_matrix(slide)
    add_right_column(slide, apparatus)
    textbox(slide, 0.75, 8.45, 6.60, 0.13, "注：当前数据为版式示例，后续可替换为真实实验数据；本页未添加触发器动画。", 6.2, MUTED, False, name="slide09_footer_note")
    end_tab = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, Inches(14.64), Inches(8.34), Inches(0.88), Inches(0.34))
    end_tab.fill.solid()
    end_tab.fill.fore_color.rgb = NAVY
    end_tab.line.fill.background()
    textbox(slide, 14.87, 8.43, 0.36, 0.12, "09", 8.2, WHITE, True, PP_ALIGN.CENTER)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
