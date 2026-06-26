from pathlib import Path
from pptx import Presentation

p = Path(r"F:\Users\user\Documents\测金属摩擦系数\outputs\metal_ice_keynote_editable.pptx")
prs = Presentation(p)
print("slides", len(prs.slides))
print("size", prs.slide_width, prs.slide_height)
for i, s in enumerate(prs.slides, 1):
    text = sum(1 for sh in s.shapes if getattr(sh, "has_text_frame", False) and sh.has_text_frame)
    pics = sum(1 for sh in s.shapes if sh.shape_type == 13)
    print(i, "shapes", len(s.shapes), "text", text, "pictures", pics)
