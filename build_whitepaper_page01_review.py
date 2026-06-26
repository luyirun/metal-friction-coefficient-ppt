from pathlib import Path

from PIL import Image, ImageEnhance
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
PREVIEW = ROOT / "outputs" / "whitepaper_preview_rebuild" / "page_01_preview.png"
ASSET_DIR = ROOT / "outputs" / "whitepaper_page_reviews" / "assets"
OUT_DIR = ROOT / "outputs" / "whitepaper_page_reviews"
OUT = OUT_DIR / "page_01_review.pptx"

W, H = 16, 9
FONT = "Microsoft YaHei"

PAPER = RGBColor(247, 250, 250)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(0, 72, 84)
TEAL_DARK = RGBColor(0, 56, 68)
DARK = RGBColor(28, 44, 49)
MUTED = RGBColor(72, 94, 101)
LINE = RGBColor(203, 222, 228)
ORANGE = RGBColor(226, 126, 39)


def crop_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    src = Image.open(PREVIEW).convert("RGB")

    hero = src.crop((545, 0, 1536, 805))
    hero_path = ASSET_DIR / "page01_hero_apparatus.png"
    hero.save(hero_path, quality=96)

    inset = src.crop((530, 505, 760, 735))
    inset_path = ASSET_DIR / "page01_inset.png"
    inset.save(inset_path, quality=96)

    bg = src.crop((0, 0, 1536, 864))
    bg = ImageEnhance.Contrast(bg).enhance(0.78)
    bg = ImageEnhance.Color(bg).enhance(0.55)
    wash = Image.new("RGB", bg.size, (247, 250, 250))
    bg = Image.blend(wash, bg, 0.28)
    bg_path = ASSET_DIR / "page01_atmosphere_wash.png"
    bg.save(bg_path, quality=94)
    return {"hero": hero_path, "inset": inset_path, "bg": bg_path}


def rect(slide, x, y, w, h, fill, line=None, radius=False, transparency=0):
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.fill.transparency = transparency
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    return shp


def text(slide, x, y, w, h, value, size, color=DARK, bold=False, align=PP_ALIGN.LEFT, name=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for idx, raw in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = raw
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def line(slide, x1, y1, x2, y2, color=TEAL, width=1.2):
    shp = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shp.line.color.rgb = color
    shp.line.width = Pt(width)
    return shp


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    assets = crop_assets()
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    rect(slide, 0, 0, W, H, PAPER)
    slide.shapes.add_picture(str(assets["bg"]), Inches(0), Inches(0), width=Inches(W), height=Inches(H))
    rect(slide, 0, 0, W, H, PAPER, transparency=18000)

    slide.shapes.add_picture(str(assets["hero"]), Inches(5.90), Inches(0.00), width=Inches(9.95), height=Inches(7.42))
    rect(slide, 5.20, 0, 1.55, H, PAPER, transparency=35000)

    rect(slide, -0.06, 0.08, 1.48, 0.72, TEAL_DARK)
    text(slide, 0.20, 0.21, 0.62, 0.32, "01", 24, WHITE, True, name="S01_page_number")

    rect(slide, 0.60, 1.42, 5.35, 2.35, PAPER, transparency=6000)
    text(slide, 0.72, 1.55, 5.10, 1.30, "金属与冰\n摩擦系数的测量", 34, TEAL, True, name="S01_main_title")
    line(slide, 0.75, 3.50, 1.34, 3.50, ORANGE, 2.4)
    line(slide, 1.46, 3.50, 2.82, 3.50, LINE, 1.2)
    text(slide, 0.75, 3.78, 3.70, 0.30, "物理实验创新竞赛答辩", 16.5, DARK, True, name="S01_subtitle")

    rect(slide, 0.70, 4.70, 3.55, 0.48, PAPER, transparency=6000)
    text(slide, 1.18, 4.84, 2.15, 0.20, "团队：极冰钥匙者", 11.5, DARK, name="S01_team")
    rect(slide, 0.70, 5.42, 3.55, 0.48, PAPER, transparency=6000)
    text(slide, 1.18, 5.56, 2.15, 0.20, "日期：2024.05.20", 11.5, DARK, name="S01_date")

    slide.shapes.add_picture(str(assets["inset"]), Inches(4.48), Inches(4.72), width=Inches(1.60), height=Inches(1.60))

    rect(slide, 0.66, 7.92, 5.25, 0.40, TEAL_DARK)
    text(slide, 0.98, 8.03, 4.55, 0.16, "物理实验  |  摩擦学研究  |  精密测量  |  工程创新", 9.5, WHITE, True, name="S01_bottom_band")
    rect(slide, 15.08, 8.51, 0.82, 0.36, TEAL_DARK)
    text(slide, 15.34, 8.59, 0.28, 0.12, "01", 10, WHITE, True, PP_ALIGN.CENTER, name="S01_footer_page")

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
