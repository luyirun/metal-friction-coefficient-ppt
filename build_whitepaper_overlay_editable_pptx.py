from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "outputs"
SRC_DIR = OUT_DIR / "single_pages"
OUT = OUT_DIR / "metal_ice_whitepaper_overlay_editable.pptx"

W, H = 16, 9
FONT = "Microsoft YaHei"

PAPER = RGBColor(247, 250, 250)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(0, 72, 84)
TEAL_2 = RGBColor(0, 94, 108)
DARK = RGBColor(28, 44, 49)
MUTED = RGBColor(80, 101, 108)
LINE = RGBColor(218, 229, 232)
ORANGE = RGBColor(226, 126, 39)
GREEN = RGBColor(61, 137, 78)
SOFT = RGBColor(246, 249, 249)
ICE = RGBColor(230, 245, 248)


def add_bg(slide, page):
    slide.shapes.add_picture(
        str(SRC_DIR / f"page_{page:02d}_industrial_whitepaper.png"),
        Inches(0),
        Inches(0),
        width=Inches(W),
        height=Inches(H),
    )


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=False, transparency=0):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.fill.transparency = transparency
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    return shp


def para_tab(slide, x, y, w, h, fill=TEAL):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = fill
    return shp


def line(slide, x1, y1, x2, y2, color=TEAL, width=1.2):
    shp = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shp.line.color.rgb = color
    shp.line.width = Pt(width)
    return shp


def txt(slide, x, y, w, h, value, size, color=DARK, bold=False, align=PP_ALIGN.LEFT, name=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, raw in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = raw
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def bullet(slide, x, y, value, size=11.5, color=DARK):
    rect(slide, x, y + 0.095, 0.055, 0.055, TEAL, radius=True)
    txt(slide, x + 0.18, y, 4.3, 0.22, value, size, color)


def wipe(slide, boxes):
    for box in boxes:
        if len(box) == 4:
            rect(slide, *box, WHITE)
        else:
            x, y, w, h, color = box
            rect(slide, x, y, w, h, color)


def page_badge(slide, page):
    para_tab(slide, -0.10, 0.05, 1.30, 0.66)
    txt(slide, 0.18, 0.17, 0.58, 0.28, f"{page:02d}", 22, WHITE, True)
    para_tab(slide, 15.12, 8.50, 0.80, 0.36)
    txt(slide, 15.36, 8.58, 0.28, 0.12, f"{page:02d}", 10, WHITE, True, PP_ALIGN.CENTER)


def header(slide, page, title, subtitle=None):
    rect(slide, 0, 0, W, 0.92, PAPER)
    page_badge(slide, page)
    txt(slide, 1.34, 0.20, 6.8, 0.34, title, 23, TEAL, True, name=f"S{page:02d}_title")
    line(slide, 0.70, 0.86, 1.23, 0.86, TEAL, 3)
    line(slide, 1.35, 0.86, 2.50, 0.86, TEAL_2, 1.5)
    if subtitle:
        txt(slide, 1.34, 0.96, 6.9, 0.22, subtitle, 13, MUTED, name=f"S{page:02d}_subtitle")


def cover_slide(slide):
    add_bg(slide, 1)
    wipe(slide, [
        (0, 0, 1.55, 0.80, PAPER),
        (0.70, 1.50, 5.20, 1.65, RGBColor(247, 250, 250)),
        (0.74, 3.35, 3.55, 0.35, RGBColor(247, 250, 250)),
        (0.70, 4.20, 2.95, 1.40, RGBColor(247, 250, 250)),
        (0.76, 7.94, 3.70, 0.33, TEAL),
        (15.10, 8.52, 0.78, 0.34, TEAL),
    ])
    page_badge(slide, 1)
    txt(slide, 0.78, 1.52, 4.55, 1.25, "金属与冰\n摩擦系数的测量", 32, TEAL, True, name="S01_main_title")
    txt(slide, 0.78, 3.35, 3.2, 0.28, "物理实验创新竞赛答辩", 15.5, DARK, True, name="S01_subtitle")
    line(slide, 0.78, 3.78, 1.36, 3.78, TEAL, 2.4)
    line(slide, 1.48, 3.78, 2.75, 3.78, LINE, 1.4)
    txt(slide, 1.17, 4.34, 2.10, 0.20, "团队：祝冰间隙者", 11.4, DARK, name="S01_team")
    txt(slide, 1.17, 5.18, 2.10, 0.20, "日期：2024.05.20", 11.4, DARK, name="S01_date")
    txt(slide, 1.02, 8.04, 3.3, 0.16, "精准测量  ·  理解机理  ·  服务工程", 9.5, WHITE, True, name="S01_footer_text")


def slide02(slide):
    add_bg(slide, 2)
    header(slide, 2, "选题背景")
    wipe(slide, [
        (8.90, 1.40, 5.30, 3.30, WHITE),
        (0.96, 5.25, 6.15, 1.85, WHITE),
        (15.10, 8.52, 0.78, 0.34, TEAL),
    ])
    txt(slide, 9.00, 1.48, 4.60, 0.30, "北极航运与冰阻力问题", 18, TEAL, True, name="S02_side_title")
    txt(slide, 9.00, 2.15, 3.80, 0.18, "航行阻力中，冰障碍阻力占比", 11.4, MUTED, name="S02_desc1")
    txt(slide, 9.00, 2.48, 2.15, 0.32, "25%-55%", 22, ORANGE, True, name="S02_num1")
    line(slide, 9.00, 3.32, 13.65, 3.32, LINE, 1)
    txt(slide, 9.00, 3.82, 4.60, 0.18, "冰面相关事故中，摩擦失控问题占比", 11.4, MUTED, name="S02_desc2")
    txt(slide, 9.00, 4.14, 2.15, 0.32, "13%-15%", 22, ORANGE, True, name="S02_num2")
    rect(slide, 1.04, 5.42, 5.95, 1.42, WHITE, LINE, radius=True)
    txt(slide, 1.35, 5.58, 1.15, 0.20, "研究意义", 12.5, TEAL, True, name="S02_research_title")
    for i, item in enumerate(["揭示船舶破冰与航行效率的关键因素之一", "为船舶材料选型、表面处理与结构设计提供依据", "推动极地装备与工程安全的进步"]):
        bullet(slide, 1.26, 5.92 + i * 0.32, item, 10.4)
    page_badge(slide, 2)


def slide03(slide):
    add_bg(slide, 3)
    header(slide, 3, "题目解读")
    wipe(slide, [
        (1.05, 1.20, 13.90, 5.95, WHITE),
        (15.10, 8.52, 0.78, 0.34, TEAL),
    ])
    sections = [
        ("机理分析", "冰的黏附性变弱\n表面微观水膜\n水膜润滑作用\n温度与压力影响", "□"),
        ("装置搭建", "斜面摩擦实验平台\n力学测量与采集\n低温环境控制\n稳定与校准", "⚒"),
        ("结果与不确定度", "摩擦系数计算\n不确定度评估\n数据性分析\n结果讨论与验证", "↗"),
    ]
    for i, (title, body, glyph) in enumerate(sections):
        x = 1.10 + i * 4.55
        rect(slide, x, 1.30, 4.10, 5.70, WHITE, LINE)
        txt(slide, x + 0.35, 1.62, 3.40, 0.24, title, 15.4, TEAL, True, PP_ALIGN.CENTER, name=f"S03_title_{i+1}")
        rect(slide, x + 1.55, 2.35, 1.00, 1.00, ICE, RGBColor(180, 214, 222), radius=True)
        txt(slide, x + 1.82, 2.65, 0.45, 0.25, glyph, 20, TEAL, True, PP_ALIGN.CENTER, name=f"S03_icon_{i+1}")
        for j, item in enumerate(body.split("\n")):
            bullet(slide, x + 0.60, 4.02 + j * 0.42, item, 11.4)
    page_badge(slide, 3)


def slide04(slide):
    add_bg(slide, 4)
    header(slide, 4, "方法选择 —— 动态斜面法")
    wipe(slide, [
        (0.98, 1.15, 2.0, 0.35, WHITE),
        (8.66, 1.15, 5.00, 5.95, WHITE),
        (15.10, 8.52, 0.78, 0.34, TEAL),
    ])
    txt(slide, 1.10, 1.26, 1.4, 0.22, "力学模型", 13.5, TEAL, True, name="S04_model_title")
    txt(slide, 8.78, 1.26, 1.4, 0.22, "理论公式", 13.5, TEAL, True, name="S04_formula_title")
    rect(slide, 9.10, 1.78, 4.30, 0.66, RGBColor(242, 244, 244), None, radius=True)
    txt(slide, 9.48, 1.98, 3.50, 0.22, "a = g(sinθ - μcosθ)", 16.4, DARK, False, PP_ALIGN.CENTER, name="S04_formula1")
    txt(slide, 11.10, 2.65, 0.25, 0.18, "↓", 18, TEAL_2, True, PP_ALIGN.CENTER)
    rect(slide, 9.10, 3.10, 4.30, 0.86, RGBColor(242, 244, 244), None, radius=True)
    txt(slide, 9.45, 3.35, 3.60, 0.24, "μ = (sinθ - a/g) / cosθ", 16.4, DARK, False, PP_ALIGN.CENTER, name="S04_formula2")
    txt(slide, 9.20, 4.34, 3.7, 1.20, "a —— 沿斜面方向的加速度\ng —— 重力加速度\nθ —— 斜面倾角\nμ —— 动摩擦系数", 11.6, DARK, name="S04_notes")
    page_badge(slide, 4)


def slide05(slide):
    add_bg(slide, 5)
    header(slide, 5, "装置设计")
    wipe(slide, [
        (0.88, 1.05, 2.40, 0.36, WHITE),
        (8.22, 1.05, 4.10, 0.36, WHITE),
        (0.72, 1.72, 1.25, 3.82, WHITE),
        (12.80, 1.65, 1.55, 3.95, WHITE),
        (0.78, 7.48, 13.90, 0.48, WHITE),
        (15.10, 8.52, 0.78, 0.34, TEAL),
    ])
    txt(slide, 1.00, 1.12, 2.15, 0.22, "斜面冰面模块", 13, TEAL, True, name="S05_left_title")
    txt(slide, 8.35, 1.12, 3.40, 0.22, "垂直冰面模块（压力加载）", 13, TEAL, True, name="S05_right_title")
    for i, label in enumerate(["导轨与滑块", "力传感器", "位移传感器", "冰面板", "角度调节机构", "支撑框架"]):
        txt(slide, 0.72, 1.85 + i * 0.48, 1.20, 0.18, label, 9.6, MUTED, name=f"S05_left_label{i+1}")
    for i, label in enumerate(["压力加载装置", "力传感器", "冰面板", "温控仓体", "底盘"]):
        txt(slide, 13.12, 1.82 + i * 0.70, 1.28, 0.18, label, 9.6, MUTED, name=f"S05_right_label{i+1}")
    for i, label in enumerate(["模块化设计", "快速更换", "低温环境模拟", "高精度测量"]):
        txt(slide, 1.15 + i * 3.35, 7.72, 1.25, 0.16, label, 10.2, TEAL, True, PP_ALIGN.CENTER, name=f"S05_feature{i+1}")
    page_badge(slide, 5)


def slide06(slide):
    add_bg(slide, 6)
    header(slide, 6, "变量设计", "六大影响变量与水平设置")
    wipe(slide, [
        (0.92, 3.05, 14.20, 2.90, WHITE),
        (15.10, 8.52, 0.78, 0.34, TEAL),
    ])
    items = [
        ("温度", "-5 ℃\n-10 ℃\n-15 ℃"),
        ("压强", "0.05 MPa\n0.10 MPa\n0.20 MPa"),
        ("材质", "铝合金\n不锈钢\n钛合金"),
        ("粗糙度", "Ra 0.2 μm\nRa 1.0 μm\nRa 3.0 μm"),
        ("面积", "1 cm²\n4 cm²\n9 cm²"),
        ("成分", "淡水冰\n海冰\n人工掺盐冰"),
    ]
    for i, (title, body) in enumerate(items):
        x = 1.25 + i * 2.40
        txt(slide, x - 0.45, 3.37, 0.90, 0.20, title, 12.4, TEAL, True, PP_ALIGN.CENTER, name=f"S06_title_{i+1}")
        txt(slide, x - 0.62, 4.15, 1.24, 0.82, body, 10.8, DARK, False, PP_ALIGN.CENTER, name=f"S06_body_{i+1}")
    page_badge(slide, 6)


def slide07(slide):
    add_bg(slide, 7)
    header(slide, 7, "问题改进")
    wipe(slide, [
        (1.35, 1.50, 3.70, 4.85, RGBColor(247, 250, 250)),
        (7.32, 1.50, 3.55, 4.85, RGBColor(247, 250, 250)),
        (15.10, 8.52, 0.78, 0.34, TEAL),
    ])
    rows = [
        ("冰面制备不平整", "表面粗糙和气泡大\n影响测量重复性", "改进方案", "定制制冰机 + 控温体系\n实现光滑透明平整冰面"),
        ("低温环境稳定性差", "温度波动导致数据漂移\n影响实验结果", "改进方案", "双层保温 + PID温控\n温度波动 ≤ ±0.2 ℃"),
        ("传感器可靠性不足", "低温漂移与噪声干扰大\n影响信号准确性", "改进方案", "低温型传感器 + 屏蔽布线\n多点校准与滤波处理"),
    ]
    for i, (p, desc, st, sol) in enumerate(rows):
        y = 1.66 + i * 1.72
        txt(slide, 1.72, y, 2.40, 0.20, p, 11.8, DARK, True, name=f"S07_problem{i+1}")
        txt(slide, 1.72, y + 0.30, 2.52, 0.40, desc, 10.0, MUTED, name=f"S07_desc{i+1}")
        txt(slide, 7.58, y, 1.20, 0.20, st, 11.8, GREEN, True, name=f"S07_solution_title{i+1}")
        txt(slide, 7.58, y + 0.30, 2.70, 0.40, sol, 10.0, DARK, name=f"S07_solution{i+1}")
    page_badge(slide, 7)


def slide08(slide):
    add_bg(slide, 8)
    header(slide, 8, "数据处理")
    wipe(slide, [
        (1.08, 2.00, 13.82, 4.45, WHITE),
        (15.10, 8.52, 0.78, 0.34, TEAL),
    ])
    steps = [
        ("数据采集", "力信号\n位移信号\n温度/压力\n采样频率 100 Hz"),
        ("加速度计算", "位移二次微分\n滤波处理\n得到 a(t)"),
        ("摩擦系数计算", "μ = (sinθ - a/g) / cosθ\n代入角度与加速度\n计算 μ"),
        ("不确定度评估", "A类不确定度\nB类不确定度\n合成标准不确定度\n扩展不确定度 U"),
    ]
    for i, (title, body) in enumerate(steps):
        x = 1.15 + i * 3.48
        txt(slide, x + 0.55, 2.18, 1.45, 0.20, title, 11.8, TEAL, True, PP_ALIGN.CENTER, name=f"S08_title_{i+1}")
        if i == 2:
            txt(slide, x + 0.28, 3.95, 1.95, 0.38, "μ = (sinθ - a/g)\n      / cosθ", 12.8, DARK, False, PP_ALIGN.CENTER, name="S08_formula")
            txt(slide, x + 0.50, 5.18, 1.50, 0.32, "代入角度与加速度\n计算 μ", 9.2, DARK, name="S08_formula_note")
        else:
            for j, item in enumerate(body.split("\n")):
                bullet(slide, x + 0.42, 5.00 + j * 0.30, item, 9.2)
    page_badge(slide, 8)


def slide09(slide):
    add_bg(slide, 9)
    header(slide, 9, "结果分析")
    wipe(slide, [
        (1.80, 1.10, 4.40, 0.32, WHITE),
        (8.80, 1.35, 4.10, 3.00, WHITE),
        (9.05, 5.55, 3.40, 0.90, RGBColor(242, 248, 248)),
        (15.10, 8.52, 0.78, 0.34, TEAL),
    ])
    txt(slide, 2.38, 1.18, 3.20, 0.22, "摩擦系数对比趋势（示例）", 12.5, TEAL, True, PP_ALIGN.CENTER, name="S09_chart_title")
    txt(slide, 9.05, 1.46, 1.30, 0.20, "数据解读", 13.0, TEAL, True, name="S09_data_title")
    for i, item in enumerate(["温度升高，μ整体下降", "粗糙度增大，μ上升", "压强增大，先升后降", "材质差异显著，钛合金综合表现最优"]):
        bullet(slide, 9.05, 2.05 + i * 0.48, item, 10.5)
    txt(slide, 9.35, 5.82, 2.70, 0.30, "结果与文献趋势一致，\n验证方法的有效性", 10.5, TEAL, True, PP_ALIGN.CENTER, name="S09_result")
    page_badge(slide, 9)


def slide10(slide):
    add_bg(slide, 10)
    header(slide, 10, "总结展望")
    wipe(slide, [
        (2.15, 1.08, 1.00, 0.25, WHITE),
        (10.35, 1.08, 2.20, 0.25, WHITE),
        (1.65, 1.55, 4.95, 4.10, WHITE),
        (11.00, 1.55, 3.30, 4.10, WHITE),
        (1.90, 7.32, 10.70, 0.32, WHITE),
        (15.10, 8.52, 0.78, 0.34, TEAL),
    ])
    txt(slide, 2.25, 1.12, 1.00, 0.20, "创新点", 12.0, TEAL, True, name="S10_innovation")
    for i, item in enumerate(["模块化多工况装置设计", "低温稳定控制与高精度测量", "多因素系统研究与量化分析", "不确定度评估与可靠性提升"]):
        txt(slide, 2.08, 1.76 + i * 0.86, 3.50, 0.18, item, 10.8, DARK, name=f"S10_point{i+1}")
    txt(slide, 10.40, 1.12, 1.70, 0.20, "工程应用前景", 12.0, TEAL, True, name="S10_app")
    apps = [("极地船舶设计优化", "降低航行阻力"), ("表面工程与减阻开发", "提升抗冰性能"), ("极地装备与结构安全", "提供理论依据")]
    for i, (title, body) in enumerate(apps):
        y = 1.82 + i * 1.22
        txt(slide, 11.10, y, 2.15, 0.18, title, 10.8, TEAL, True, name=f"S10_app_title{i+1}")
        txt(slide, 11.10, y + 0.30, 1.60, 0.16, body, 9.2, MUTED, name=f"S10_app_body{i+1}")
    txt(slide, 2.00, 7.40, 9.80, 0.18, "未来将拓展更多材料与工况，建立摩擦数据库，服务极地工程与安全航行。", 9.2, DARK, name="S10_future")
    page_badge(slide, 10)


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slides = [cover_slide, slide02, slide03, slide04, slide05, slide06, slide07, slide08, slide09, slide10]
    for builder in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        builder(slide)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
