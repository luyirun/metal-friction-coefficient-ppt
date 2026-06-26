from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(r"F:\Users\user\Documents\测金属摩擦系数")
SRC_DIR = ROOT / "outputs" / "single_pages"
ASSET_DIR = ROOT / "outputs" / "hybrid_editable_assets"
OUT = ROOT / "outputs" / "metal_ice_hybrid_high_fidelity_editable.pptx"

SLIDE_W, SLIDE_H = 16, 9
COLS, ROWS = 8, 6

FONT = "Microsoft YaHei"
TEAL = RGBColor(0, 74, 87)
DARK = RGBColor(0, 76, 89)
TEXT = RGBColor(30, 45, 50)
MUTED = RGBColor(76, 98, 105)
ORANGE = RGBColor(232, 126, 45)
WHITE = RGBColor(255, 255, 255)
BG = RGBColor(247, 250, 250)
PAPER = RGBColor(246, 249, 248)


def add_sliced_page(slide, page_no):
    src = SRC_DIR / f"page_{page_no:02d}_industrial_whitepaper.png"
    im = Image.open(src).convert("RGB")
    w, h = im.size
    tw, th = w // COLS, h // ROWS
    for r in range(ROWS):
        for c in range(COLS):
            left = c * tw
            top = r * th
            right = w if c == COLS - 1 else (c + 1) * tw
            bottom = h if r == ROWS - 1 else (r + 1) * th
            tile = im.crop((left, top, right, bottom))
            path = ASSET_DIR / f"p{page_no:02d}_{r:02d}_{c:02d}.png"
            tile.save(path, optimize=True)
            slide.shapes.add_picture(
                str(path),
                Inches(SLIDE_W * left / w),
                Inches(SLIDE_H * top / h),
                width=Inches(SLIDE_W * (right - left) / w),
                height=Inches(SLIDE_H * (bottom - top) / h),
            )


def rect(slide, x, y, w, h, fill=PAPER, line=None, transparency=0, radius=False):
    typ = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shp = slide.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.fill.transparency = transparency
    if line:
        shp.line.color.rgb = line
    else:
        shp.line.fill.background()
    return shp


def textbox(slide, x, y, w, h, text, size, color=TEXT, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def bullet_lines(slide, x, y, lines, size=12.5, gap=0.34, color=TEXT):
    for i, line in enumerate(lines):
        yy = y + i * gap
        rect(slide, x, yy + 0.10, 0.055, 0.055, TEAL, radius=True)
        textbox(slide, x + 0.18, yy, 3.6, 0.24, line, size, color)


def title_mask(slide, num, title, subtitle=None):
    rect(slide, 0.92, 0.12, 7.3, 0.9, PAPER)
    textbox(slide, 1.28, 0.22, 6.0, 0.36, title, 25, TEAL, True)
    if subtitle:
        textbox(slide, 1.28, 0.82, 6.3, 0.28, subtitle, 12.5, TEAL, False)
    rect(slide, 15.28, 8.53, 0.42, 0.20, DARK)
    textbox(slide, 15.35, 8.52, 0.35, 0.18, f"{num:02d}", 11.5, WHITE, True, PP_ALIGN.CENTER)


def page1(slide):
    rect(slide, 0.42, 1.70, 7.15, 4.65, RGBColor(246, 250, 251), transparency=12000)
    textbox(slide, 0.98, 2.08, 6.10, 1.92, "金属与冰\n摩擦系数的测量", 43, TEAL, True)
    rect(slide, 0.50, 4.05, 5.35, 1.05, PAPER)
    textbox(slide, 0.58, 4.16, 5.0, 0.32, "青少年科学素养大赛 · 实验答辩", 17, TEAL, True)
    textbox(slide, 0.58, 4.58, 5.1, 0.26, "参赛选手：金属与冰界面摩擦实验方案汇报", 11.5, MUTED)
    rect(slide, 0.75, 5.62, 4.30, 0.46, PAPER)
    textbox(slide, 1.16, 5.72, 3.45, 0.24, "参赛选手：实验项目组", 12.5, TEXT)
    rect(slide, 0.60, 8.36, 6.85, 0.48, DARK)
    textbox(slide, 0.88, 8.48, 5.4, 0.22, "实验测量 · 误差分析 · 答辩汇报", 10.5, WHITE, True)


def page2(slide):
    title_mask(slide, 2, "选题背景")
    rect(slide, 9.45, 1.40, 5.65, 4.22, WHITE)
    textbox(slide, 9.58, 1.55, 4.0, 0.32, "北极航运与冰阻力问题", 19, TEAL, True)
    textbox(slide, 9.58, 2.38, 4.2, 0.25, "航行阻力中，冰障碍阻力占比", 12.5, TEXT)
    textbox(slide, 9.58, 2.76, 2.2, 0.44, "25%-55%", 23, ORANGE, True)
    textbox(slide, 9.58, 3.72, 4.6, 0.25, "冰面相关事故中，摩擦失控问题占比", 12.5, TEXT)
    textbox(slide, 9.58, 4.10, 2.2, 0.44, "13%-15%", 23, ORANGE, True)
    rect(slide, 0.72, 5.76, 7.95, 1.90, WHITE)
    textbox(slide, 1.08, 5.94, 1.5, 0.25, "研究意义", 13.5, TEAL, True)
    bullet_lines(slide, 1.02, 6.36, ["揭示船舶破冰与航行效率的关键因素之一", "为船舶材料选型、表面处理与结构设计提供依据", "推动极地装备与工程安全的进步"], 10.5, 0.33)


def page3(slide):
    title_mask(slide, 3, "题目解读")
    cols = [
        (1.35, "机理分析", ["冰的黏附性变弱", "表面微观水膜", "水膜润滑作用", "温度与压力影响"]),
        (5.75, "装置搭建", ["斜面摩擦实验平台", "力学测量与采集", "低温环境控制", "稳定与校准"]),
        (10.15, "结果与不确定度", ["摩擦系数计算", "不确定度评估", "数据性分析", "结果讨论与验证"]),
    ]
    for x, head, lines in cols:
        rect(slide, x + 0.20, 1.42, 3.55, 5.65, RGBColor(252, 253, 253), transparency=5000)
        textbox(slide, x + 0.78, 1.55, 2.3, 0.3, head, 16, TEAL, True, PP_ALIGN.CENTER)
        bullet_lines(slide, x + 0.55, 4.05, lines, 11.5, 0.43)


def page4(slide):
    title_mask(slide, 4, "方法选择 —— 动态斜面法")
    rect(slide, 8.92, 1.18, 5.65, 5.60, PAPER)
    textbox(slide, 9.05, 1.35, 1.8, 0.28, "理论公式", 14, TEAL, True)
    rect(slide, 9.12, 1.92, 4.80, 0.68, RGBColor(241, 245, 245), radius=True)
    textbox(slide, 9.62, 2.08, 3.5, 0.28, "a = g(sinθ - μcosθ)", 17, TEXT, False, PP_ALIGN.CENTER)
    textbox(slide, 11.30, 2.92, 0.4, 0.28, "↓", 22, RGBColor(145, 165, 172), True, PP_ALIGN.CENTER)
    rect(slide, 9.12, 3.40, 4.80, 0.90, RGBColor(241, 245, 245), radius=True)
    textbox(slide, 9.52, 3.62, 3.8, 0.32, "μ = (sinθ - a/g) / cosθ", 17, TEXT, False, PP_ALIGN.CENTER)
    textbox(slide, 9.28, 4.70, 4.35, 1.38, "a —— 沿斜面方向的加速度\ng —— 重力加速度\nθ —— 斜面倾角\nμ —— 动摩擦系数", 12.5, TEXT)


def page5(slide):
    title_mask(slide, 5, "装置设计")
    rect(slide, 0.95, 1.04, 3.0, 0.36, PAPER)
    textbox(slide, 1.18, 1.13, 2.0, 0.22, "斜面冰面模块", 13.5, TEAL, True)
    rect(slide, 8.68, 1.04, 4.2, 0.36, PAPER)
    textbox(slide, 8.92, 1.13, 3.6, 0.22, "垂直冰面模块（压力加载）", 13.5, TEAL, True)
    rect(slide, 1.05, 7.08, 13.60, 0.55, PAPER)
    for i, lab in enumerate(["模块化设计", "快速更换", "低温环境模拟", "高精度测量"]):
        textbox(slide, 1.42 + i * 3.35, 7.25, 1.6, 0.20, lab, 10.5, TEAL, False, PP_ALIGN.CENTER)


def page6(slide):
    title_mask(slide, 6, "变量设计", "六大影响变量与水平设置")
    items = [
        ("温度", "-5 ℃\n-10 ℃\n-15 ℃"), ("压强", "0.05 MPa\n0.10 MPa\n0.20 MPa"), ("材质", "铝合金\n不锈钢\n钛合金"),
        ("粗糙度", "Ra 0.2 μm\nRa 1.0 μm\nRa 3.0 μm"), ("面积", "1 cm²\n4 cm²\n9 cm²"), ("成分", "淡水冰\n海冰\n人工掺盐冰"),
    ]
    for i, (head, body) in enumerate(items):
        x = 0.95 + i * 2.35
        rect(slide, x + 0.22, 3.74, 1.4, 1.58, PAPER)
        textbox(slide, x + 0.42, 3.95, 1.0, 0.24, head, 14, TEAL, True, PP_ALIGN.CENTER)
        textbox(slide, x + 0.26, 4.48, 1.34, 0.90, body, 10.5, TEXT, False, PP_ALIGN.CENTER)


def page7(slide):
    title_mask(slide, 7, "问题改进")
    rows = [
        ("冰面制备不平整", "表面粗糙和沉积过大\n影响测量重复性", "改进方案", "定制制冰机 + 控温体系\n实现光滑透明平整冰面"),
        ("低温环境稳定性差", "温度波动导致数据漂移\n影响实验结果", "改进方案", "双层保温 + PID温控\n温度波动 ≤ ±0.2 ℃"),
        ("传感器可靠性不足", "低温漂移与噪声干扰大\n影响信号准确性", "改进方案", "低温型传感器 + 屏蔽布线\n多点校准与滤波处理"),
    ]
    for i, (a, b, c, d) in enumerate(rows):
        y = 1.42 + i * 1.86
        rect(slide, 2.35, y, 3.40, 0.98, RGBColor(248, 251, 251), radius=True)
        textbox(slide, 3.35, y + 0.16, 1.9, 0.22, a, 12.5, TEAL, True)
        textbox(slide, 3.35, y + 0.48, 2.05, 0.32, b, 9.2, MUTED)
        rect(slide, 8.20, y, 4.32, 0.98, RGBColor(250, 252, 250), radius=True)
        textbox(slide, 8.55, y + 0.14, 1.4, 0.22, c, 12.5, RGBColor(75, 145, 78), True)
        textbox(slide, 8.55, y + 0.48, 3.0, 0.32, d, 9.2, TEXT)


def page8(slide):
    title_mask(slide, 8, "数据处理")
    steps = [
        ("数据采集", ["力信号", "位移信号", "温度/压力", "采样频率 100 Hz"]),
        ("加速度计算", ["位移二次微分", "滤波处理", "得到 a(t)"]),
        ("摩擦系数计算", ["代入角度与加速度", "计算 μ"]),
        ("不确定度评估", ["A类不确定度", "B类不确定度", "合成标准不确定度", "扩展不确定度 U"]),
    ]
    for i, (head, lines) in enumerate(steps):
        x = 0.92 + i * 3.65
        rect(slide, x + 0.20, 1.50, 2.55, 5.55, RGBColor(251, 253, 253), radius=True)
        textbox(slide, x + 0.58, 1.83, 1.8, 0.25, head, 12.5, TEAL, True, PP_ALIGN.CENTER)
        bullet_lines(slide, x + 0.55, 5.24, lines, 9.6, 0.32)
    rect(slide, 7.98, 3.50, 2.15, 0.82, RGBColor(250, 252, 252))
    textbox(slide, 8.08, 3.67, 1.95, 0.28, "μ = (sinθ - a/g) / cosθ", 13.5, TEXT, False, PP_ALIGN.CENTER)


def page9(slide):
    title_mask(slide, 9, "结果分析")
    rect(slide, 9.18, 1.25, 4.35, 4.60, PAPER)
    textbox(slide, 9.52, 1.52, 1.6, 0.25, "数据解读", 14, TEAL, True)
    bullet_lines(slide, 9.55, 2.18, ["温度升高，μ整体下降", "粗糙度增大，μ上升", "压强增大，先升后降", "材质差异显著，钛合金综合表现最优"], 11.5, 0.52)
    rect(slide, 9.68, 5.86, 3.42, 0.86, RGBColor(239, 247, 247), radius=True)
    textbox(slide, 10.08, 6.08, 2.65, 0.28, "结果与文献趋势一致，验证方法的有效性", 11.0, TEAL, True, PP_ALIGN.CENTER)


def page10(slide):
    title_mask(slide, 10, "总结展望")
    rect(slide, 2.18, 0.98, 1.80, 0.30, PAPER)
    textbox(slide, 2.75, 1.04, 0.9, 0.20, "创新点", 11.5, TEAL, True)
    points = ["模块化多工况装置设计", "低温稳定控制与高精度测量", "多因素系统研究与量化分析", "不确定度评估与可靠性提升"]
    for i, point in enumerate(points):
        rect(slide, 1.40, 1.44 + i * 1.13, 4.55, 0.62, WHITE, radius=True)
        textbox(slide, 2.18, 1.58 + i * 1.13, 3.2, 0.20, point, 11.5, TEXT)
    rect(slide, 10.75, 0.98, 2.20, 0.30, PAPER)
    textbox(slide, 11.05, 1.04, 1.5, 0.20, "工程应用前景", 11.5, TEAL, True)
    labels = [("极地船舶设计优化", "降低航行阻力"), ("表面工程与减阻开发", "提升抗冰性能"), ("极地装备与结构安全", "提供理论依据")]
    for i, (h, b) in enumerate(labels):
        y = 1.55 + i * 1.48
        rect(slide, 11.22, y, 3.12, 0.78, WHITE, radius=True)
        textbox(slide, 11.58, y + 0.16, 1.8, 0.20, h, 11.2, TEAL, True)
        textbox(slide, 11.58, y + 0.43, 1.5, 0.18, b, 9.4, MUTED)
    rect(slide, 1.50, 7.30, 12.80, 0.48, PAPER)
    textbox(slide, 2.18, 7.42, 10.6, 0.20, "未来将拓展更多材料与工况，建立摩擦数据库，服务极地工程与安全航行。", 10.5, TEXT)


PAGE_EDITORS = {
    1: page1, 2: page2, 3: page3, 4: page4, 5: page5,
    6: page6, 7: page7, 8: page8, 9: page9, 10: page10,
}


def build():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    for page_no in range(1, 11):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_sliced_page(slide, page_no)
        PAGE_EDITORS[page_no](slide)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
