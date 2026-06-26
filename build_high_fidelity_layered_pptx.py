from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(r"F:\Users\user\Documents\测金属摩擦系数")
SRC_DIR = ROOT / "outputs" / "single_pages"
ASSET_DIR = ROOT / "outputs" / "layered_ppt_assets"
OUT = ROOT / "outputs" / "metal_ice_high_fidelity_layered.pptx"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

SLIDE_W, SLIDE_H = 16, 9
COLS, ROWS = 4, 3


def add_sliced_page(slide, page_no):
    src = SRC_DIR / f"page_{page_no:02d}_industrial_whitepaper.png"
    im = Image.open(src).convert("RGB")
    w, h = im.size
    tile_w, tile_h = w // COLS, h // ROWS
    for r in range(ROWS):
        for c in range(COLS):
            left = c * tile_w
            top = r * tile_h
            right = w if c == COLS - 1 else (c + 1) * tile_w
            bottom = h if r == ROWS - 1 else (r + 1) * tile_h
            tile = im.crop((left, top, right, bottom))
            tile_path = ASSET_DIR / f"p{page_no:02d}_tile_{r}_{c}.png"
            tile.save(tile_path, quality=95)
            x = SLIDE_W * left / w
            y = SLIDE_H * top / h
            tw = SLIDE_W * (right - left) / w
            th = SLIDE_H * (bottom - top) / h
            slide.shapes.add_picture(str(tile_path), Inches(x), Inches(y), width=Inches(tw), height=Inches(th))


def build():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    for i in range(1, 11):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_sliced_page(slide, i)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
