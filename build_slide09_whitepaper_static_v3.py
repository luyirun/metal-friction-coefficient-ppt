from pathlib import Path
from math import cos, sin, pi

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
PROJECT = ROOT / "projects" / "slide09_whitepaper_static_v3"
ASSET_DIR = PROJECT / "assets"
OUT_DIR = PROJECT / "exports"
VALIDATION_DIR = PROJECT / "validation"
OUT = OUT_DIR / "slide09_whitepaper_static_v3.pptx"

W, H = 16, 9
FONT = "Microsoft YaHei"

PAPER = RGBColor(248, 250, 250)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(0, 72, 84)
TEAL_2 = RGBColor(0, 96, 112)
ICE = RGBColor(229, 244, 247)
ICE_2 = RGBColor(216, 237, 242)
ICE_LINE = RGBColor(196, 222, 228)
BLUE = RGBColor(75, 172, 198)
DARK = RGBColor(28, 44, 49)
MUTED = RGBColor(88, 108, 114)
LINE = RGBColor(211, 225, 229)
SOFT_LINE = RGBColor(228, 238, 241)
GRAY = RGBColor(244, 247, 247)


FACTORS = [
    ("F01", "温度", "Temperature", "℃", "-15 / -10 / -5 / 0 ℃"),
    ("F02", "压强", "Pressure", "P", "0.05 / 0.10 / 0.20 MPa"),
    ("F03", "金属材质", "Metal Material", "▰", "Al / Steel / Ti"),
    ("F04", "表面粗糙度", "Surface Roughness", "≈", "Ra 0.2 / 1.0 / 3.0 μm"),
    ("F05", "接触面积", "Contact Area", "□", "1 / 4 / 9 cm²"),
    ("F06", "冰的成分", "Ice Composition", "⌬", "淡水冰 / 海冰 / 盐冰"),
]


def make_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)

    bg = Image.new("RGB", (2048, 1152), (248, 250, 250))
    d = ImageDraw.Draw(bg, "RGBA")
    for x in range(-180, 2260, 220):
        d.line([(x, 130), (x + 110, 1060)], fill=(180, 205, 212, 18), width=2)
    for y in range(170, 1040, 170):
        d.line([(110, y), (1985, y)], fill=(180, 205, 212, 14), width=2)
    for box, alpha in [((490, -360, 1560, 710), 18), ((590, -250, 1460, 620), 14), ((710, -120, 1340, 510), 11)]:
        d.ellipse(box, outline=(75, 172, 198, alpha), width=4)
    bg = bg.filter(ImageFilter.GaussianBlur(0.45))
    texture = ASSET_DIR / "whitepaper_background_texture_v3.png"
    bg.save(texture, quality=96)

    scene = Image.new("RGB", (1400, 760), (247, 250, 250))
    d = ImageDraw.Draw(scene, "RGBA")
    for y in range(760):
        shade = int(252 - y * 0.018)
        d.line([(0, y), (1400, y)], fill=(shade, shade, shade, 255))
    d.rounded_rectangle((70, 58, 1330, 702), radius=42, outline=(207, 226, 231, 190), width=4)
    d.rounded_rectangle((150, 455, 1250, 530), radius=20, fill=(220, 238, 242, 235), outline=(148, 180, 188, 225), width=3)
    d.rectangle((205, 430, 1195, 463), fill=(241, 250, 252, 175))
    d.ellipse((500, 430, 900, 520), fill=(75, 172, 198, 28))
    d.rounded_rectangle((520, 238, 880, 432), radius=12, fill=(178, 190, 192, 255), outline=(112, 132, 136, 255), width=4)
    for i in range(22):
        x = 535 + i * 15
        d.line([(x, 252), (x + 96, 420)], fill=(236, 242, 243, 34), width=2)
    d.line((700, 95, 700, 238), fill=(0, 72, 84, 230), width=8)
    d.polygon([(678, 116), (722, 116), (700, 82)], fill=(0, 72, 84, 230))
    for x in [260, 380, 1000, 1125]:
        d.line((x, 495, x + 74, 450), fill=(255, 255, 255, 155), width=2)
    d.rectangle((145, 540, 1255, 548), fill=(215, 229, 233, 90))
    scene = scene.filter(ImageFilter.UnsharpMask(radius=1.1, percent=145, threshold=3))
    scene_path = ASSET_DIR / "metal_ice_contact_inset_v3.png"
    scene.save(scene_path, quality=97)
    return texture, scene_path


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=False, transparency=0, name=None):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    if name:
        shp.name = name
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.fill.transparency = transparency
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    return shp


def line(slide, x1, y1, x2, y2, color=LINE, width=1.0, name=None, dash=False):
    shp = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    if name:
        shp.name = name
    shp.line.color.rgb = color
    shp.line.width = Pt(width)
    if dash:
        shp.line.dash_style = 4
    return shp


def text(slide, x, y, w, h, value, size, color=DARK, bold=False, align=PP_ALIGN.LEFT, name=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for idx, raw in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = raw
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def tab(slide, x, y, w, h, fill=TEAL):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = fill
    return shp


def ellipse(slide, x, y, w, h, fill=None, outline=ICE_LINE, width=1.0, transparency=100000, name=None):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shp.name = name
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    else:
        shp.fill.background()
    shp.fill.transparency = transparency
    shp.line.color.rgb = outline
    shp.line.width = Pt(width)
    return shp


def add_header(slide):
    tab(slide, -0.10, 0.08, 1.35, 0.66)
    text(slide, 0.18, 0.20, 0.58, 0.32, "09", 22, WHITE, True, name="slide09_title_index")
    text(slide, 1.42, 0.22, 3.1, 0.36, "结果分析", 23.5, TEAL, True, name="slide09_title_main")
    line(slide, 0.72, 0.86, 1.28, 0.86, TEAL, 3)
    line(slide, 1.40, 0.86, 2.62, 0.86, TEAL_2, 1.4)
    text(
        slide, 1.42, 0.98, 7.2, 0.28,
        "控制变量法下的摩擦系数变化规律  /  Controlled Variable Analysis of Metal-Ice Friction Coefficient",
        11.6, MUTED, name="slide09_title_sub"
    )
    tab(slide, 15.12, 8.50, 0.80, 0.36)
    text(slide, 15.36, 8.58, 0.27, 0.12, "09", 10, WHITE, True, PP_ALIGN.CENTER)
    text(slide, 0.92, 8.48, 8.4, 0.16, "注：当前数据为版式示例，后续可替换为真实实验数据。", 7.5, MUTED, name="slide09_data_source_note")


def add_node(slide, fid, cn, en, glyph, levels, x, y, side):
    rect(slide, x, y, 1.92, 0.72, WHITE, ICE_LINE, True, name=f"factor_node_{fid}")
    rect(slide, x + 0.10, y + 0.13, 0.40, 0.40, ICE, ICE_LINE, True)
    text(slide, x + 0.19, y + 0.255, 0.22, 0.08, glyph, 7.7, TEAL, True, PP_ALIGN.CENTER)
    text(slide, x + 0.62, y + 0.13, 0.42, 0.10, fid, 6.9, BLUE, True)
    text(slide, x + 0.62, y + 0.30, 0.90, 0.13, cn, 9.0, TEAL, True)
    text(slide, x + 0.62, y + 0.49, 1.04, 0.09, en, 5.4, MUTED)
    text(slide, x + (0.08 if side == "left" else 0.70), y + 0.78, 1.1, 0.08, levels, 5.4, MUTED, False, PP_ALIGN.CENTER)


def add_model_framework(slide):
    text(slide, 0.92, 1.42, 2.28, 0.24, "六因素矩阵 / 模型框架", 13.2, TEAL, True)
    text(slide, 5.32, 1.44, 1.40, 0.16, "Response variable: μ", 7.6, BLUE, True, PP_ALIGN.RIGHT)
    rect(slide, 0.88, 1.82, 6.45, 5.72, WHITE, LINE, True, name="slide09_factor_matrix")
    rect(slide, 1.10, 2.08, 6.00, 5.08, GRAY, SOFT_LINE, True, transparency=24000, name="model_inner_plate")

    cx, cy = 4.08, 4.65
    ellipse(slide, cx - 2.05, cy - 1.70, 4.10, 3.40, None, ICE_LINE, 0.8, name="model_orbit_outer")
    ellipse(slide, cx - 1.52, cy - 1.20, 3.04, 2.40, None, SOFT_LINE, 0.7, name="model_orbit_inner")
    line(slide, cx - 2.10, cy, cx + 2.10, cy, ICE_LINE, 0.75)
    line(slide, cx, cy - 1.72, cx, cy + 1.72, ICE_LINE, 0.75)
    for deg in [35, 145, 215, 325]:
        a = pi * deg / 180
        line(slide, cx, cy, cx + 2.0 * cos(a), cy + 1.58 * sin(a), SOFT_LINE, 0.65)

    rect(slide, cx - 0.68, cy - 0.38, 1.36, 0.76, ICE_2, ICE_LINE, True, name="model_response_core")
    rect(slide, cx - 0.50, cy - 0.25, 1.00, 0.50, WHITE, ICE_LINE, True, transparency=16000)
    text(slide, cx - 0.36, cy - 0.20, 0.72, 0.20, "μ", 21, TEAL, True, PP_ALIGN.CENTER, name="model_response_mu")
    text(slide, cx - 0.48, cy + 0.12, 0.96, 0.08, "response", 5.8, MUTED, False, PP_ALIGN.CENTER)

    nodes = [
        (*FACTORS[0], 2.10, 2.70, "left"),
        (*FACTORS[1], 4.22, 2.70, "right"),
        (*FACTORS[5], 1.22, 4.26, "left"),
        (*FACTORS[2], 5.10, 4.26, "right"),
        (*FACTORS[4], 2.10, 5.82, "left"),
        (*FACTORS[3], 4.22, 5.82, "right"),
    ]
    anchors = [(3.06, 3.06), (5.18, 3.06), (2.18, 4.62), (6.06, 4.62), (3.06, 6.18), (5.18, 6.18)]
    for ax, ay in anchors:
        line(slide, cx, cy, ax, ay, SOFT_LINE, 0.65)
    for node in nodes:
        add_node(slide, *node)

    text(slide, 1.22, 7.20, 5.70, 0.16, "模型图用于定位控制变量与响应量关系；交互阶段再在该结构上叠加分析态。", 8.0, MUTED)


def add_right_overview(slide, scene_path):
    rect(slide, 7.86, 1.82, 7.06, 5.72, WHITE, LINE, True, name="slide09_experiment_model_overview")
    rect(slide, 8.12, 2.08, 6.54, 5.16, GRAY, SOFT_LINE, True, transparency=26000)
    text(slide, 8.26, 2.20, 3.15, 0.25, "实验手段选择 + 模型构建", 13.2, TEAL, True)
    pic = slide.shapes.add_picture(str(scene_path), Inches(8.30), Inches(2.66), width=Inches(3.18), height=Inches(1.72))
    pic.name = "slide09_metal_ice_inset_image"
    text(slide, 8.54, 4.50, 2.72, 0.15, "金属-冰接触实验场景（局部视觉资产）", 7.5, MUTED, False, PP_ALIGN.CENTER)
    text(
        slide, 11.86, 2.83, 2.35, 1.04,
        "以控制变量法组织实验水平，在同一测量框架下改变单一因素，提取 μ 的变化趋势并回到物理机理解释。",
        10.2, DARK, name="overview_method_text"
    )
    line(slide, 11.62, 2.78, 11.62, 4.18, ICE_LINE, 0.8)
    rect(slide, 11.86, 4.12, 0.72, 0.26, WHITE, ICE_LINE, True)
    text(slide, 12.00, 4.205, 0.44, 0.06, "input", 5.4, BLUE, True, PP_ALIGN.CENTER)
    rect(slide, 12.72, 4.12, 0.72, 0.26, WHITE, ICE_LINE, True)
    text(slide, 12.84, 4.205, 0.48, 0.06, "model", 5.4, BLUE, True, PP_ALIGN.CENTER)
    rect(slide, 13.58, 4.12, 0.72, 0.26, WHITE, ICE_LINE, True)
    text(slide, 13.70, 4.205, 0.48, 0.06, "output", 5.4, BLUE, True, PP_ALIGN.CENTER)

    rect(slide, 8.34, 4.86, 5.96, 0.34, WHITE, SOFT_LINE, True)
    text(slide, 8.56, 4.97, 1.20, 0.08, "变量矩阵", 5.8, MUTED, False, PP_ALIGN.CENTER)
    text(slide, 9.86, 4.97, 1.20, 0.08, "实验水平", 5.8, MUTED, False, PP_ALIGN.CENTER)
    text(slide, 11.16, 4.97, 1.20, 0.08, "μ 数据", 5.8, MUTED, False, PP_ALIGN.CENTER)
    text(slide, 12.46, 4.97, 1.20, 0.08, "物理解释", 5.8, MUTED, False, PP_ALIGN.CENTER)

    bus_y = 5.82
    line(slide, 8.55, bus_y, 14.05, bus_y, ICE_LINE, 1.0)
    labels = [("01", "控制变量法"), ("02", "水平设计"), ("03", "μ 数据采集"), ("04", "趋势建模"), ("05", "结果分析")]
    widths = [1.02, 1.02, 1.10, 1.02, 1.02]
    x = 8.36
    for i, ((num, label), w) in enumerate(zip(labels, widths)):
        fill = TEAL if i == 4 else WHITE
        fg = WHITE if i == 4 else TEAL
        rect(slide, x, 5.42, w, 0.78, fill, ICE_LINE, True, name=f"overview_step_{i+1}")
        text(slide, x + 0.10, 5.54, 0.24, 0.08, num, 5.6, BLUE if i < 4 else ICE, True)
        text(slide, x + 0.12, 5.78, w - 0.24, 0.12, label, 7.2, fg, True, PP_ALIGN.CENTER)
        x += w + 0.32

    rect(slide, 8.34, 6.62, 5.95, 0.46, WHITE, LINE, True)
    text(slide, 8.78, 6.77, 5.10, 0.13, "静态阶段只呈现分析入口：模型总览、实验方法与建模流程。", 8.4, MUTED)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    VALIDATION_DIR.mkdir(parents=True, exist_ok=True)
    texture, scene = make_assets()
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(texture), Inches(0), Inches(0), width=Inches(W), height=Inches(H)).name = "slide09_background_texture"
    add_header(slide)
    add_model_framework(slide)
    add_right_overview(slide, scene)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
