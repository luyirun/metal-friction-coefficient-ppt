from pathlib import Path
from PIL import Image

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(r"F:\Users\user\Documents\测金属摩擦系数")
SRC = ROOT / "outputs" / "keynote_single_pages"
ASSETS = ROOT / "outputs" / "keynote_pptx_assets"
OUT = ROOT / "outputs" / "metal_ice_keynote_editable.pptx"
ASSETS.mkdir(exist_ok=True)

W, H = 13.333333, 7.5

NAVY = RGBColor(5, 54, 101)
TEAL = RGBColor(0, 92, 127)
BLUE = RGBColor(19, 103, 168)
ICE = RGBColor(232, 244, 250)
LIGHT = RGBColor(246, 250, 253)
MID = RGBColor(203, 226, 238)
ORANGE = RGBColor(237, 115, 35)
GRAY = RGBColor(78, 96, 108)
WHITE = RGBColor(255, 255, 255)


def crop_asset(slide_no, name, box):
    src = Image.open(SRC / f"slide_{slide_no:02d}.png").convert("RGB")
    w, h = src.size
    x1, y1, x2, y2 = box
    box_px = (int(x1 * w), int(y1 * h), int(x2 * w), int(y2 * h))
    out = ASSETS / f"s{slide_no:02d}_{name}.png"
    src.crop(box_px).save(out)
    return out


def add_text(slide, text, x, y, w, h, size=24, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def add_multiline(slide, lines, x, y, w, h, size=13, color=GRAY, bullet=False, gap=1.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.line_spacing = gap
        p.level = 0
        if bullet:
            p.text = "• " + line
    return box


def add_round_rect(slide, x, y, w, h, fill=WHITE, line=MID, radius=True):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(1)
    return shp


def add_circle_icon(slide, x, y, label, fill=WHITE, line=BLUE, txt=BLUE):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.42), Inches(0.42))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(1.5)
    add_text(slide, label, x + 0.08, y + 0.08, 0.26, 0.2, size=10, bold=True, color=txt, align=PP_ALIGN.CENTER)
    return shp


def footer(slide, page):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.12), Inches(W), Inches(0.38))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    add_text(slide, "金属与冰摩擦系数测量", 0.65, 7.22, 3.0, 0.18, size=9, bold=False, color=WHITE)
    add_text(slide, f"{page:02d} / 10", 11.75, 7.21, 0.9, 0.2, size=11, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)


def bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    # soft ice band
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
    shp.fill.solid()
    shp.fill.fore_color.rgb = LIGHT
    shp.line.fill.background()


def title_block(slide, page, title, subtitle=None, x=0.55, y=0.35):
    add_text(slide, f"{page:02d}", x, y, 1.0, 0.45, size=34, bold=True, color=NAVY)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.02), Inches(y + 0.73), Inches(0.72), Inches(0.05))
    line.fill.solid(); line.fill.fore_color.rgb = NAVY; line.line.fill.background()
    add_text(slide, title, x, y + 1.05, 4.6, 0.5, size=25, bold=True, color=NAVY)
    if subtitle:
        add_text(slide, subtitle, x, y + 1.65, 4.6, 0.35, size=13, bold=True, color=TEAL)


def add_picture(slide, path, x, y, w, h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def slide01(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    hero = crop_asset(1, "hero", (0.56, 0.02, 1.0, 0.94))
    add_picture(s, hero, 5.25, 0.0, 8.08, 7.1)
    title_block(s, 1, "金属与冰\n摩擦系数的测量")
    add_text(s, "可重复、可调控的测量系统", 1.45, 3.92, 3.7, 0.25, size=16, bold=True, color=NAVY)
    add_multiline(s, ["围绕金属-冰接触面的滑动摩擦特性，搭建实验装置，", "分析不同条件下摩擦系数的变化规律。"], 1.45, 4.45, 4.2, 0.6, size=13, color=RGBColor(25, 54, 86))
    add_round_rect(s, 0.7, 5.38, 4.5, 0.98, fill=RGBColor(238, 246, 251), line=RGBColor(188, 214, 228))
    add_text(s, "研究要点", 0.95, 5.55, 1.2, 0.2, size=12, bold=True, color=NAVY)
    add_multiline(s, ["金属-冰接触面", "实验装置", "规律分析"], 0.98, 5.82, 2.2, 0.5, size=10, color=RGBColor(20, 51, 78), bullet=True)
    add_circle_icon(s, 0.7, 3.85, "◎")
    footer(s, 1)


def slide02(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    ship = crop_asset(2, "ship", (0.50, 0.0, 1.0, 0.78))
    add_picture(s, ship, 6.15, 0.0, 7.18, 5.95)
    title_block(s, 2, "选题背景", "冰上摩擦影响极地航运阻力与能耗")
    add_multiline(s, ["北极航道可缩短航程 25%-55%，但冰阻力会使", "燃油消耗增加约 13%-15%。金属与冰摩擦系数测量", "可为船冰阻力分析和极地装备设计提供依据。"], 0.7, 2.25, 5.0, 1.1, size=12)
    for i, (num, lab) in enumerate([("25%-55%", "北极航道可缩短航程"), ("13%-15%", "冰阻力会使燃油消耗增加")]):
        y = 3.75 + i * 0.9
        add_circle_icon(s, 0.8, y + 0.05, "↗")
        add_text(s, num, 1.45, y, 2.2, 0.28, size=23, bold=True, color=ORANGE)
        add_text(s, lab, 1.45, y + 0.38, 3.1, 0.22, size=10, color=GRAY)
    footer(s, 2)


def slide03(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title_block(s, 3, "题目解读", "完成机制分析、装置搭建、结果评估")
    add_multiline(s, ["研究任务包括分析金属与冰之间的滑动摩擦机制，设计并制作实验装置，获得摩擦系数测量结果并讨论不确定度。"], 0.8, 1.75, 11.3, 0.45, size=11, color=GRAY)
    steps = [("机制分析", "研究金属与冰之间的\n滑动摩擦机制"), ("装置搭建", "设计并制作实验装置"), ("结果评估", "获得摩擦系数测量结果\n并讨论不确定度")]
    xs = [1.45, 5.3, 9.15]
    for i, (t, b) in enumerate(steps):
        add_shape = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(xs[i]), Inches(2.95), Inches(1.25), Inches(1.25))
        add_shape.fill.solid(); add_shape.fill.fore_color.rgb = RGBColor(235, 247, 252); add_shape.line.color.rgb = BLUE
        add_text(s, ["⚙", "▣", "▥"][i], xs[i] + 0.39, 3.22, 0.45, 0.3, size=22, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
        add_text(s, t, xs[i] - 0.28, 4.45, 1.8, 0.25, size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_multiline(s, b.split("\n"), xs[i] - 0.35, 4.9, 1.95, 0.5, size=10, color=GRAY)
        if i < 2:
            line = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(xs[i] + 1.7), Inches(3.35), Inches(1.25), Inches(0.22))
            line.fill.solid(); line.fill.fore_color.rgb = MID; line.line.fill.background()
    footer(s, 3)


def slide04(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    diagram = crop_asset(4, "diagram", (0.45, 0.10, 1.0, 0.86))
    add_picture(s, diagram, 6.1, 0.85, 6.75, 5.9)
    title_block(s, 4, "方法选择", "动态斜面法反推出动摩擦系数")
    add_round_rect(s, 0.8, 2.35, 3.7, 0.55, fill=RGBColor(230, 243, 250), line=MID)
    add_text(s, "ma = mg sinθ - μmg cosθ", 1.05, 2.52, 3.15, 0.2, size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_multiline(s, ["通过记录物块沿冰面下滑过程中的运动参数，", "结合斜面方向动力学方程，反推出动摩擦系数 μ。"], 0.85, 3.25, 4.5, 0.8, size=12)
    footer(s, 4)


def slide05(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    app = crop_asset(5, "apparatus", (0.28, 0.18, 0.78, 0.78))
    add_picture(s, app, 4.3, 1.35, 5.05, 4.55)
    title_block(s, 5, "装置设计", "斜面冰面与竖直冰面双结构验证")
    add_multiline(s, ["实验装置由斜面冰面结构和竖直冰面结构组成，", "用于动态滑动测量与补充验证，提高实验结果可靠性。"], 0.8, 2.15, 4.0, 0.65, size=11)
    items = ["斜面冰面结构", "竖直冰面结构", "低温环境模拟", "摩擦系统测量"]
    for i, item in enumerate(items):
        y = 3.2 + i * 0.7
        add_circle_icon(s, 0.82, y, str(i + 1))
        add_text(s, item, 1.35, y + 0.08, 2.2, 0.2, size=11, bold=True, color=TEAL)
    for i, item in enumerate(["高精度传感器", "温控系统", "数据采集", "模块化结构"]):
        y = 2.2 + i * 0.78
        add_circle_icon(s, 10.15, y, "▣")
        add_text(s, item, 10.7, y + 0.08, 1.6, 0.2, size=10, color=GRAY)
    footer(s, 5)


def slide06(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title_block(s, 6, "变量设计", "冰体维度与金属维度分变量控制")
    add_multiline(s, ["围绕温度、压强、材质、粗糙度、接触面积和冰的成分建立因素矩阵，通过单因素控制比较摩擦系数变化。"], 0.7, 1.85, 11.8, 0.4, size=10)
    cx, cy = 6.65, 4.0
    center = s.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(cx - 0.55), Inches(cy - 0.48), Inches(1.1), Inches(0.96))
    center.fill.solid(); center.fill.fore_color.rgb = NAVY; center.line.color.rgb = NAVY
    add_text(s, "μ", cx - 0.18, cy - 0.18, 0.36, 0.25, size=25, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    factors = [("温度", "−30℃~0℃", 6.55, 2.35), ("压强", "0.05C–0.5 MPa", 9.7, 3.1), ("粗糙度", "Ra 0.02–1.6 μm", 9.65, 4.85), ("成分", "纯水/海水/含盐冰", 7.65, 5.75), ("面积", "接触面积变化", 4.4, 5.75), ("材质", "铝/铜/钢/铁", 3.55, 4.0)]
    for name, val, x, y in factors:
        line = s.shapes.add_connector(1, Inches(cx), Inches(cy), Inches(x + 0.35), Inches(y + 0.28))
        line.line.color.rgb = MID; line.line.width = Pt(1)
        add_circle_icon(s, x, y, "◎")
        add_text(s, name, x + 0.55, y + 0.02, 1.1, 0.2, size=12, bold=True, color=NAVY)
        add_text(s, val, x + 0.55, y + 0.30, 1.8, 0.18, size=8, color=GRAY)
    footer(s, 6)


def slide07(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title_block(s, 7, "问题改进", "解决冰面平整、低温保持与传感稳定")
    add_multiline(s, ["针对冰面气泡和凹坑、实验温度波动、低温传感器稳定性等问题，采用去离子水、平整制冰、低温控制和外置传感器等改进方案。"], 0.75, 1.75, 11.5, 0.4, size=10)
    crops = [
        crop_asset(7, "ice", (0.06, 0.33, 0.33, 0.72)),
        crop_asset(7, "temp", (0.35, 0.33, 0.63, 0.72)),
        crop_asset(7, "sensor", (0.66, 0.33, 0.94, 0.72)),
    ]
    cards = [("冰面制备", "去离子水、均匀滴加\n保障冰面平整无气泡"), ("平整度与温控", "采用制冷和温控平台\n控制温度 ±0.5℃"), ("传感器可靠性", "外置传感器方案\n保证低温稳定采集")]
    for i, (t, b) in enumerate(cards):
        x = 0.75 + i * 4.08
        add_round_rect(s, x, 2.55, 3.45, 3.65, fill=WHITE, line=RGBColor(220, 232, 238))
        add_picture(s, crops[i], x + 0.18, 2.75, 3.08, 1.6)
        add_text(s, t, x + 0.25, 4.62, 2.4, 0.22, size=13, bold=True, color=NAVY)
        add_multiline(s, b.split("\n"), x + 0.25, 5.02, 2.8, 0.5, size=9)
    footer(s, 7)


def slide08(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title_block(s, 8, "数据处理", "由 x、t、θ 得到 a，再计算 μ")
    add_multiline(s, ["实验记录位移、时间、倾角等参数，计算加速度后代入动力学模型得到摩擦系数，并进行不确定度评估。"], 0.75, 1.75, 11.5, 0.35, size=10)
    steps = [("数据采集\nx,t,θ", "记录 x\n时间 t\n倾角 θ"), ("加速度计算\na", "数据滤波\n求取加速度"), ("摩擦系数计算\nμ", "代入模型\n计算 μ"), ("不确定度分析\nU(μ)", "误差传播\n结果评估")]
    for i, (t, b) in enumerate(steps):
        x = 0.9 + i * 3.0
        add_round_rect(s, x, 3.0, 2.15, 2.25, fill=WHITE, line=MID)
        add_text(s, str(i + 1), x + 0.18, 3.25, 0.3, 0.25, size=14, bold=True, color=RGBColor(126, 179, 213))
        add_text(s, t, x + 0.55, 3.22, 1.3, 0.55, size=12, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
        add_multiline(s, b.split("\n"), x + 0.52, 4.35, 1.2, 0.42, size=8, color=GRAY)
        if i < 3:
            arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.25), Inches(3.9), Inches(0.55), Inches(0.26))
            arr.fill.solid(); arr.fill.fore_color.rgb = RGBColor(113, 170, 207); arr.line.fill.background()
    footer(s, 8)


def slide09(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title_block(s, 9, "结果分析", "比较不同条件下摩擦系数变化趋势")
    chart_area = crop_asset(9, "charts", (0.06, 0.27, 0.96, 0.74))
    add_picture(s, chart_area, 0.72, 2.15, 11.85, 3.45)
    add_round_rect(s, 0.9, 5.95, 11.5, 0.7, fill=RGBColor(229, 242, 249), line=MID)
    add_text(s, "通过 μ-温度、μ-压力、μ-粗糙度等趋势图分析变量影响，不凭空给出精确数据，只展示对比分析框架。", 1.25, 6.18, 10.5, 0.2, size=11, color=NAVY)
    footer(s, 9)


def slide10(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    title_block(s, 10, "总结展望", "自制装置、多变量测量、工程应用")
    app = crop_asset(10, "apparatus", (0.52, 0.18, 0.96, 0.52))
    ship = crop_asset(10, "ship", (0.52, 0.55, 0.96, 0.90))
    add_picture(s, app, 6.9, 1.25, 5.4, 2.25)
    add_picture(s, ship, 6.9, 3.85, 5.4, 2.15)
    add_round_rect(s, 0.8, 2.05, 4.95, 1.35, fill=WHITE, line=MID)
    add_text(s, "研究总结", 1.15, 2.28, 1.4, 0.2, size=13, bold=True, color=NAVY)
    add_multiline(s, ["搭建了双结构实验测量装置", "完成多变量影响规律分析", "建立了数据处理与不确定度评估方法"], 1.15, 2.65, 3.7, 0.48, size=9, bullet=True)
    add_round_rect(s, 0.8, 3.8, 4.95, 1.35, fill=WHITE, line=MID)
    add_text(s, "应用前景", 1.15, 4.03, 1.4, 0.2, size=13, bold=True, color=NAVY)
    add_multiline(s, ["船体设计与减阻优化", "冰区结构与表面安全评估", "极地机器人与交通工具研发"], 1.15, 4.4, 3.7, 0.48, size=9, bullet=True)
    footer(s, 10)


def main():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    for fn in [slide01, slide02, slide03, slide04, slide05, slide06, slide07, slide08, slide09, slide10]:
        fn(prs)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
