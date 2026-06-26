from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "outputs" / "page01_02_hybrid_test_v1"
ASSET_DIR = OUT_DIR / "assets"
PPTX_OUT = OUT_DIR / "page01_02_hybrid_test_v1.pptx"

SLIDE_W = Inches(13.333333)
SLIDE_H = Inches(7.5)

TEAL = RGBColor(8, 75, 92)
TEAL_DARK = RGBColor(4, 58, 73)
TEAL_LIGHT = RGBColor(82, 151, 169)
ORANGE = RGBColor(239, 119, 35)
TEXT = RGBColor(24, 54, 64)
MUTED = RGBColor(82, 104, 112)
LINE = RGBColor(191, 215, 222)
PANEL = RGBColor(250, 253, 253)
GREEN = RGBColor(48, 166, 79)


def px(value: float):
    return Pt(value / 4)


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=20,
    bold=False,
    color=TEXT,
    align=PP_ALIGN.LEFT,
    font="Microsoft YaHei",
):
    box = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    box.fill.background()
    box.line.fill.background()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_line(slide, x1, y1, x2, y2, color=LINE, width=1.0):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2)
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def add_rect(
    slide,
    x,
    y,
    w,
    h,
    fill=PANEL,
    line=LINE,
    radius=False,
    transparency=0,
    line_width=0.75,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        px(x),
        px(y),
        px(w),
        px(h),
    )
    shape.fill.background()
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    return shape


def add_icon_circle(slide, cx, cy, label, size=52):
    circ = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, px(cx - size / 2), px(cy - size / 2), px(size), px(size)
    )
    circ.fill.background()
    circ.line.color.rgb = LINE
    circ.line.width = Pt(0.8)
    add_text(
        slide,
        label,
        cx - size / 2,
        cy - 11,
        size,
        24,
        size=15,
        bold=True,
        color=TEAL,
        align=PP_ALIGN.CENTER,
        font="Arial",
    )


def add_page_marker(slide, num, small=False):
    if small:
        add_text(slide, num, 30, 13, 58, 42, size=18, bold=True, color=RGBColor(255, 255, 255))
        add_line(slide, 100, 42, 315, 42, color=TEAL_LIGHT, width=1.2)
    else:
        add_text(slide, num, 60, 30, 148, 70, size=38, bold=True, color=RGBColor(255, 255, 255))


def add_background(slide, filename):
    slide.shapes.add_picture(str(ASSET_DIR / filename), 0, 0, width=SLIDE_W, height=SLIDE_H)


def build_page01(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, "page01_bg_notext_v2.png")
    add_page_marker(slide, "01")

    add_text(slide, "金属与冰", 118, 184, 710, 112, size=38, bold=True, color=TEAL_DARK)
    add_text(slide, "摩擦系数的测量", 118, 300, 860, 126, size=38, bold=True, color=TEAL_DARK)
    add_line(slide, 120, 437, 408, 437, color=ORANGE, width=2.0)
    add_text(slide, "物理实验创新竞赛答辩", 120, 464, 560, 42, size=17, color=TEXT)

    add_icon_circle(slide, 144, 603, "●", size=28)
    add_text(slide, "团队：", 202, 583, 90, 32, size=13, bold=True, color=TEAL_DARK)
    add_text(slide, "极冰组陈者", 288, 583, 230, 32, size=13, color=TEXT)

    add_text(slide, "▣", 131, 674, 28, 28, size=14, bold=True, color=TEAL)
    add_text(slide, "日期：", 202, 670, 90, 32, size=13, bold=True, color=TEAL_DARK)
    add_text(slide, "2024.05.20", 288, 670, 230, 32, size=13, color=TEXT)

    footer_items = ["物理实验", "摩擦学研究", "精密测量", "工程创新"]
    x = 520
    for idx, item in enumerate(footer_items):
        add_text(slide, item, x, 2006, 220, 28, size=10, color=RGBColor(255, 255, 255), align=PP_ALIGN.CENTER)
        if idx < len(footer_items) - 1:
            add_line(slide, x + 230, 1998, x + 230, 2038, color=RGBColor(126, 173, 184), width=0.7)
        x += 270


def build_page02(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, "page02_bg_notext.png")
    add_page_marker(slide, "02", small=True)
    add_text(slide, "选题背景", 372, 43, 390, 56, size=22, bold=True, color=TEAL_DARK)
    add_line(slide, 374, 118, 760, 118, color=TEAL_LIGHT, width=1.3)

    add_text(slide, "北极航运与冰阻力问题", 2580, 300, 720, 50, size=19, bold=True, color=TEAL_DARK)
    add_line(slide, 2582, 380, 3332, 380, color=LINE, width=0.9)
    add_text(slide, "航行阻力中，冰摩擦阻力占比", 2580, 465, 650, 34, size=13, color=MUTED)
    add_text(slide, "25%-55%", 2580, 535, 430, 64, size=30, bold=True, color=ORANGE)
    add_text(slide, "冰面相关事故中，摩擦失控因素占比", 2580, 690, 720, 34, size=13, color=MUTED)
    add_text(slide, "13%-15%", 2580, 760, 430, 64, size=30, bold=True, color=ORANGE)

    add_text(slide, "研究意义", 466, 1488, 250, 40, size=13, bold=True, color=TEAL_DARK)
    items = [
        "明确金属滑冰与航行过程中的关键摩擦机理。",
        "为极地材料选型、表面处理与结构设计提供依据。",
        "提升极地装备在低温与工程安全的可靠性。",
    ]
    y = 1572
    for item in items:
        add_icon_circle(slide, 358, y + 13, "◎", size=42)
        add_text(slide, item, 466, y, 2200, 36, size=12, color=TEXT)
        y += 78


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    build_page01(prs)
    build_page02(prs)
    prs.save(PPTX_OUT)
    print(PPTX_OUT)


if __name__ == "__main__":
    main()
