from pathlib import Path
from PIL import Image

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(r"F:\Users\user\Documents\测金属摩擦系数")
SRC = ROOT / "outputs" / "keynote_single_pages"
ASSETS = ROOT / "outputs" / "keynote_pptx_assets_v2"
OUT = ROOT / "outputs" / "metal_ice_keynote_editable_v2.pptx"
ASSETS.mkdir(exist_ok=True)

W, H = 13.333333, 7.5
NAVY = RGBColor(4, 55, 103)
BLUE = RGBColor(0, 95, 158)
TEAL = RGBColor(0, 101, 134)
ICE = RGBColor(235, 247, 252)
PAPER = RGBColor(247, 251, 254)
MID = RGBColor(190, 216, 230)
ORANGE = RGBColor(240, 113, 28)
GRAY = RGBColor(67, 88, 103)
WHITE = RGBColor(255, 255, 255)


def crop(slide_no, name, box):
    im = Image.open(SRC / f"slide_{slide_no:02d}.png").convert("RGB")
    w, h = im.size
    x1, y1, x2, y2 = box
    out = ASSETS / f"s{slide_no:02d}_{name}.png"
    im.crop((int(x1 * w), int(y1 * h), int(x2 * w), int(y2 * h))).save(out)
    return out


def add_text(slide, text, x, y, w, h, size, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def add_lines(slide, lines, x, y, w, h, size=14, color=GRAY, bullet=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("• " if bullet else "") + line
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.line_spacing = 1.05
    return box


def rect(slide, x, y, w, h, fill, line=None, rounded=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    return shp


def pic(slide, path, x, y, w, h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PAPER


def footer(slide, page):
    rect(slide, 0, 7.12, W, 0.38, NAVY)
    add_text(slide, "金属与冰摩擦系数测量", 0.58, 7.22, 3.2, 0.18, 9, color=WHITE)
    add_text(slide, f"{page:02d} / 10", 11.65, 7.20, 1.0, 0.2, 12, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)


def page_title(slide, page, title, subtitle, large=False):
    add_text(slide, f"{page:02d}", 0.58, 0.45, 1.25, 0.55, 42 if large else 34, bold=True, color=ORANGE if page == 9 else NAVY)
    rect(slide, 0.62, 1.18, 0.75, 0.055, NAVY)
    add_text(slide, title, 0.58, 1.52 if large else 1.35, 5.0, 1.25, 36 if large else 28, bold=True, color=NAVY)
    if subtitle:
        add_text(slide, subtitle, 0.62, 3.35 if large else 2.28, 5.1, 0.35, 16, bold=True, color=TEAL)


def chip(slide, x, y, text, w=2.3):
    rect(slide, x, y, w, 0.45, WHITE, MID, rounded=True)
    add_text(slide, text, x + 0.18, y + 0.12, w - 0.28, 0.16, 10, bold=True, color=TEAL)


def slide01(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    pic(s, crop(1, "hero", (0.43, 0.02, 1.0, 0.93)), 5.15, 0.0, 8.18, 7.12)
    page_title(s, 1, "金属与冰\n摩擦系数的测量", "可重复、可调控的测量系统", large=True)
    add_lines(s, ["围绕金属-冰接触面的滑动摩擦特性，搭建实验装置，", "分析不同条件下摩擦系数的变化规律。"], 0.62, 4.05, 4.9, 0.65, 15, RGBColor(23, 55, 88))
    rect(s, 0.62, 5.35, 3.6, 0.85, RGBColor(238, 247, 252), MID, rounded=True)
    add_text(s, "研究要点", 0.86, 5.52, 1.2, 0.18, 12, bold=True)
    add_lines(s, ["金属-冰接触面", "实验装置", "规律分析"], 0.9, 5.78, 2.1, 0.4, 9, RGBColor(24, 57, 86), bullet=True)
    footer(s, 1)


def slide02(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    pic(s, crop(2, "ship", (0.43, 0.00, 1.0, 0.82)), 5.3, 0.0, 8.03, 6.25)
    page_title(s, 2, "选题背景", "冰上摩擦影响极地航运阻力与能耗")
    add_lines(s, ["北极航道缩短航程，但海冰显著增加船冰阻力。", "测量金属与冰摩擦系数，可支撑极地装备减阻设计。"], 0.65, 2.9, 4.85, 0.7, 13)
    for idx, (num, label) in enumerate([("25%-55%", "航程缩短范围"), ("13%-15%", "燃油消耗增加")]):
        y = 4.08 + idx * 0.82
        add_text(s, num, 1.05, y, 2.2, 0.35, 24, bold=True, color=ORANGE)
        add_text(s, label, 1.08, y + 0.42, 2.4, 0.18, 10, color=GRAY)
    footer(s, 2)


def slide03(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    page_title(s, 3, "题目解读", "完成机制分析、装置搭建、结果评估")
    add_lines(s, ["研究任务包括分析摩擦机制、设计实验装置、获得测量结果并讨论不确定度。"], 0.68, 2.08, 11.5, 0.3, 12)
    items = [("机制分析", "滑动摩擦机制"), ("装置搭建", "实验装置设计"), ("结果评估", "数据与不确定度")]
    for i, (a, b) in enumerate(items):
        x = 1.45 + i * 3.95
        oval = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(3.05), Inches(1.45), Inches(1.45))
        oval.fill.solid(); oval.fill.fore_color.rgb = RGBColor(232, 246, 252); oval.line.color.rgb = BLUE
        add_text(s, str(i + 1), x + 0.55, 3.48, 0.3, 0.3, 20, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
        add_text(s, a, x - 0.25, 4.82, 2.0, 0.25, 15, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, b, x - 0.25, 5.25, 2.0, 0.2, 10, color=GRAY, align=PP_ALIGN.CENTER)
        if i < 2:
            arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 1.85), Inches(3.62), Inches(1.15), Inches(0.25))
            arr.fill.solid(); arr.fill.fore_color.rgb = MID; arr.line.fill.background()
    footer(s, 3)


def slide04(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    pic(s, crop(4, "diagram", (0.38, 0.08, 1.0, 0.86)), 5.45, 0.85, 7.35, 5.95)
    page_title(s, 4, "方法选择", "动态斜面法反推出动摩擦系数")
    rect(s, 0.72, 3.0, 4.1, 0.58, RGBColor(231, 244, 250), MID, rounded=True)
    add_text(s, "ma = mg sinθ - μmg cosθ", 0.95, 3.17, 3.6, 0.2, 16, bold=True, align=PP_ALIGN.CENTER)
    add_lines(s, ["记录物块沿冰面下滑过程中的运动参数，", "结合斜面方向动力学方程反推出 μ。"], 0.75, 3.95, 4.6, 0.65, 13)
    footer(s, 4)


def slide05(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    pic(s, crop(5, "apparatus", (0.25, 0.16, 0.82, 0.80)), 3.55, 1.25, 6.2, 4.95)
    page_title(s, 5, "装置设计", "斜面冰面与竖直冰面双结构验证")
    add_lines(s, ["双结构实验装置用于动态滑动测量与补充验证，", "提高实验结果可靠性。"], 0.65, 2.85, 3.65, 0.6, 12)
    for i, t in enumerate(["斜面冰面结构", "竖直冰面结构", "低温环境模拟", "摩擦系统测量"]):
        chip(s, 0.72, 4.0 + i * 0.52, t, 2.2)
    for i, t in enumerate(["高精度传感器", "温控系统", "数据采集", "模块化结构"]):
        chip(s, 10.2, 2.45 + i * 0.72, t, 2.0)
    footer(s, 5)


def slide06(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    page_title(s, 6, "变量设计", "冰体维度与金属维度分变量控制")
    add_lines(s, ["围绕温度、压强、材质、粗糙度、接触面积和冰的成分建立因素矩阵。"], 0.7, 2.1, 9.8, 0.3, 12)
    cx, cy = 6.7, 4.15
    center = s.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(cx - 0.58), Inches(cy - 0.5), Inches(1.16), Inches(1.0))
    center.fill.solid(); center.fill.fore_color.rgb = NAVY; center.line.color.rgb = NAVY
    add_text(s, "μ", cx - 0.18, cy - 0.17, 0.36, 0.25, 26, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    factors = [("温度", "−30℃~0℃", 6.45, 2.35), ("压强", "0.05–0.5 MPa", 9.5, 3.2), ("粗糙度", "Ra 0.02–1.6 μm", 9.35, 5.0), ("成分", "纯水/海水/含盐冰", 7.35, 5.85), ("面积", "接触面积变化", 4.25, 5.85), ("材质", "铝/铜/钢/铁", 3.55, 4.0)]
    for name, val, x, y in factors:
        conn = s.shapes.add_connector(1, Inches(cx), Inches(cy), Inches(x + 0.2), Inches(y + 0.2))
        conn.line.color.rgb = MID
        conn.line.width = Pt(1)
        add_text(s, name, x, y, 1.0, 0.22, 13, bold=True)
        add_text(s, val, x, y + 0.32, 1.65, 0.18, 8, color=GRAY)
    footer(s, 6)


def slide07(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    page_title(s, 7, "问题改进", "解决冰面平整、低温保持与传感稳定")
    add_lines(s, ["针对冰面气泡、温度波动、低温传感器稳定性等问题，采用制冰与采集方案优化。"], 0.68, 2.05, 10.5, 0.3, 12)
    imgs = [
        crop(7, "ice", (0.06, 0.33, 0.33, 0.72)),
        crop(7, "temp", (0.35, 0.33, 0.63, 0.72)),
        crop(7, "sensor", (0.66, 0.33, 0.94, 0.72)),
    ]
    cards = [("冰面制备", "去离子水、均匀滴加"), ("平整度与温控", "控制温度 ±0.5℃"), ("传感器可靠性", "外置传感器稳定采集")]
    for i, (title, text) in enumerate(cards):
        x = 0.72 + i * 4.12
        rect(s, x, 2.75, 3.55, 3.35, WHITE, RGBColor(220, 232, 238), rounded=True)
        pic(s, imgs[i], x + 0.16, 2.95, 3.22, 1.75)
        add_text(s, title, x + 0.25, 4.92, 2.3, 0.22, 14, bold=True)
        add_text(s, text, x + 0.25, 5.34, 2.7, 0.18, 10, color=GRAY)
    footer(s, 7)


def slide08(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    page_title(s, 8, "数据处理", "由 x、t、θ 得到 a，再计算 μ")
    add_lines(s, ["采集位移、时间、倾角等参数，代入动力学模型得到摩擦系数，并进行不确定度评估。"], 0.68, 2.05, 11.0, 0.3, 12)
    steps = [("数据采集", "x,t,θ"), ("加速度计算", "a"), ("摩擦系数计算", "μ"), ("不确定度分析", "U(μ)")]
    for i, (a, b) in enumerate(steps):
        x = 0.95 + i * 3.05
        rect(s, x, 3.25, 2.25, 1.75, WHITE, MID, rounded=True)
        add_text(s, str(i + 1), x + 0.18, 3.48, 0.3, 0.22, 12, color=RGBColor(130, 180, 213))
        add_text(s, a, x + 0.55, 3.52, 1.2, 0.2, 12, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
        add_text(s, b, x + 0.55, 4.12, 1.2, 0.2, 16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        if i < 3:
            arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.33), Inches(3.95), Inches(0.58), Inches(0.24))
            arr.fill.solid(); arr.fill.fore_color.rgb = RGBColor(113, 170, 207); arr.line.fill.background()
    footer(s, 8)


def slide09(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    # Recreate the source-page rhythm: big orange page number + wide chart region.
    add_text(s, "09", 0.82, 0.55, 1.3, 0.55, 42, bold=True, color=ORANGE)
    rect(s, 0.86, 1.35, 0.75, 0.055, NAVY)
    add_text(s, "结果分析", 2.1, 0.82, 3.2, 0.45, 30, bold=True)
    add_text(s, "比较不同条件下摩擦系数变化趋势", 2.12, 1.62, 5.0, 0.28, 16, bold=True)
    pic(s, crop(9, "charts", (0.06, 0.28, 0.96, 0.73)), 0.8, 2.22, 11.75, 3.35)
    rect(s, 0.78, 5.9, 11.7, 0.78, RGBColor(230, 243, 250), MID, rounded=True)
    add_text(s, "通过 μ-温度、μ-压力、μ-粗糙度等趋势图分析变量影响，只展示对比分析框架。", 1.25, 6.18, 10.4, 0.24, 14, bold=True)
    footer(s, 9)


def slide10(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    pic(s, crop(10, "apparatus", (0.50, 0.16, 0.97, 0.53)), 6.25, 1.1, 5.95, 2.55)
    pic(s, crop(10, "ship", (0.50, 0.54, 0.97, 0.91)), 6.25, 4.02, 5.95, 2.25)
    page_title(s, 10, "总结展望", "自制装置、多变量测量、工程应用")
    rect(s, 0.72, 2.85, 4.75, 1.15, WHITE, MID, rounded=True)
    add_text(s, "研究总结", 1.02, 3.08, 1.2, 0.2, 13, bold=True)
    add_lines(s, ["双结构实验测量装置", "多变量影响规律分析", "数据处理与不确定度评估"], 1.02, 3.42, 3.4, 0.44, 9, bullet=True)
    rect(s, 0.72, 4.35, 4.75, 1.15, WHITE, MID, rounded=True)
    add_text(s, "应用前景", 1.02, 4.58, 1.2, 0.2, 13, bold=True)
    add_lines(s, ["船体减阻优化", "冰区结构安全评估", "极地装备研发"], 1.02, 4.92, 3.3, 0.44, 9, bullet=True)
    footer(s, 10)


def main():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    for fn in [slide01, slide02, slide03, slide04, slide05, slide06, slide07, slide08, slide09, slide10]:
        fn(prs)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
