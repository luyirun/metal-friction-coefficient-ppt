from pathlib import Path
from PIL import Image

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(r"F:\Users\user\Documents\测金属摩擦系数")
SRC_DIR = ROOT / "outputs" / "single_pages"
ASSET_DIR = ROOT / "outputs" / "editable_ppt_assets" / "full_deck"
OUT = ROOT / "outputs" / "metal_ice_full_editable.pptx"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

FONT_NAME = "Microsoft YaHei"

W, H = 16, 9
TEAL = RGBColor(4, 74, 87)
DARK = RGBColor(3, 63, 75)
TEXT = RGBColor(32, 45, 50)
MUTED = RGBColor(93, 112, 119)
BLUE = RGBColor(188, 224, 235)
ORANGE = RGBColor(232, 126, 45)
LINE = RGBColor(205, 219, 224)
BG = RGBColor(247, 250, 250)


def crop(page, name, box):
    im = Image.open(SRC_DIR / f"page_{page:02d}_industrial_whitepaper.png").convert("RGBA")
    p = ASSET_DIR / f"p{page:02d}_{name}.png"
    im.crop(box).save(p)
    return p


def prep_assets():
    return {
        "p01_mountains": crop(1, "mountains", (0, 0, 2048, 500)),
        "p01_apparatus": crop(1, "apparatus", (1100, 55, 2048, 1070)),
        "p01_inset": crop(1, "inset", (710, 692, 1040, 1010)),
        "p02_ship": crop(2, "ship", (55, 180, 1135, 680)),
        "p02_ice": crop(2, "ice", (1190, 720, 1998, 1040)),
        "p05_left": crop(5, "inclined", (120, 210, 1088, 865)),
        "p05_right": crop(5, "vertical", (1180, 225, 2048, 890)),
        "p10_ship1": crop(10, "ship1", (910, 240, 1425, 450)),
        "p10_iceblock": crop(10, "iceblock", (910, 470, 1425, 680)),
        "p10_ship2": crop(10, "ship2", (910, 700, 1425, 910)),
    }


def add_text(slide, x, y, w, h, s, size, color=TEXT, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    r.font.name = FONT_NAME
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def add_rect(slide, x, y, w, h, fill, radius=False, line=None, trans=0):
    typ = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shp = slide.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.fill.transparency = trans
    shp.line.color.rgb = line if line else fill
    if line is None:
        shp.line.transparency = 100000
    return shp


def add_line(slide, x1, y1, x2, y2, color=LINE, width=1.3):
    l = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    l.line.color.rgb = color
    l.line.width = Pt(width)
    return l


def base(slide, num, title, subtitle=None):
    add_rect(slide, 0, 0, W, H, BG)
    tab = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, Inches(-0.06), Inches(0), Inches(1.22), Inches(0.62))
    tab.fill.solid(); tab.fill.fore_color.rgb = DARK; tab.line.color.rgb = DARK
    add_text(slide, 0.08, 0.09, 0.70, 0.28, f"{num:02d}", 23, RGBColor(255, 255, 255), True)
    add_text(slide, 1.32, 0.25, 7.2, 0.46, title, 27, TEAL, True)
    add_line(slide, 1.32, 0.85, 1.97, 0.85, TEAL, 3.0)
    add_line(slide, 2.02, 0.85, 3.13, 0.85, LINE, 1.6)
    if subtitle:
        add_text(slide, 1.32, 0.96, 6.4, 0.30, subtitle, 13, TEAL, True)
    marker = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, Inches(15.12), Inches(8.47), Inches(0.80), Inches(0.38))
    marker.fill.solid(); marker.fill.fore_color.rgb = DARK; marker.line.color.rgb = DARK
    add_text(slide, 15.36, 8.54, 0.32, 0.16, f"{num:02d}", 12, RGBColor(255, 255, 255), True, PP_ALIGN.CENTER)


def bullet(slide, x, y, text, size=12):
    add_rect(slide, x, y + 0.08, 0.07, 0.07, TEAL, radius=True)
    add_text(slide, x + 0.20, y, 3.3, 0.28, text, size, TEXT)


def card(slide, x, y, w, h, title, body=None):
    add_rect(slide, x, y, w, h, RGBColor(255, 255, 255), radius=True, line=RGBColor(225, 232, 234))
    add_text(slide, x + 0.22, y + 0.18, w - 0.44, 0.32, title, 15, TEAL, True)
    if body:
        add_text(slide, x + 0.22, y + 0.62, w - 0.44, h - 0.65, body, 11, MUTED)


def icon_circle(slide, cx, cy, label, sub=None):
    c = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(cx - 0.46), Inches(cy - 0.46), Inches(0.92), Inches(0.92))
    c.fill.solid(); c.fill.fore_color.rgb = RGBColor(219, 239, 245); c.line.color.rgb = RGBColor(184, 215, 224)
    add_text(slide, cx - 0.32, cy - 0.12, 0.64, 0.22, label, 12, TEAL, True, PP_ALIGN.CENTER)
    if sub:
        add_text(slide, cx - 0.55, cy + 0.58, 1.10, 0.50, sub, 10.5, TEXT, False, PP_ALIGN.CENTER)


def build_page01(slide, a):
    add_rect(slide, 0, 0, W, H, RGBColor(246, 250, 251))
    slide.shapes.add_picture(str(a["p01_mountains"]), Inches(0), Inches(0), width=Inches(16), height=Inches(3.95))
    wash = add_rect(slide, 0, 0, W, H, RGBColor(255, 255, 255)); wash.fill.transparency = 39000
    slide.shapes.add_picture(str(a["p01_apparatus"]), Inches(8.55), Inches(0.30), width=Inches(7.42), height=Inches(8.05))
    tab = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, Inches(-0.18), Inches(-0.03), Inches(2.62), Inches(1.32))
    tab.fill.solid(); tab.fill.fore_color.rgb = DARK; tab.line.color.rgb = DARK
    add_text(slide, 0.48, 0.28, 1.15, 0.62, "01", 48, RGBColor(255, 255, 255), True)
    add_text(slide, 1.00, 2.10, 5.80, 1.95, "金属与冰\n摩擦系数的测量", 40, TEAL, True)
    add_text(slide, 1.04, 4.46, 4.60, 0.40, "大学生物理实验创新赛答辩", 22, TEXT, True)
    add_line(slide, 1.04, 5.10, 4.95, 5.10, RGBColor(216,226,229), 2)
    add_line(slide, 1.04, 5.10, 1.60, 5.10, TEAL, 3)
    icon_circle(slide, 1.14, 6.15, "人"); add_text(slide, 1.72, 5.97, 3.25, 0.38, "参赛选手：请填写姓名", 18)
    icon_circle(slide, 1.14, 7.15, "校"); add_text(slide, 1.72, 6.97, 3.20, 0.38, "参赛单位：请填写学校", 18)
    slide.shapes.add_picture(str(a["p01_inset"]), Inches(5.78), Inches(6.02), width=Inches(2.40), height=Inches(2.40))
    banner = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, Inches(0.68), Inches(8.10), Inches(7.00), Inches(0.58))
    banner.fill.solid(); banner.fill.fore_color.rgb = DARK; banner.line.color.rgb = DARK
    add_text(slide, 1.08, 8.22, 1.45, 0.32, "实验测量", 19, RGBColor(255,255,255), True)
    add_text(slide, 3.05, 8.22, 1.45, 0.32, "机理分析", 19, RGBColor(255,255,255), True)
    add_text(slide, 5.02, 8.22, 1.45, 0.32, "创新设计", 19, RGBColor(255,255,255), True)
    marker = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, Inches(14.87), Inches(8.36), Inches(1.20), Inches(0.64))
    marker.fill.solid(); marker.fill.fore_color.rgb = DARK; marker.line.color.rgb = DARK
    add_text(slide, 15.20, 8.53, 0.55, 0.28, "01", 23, RGBColor(255,255,255), True, PP_ALIGN.CENTER)


def page02(slide, a):
    base(slide, 2, "选题背景")
    slide.shapes.add_picture(str(a["p02_ship"]), Inches(0.58), Inches(1.45), width=Inches(8.28), height=Inches(3.85))
    card(slide, 0.58, 5.72, 8.28, 2.34, "研究意义", "揭示船舶破冰与航行效率的关键因素之一\n为船舶材料选型、表面处理与结构设计提供依据\n推动极地装备与工程安全的进步")
    add_text(slide, 9.65, 1.55, 5.2, 0.5, "北极航运与冰阻力问题", 25, TEAL, True)
    add_text(slide, 9.65, 2.45, 4.4, 0.38, "航行阻力中，冰障碍阻力占比", 15, TEXT)
    add_text(slide, 9.65, 2.94, 3.6, 0.70, "25%-55%", 31, ORANGE, True)
    add_line(slide, 9.65, 3.92, 14.85, 3.92, RGBColor(205, 232, 238), 1.2)
    add_text(slide, 9.65, 4.40, 4.7, 0.38, "冰面相关事故中，摩擦失控问题占比", 15, TEXT)
    add_text(slide, 9.65, 4.88, 3.4, 0.70, "13%-15%", 31, ORANGE, True)
    slide.shapes.add_picture(str(a["p02_ice"]), Inches(9.65), Inches(6.05), width=Inches(5.50), height=Inches(1.95))


def page03(slide):
    base(slide, 3, "题目解读")
    cols = [("机理分析", "冰的黏附性变弱\n表面微观水膜\n水膜润滑作用\n温度与压力影响"),
            ("装置搭建", "斜面摩擦实验平台\n力学测量与采集\n低温环境控制\n稳定与校准"),
            ("结果与不确定度", "摩擦系数计算\n不确定度评估\n数据性分析\n结果讨论与验证")]
    for i, (t, b) in enumerate(cols):
        x = 1.15 + i * 4.55
        card(slide, x, 1.35, 4.15, 6.45, t, "")
        icon_circle(slide, x + 2.05, 2.55, ["冰", "器", "图"][i])
        y = 3.70
        for line in b.split("\n"):
            bullet(slide, x + 0.55, y, line, 13); y += 0.58
        if i < 2:
            add_line(slide, x + 4.35, 1.45, x + 4.35, 7.25, RGBColor(190, 207, 212), 1.1)


def page04(slide):
    base(slide, 4, "方法选择 —— 动态斜面法")
    add_text(slide, 1.25, 1.35, 2, 0.34, "力学模型", 16, TEAL, True)
    add_line(slide, 2.0, 6.35, 7.6, 3.88, TEAL, 2.6)
    add_rect(slide, 4.30, 4.52, 1.25, 0.62, RGBColor(210, 215, 216), line=RGBColor(160,165,166))
    add_line(slide, 4.92, 4.82, 4.92, 2.15, TEAL, 2.1); add_text(slide, 5.05, 2.02, 0.4, 0.3, "N", 14, TEAL, True)
    add_line(slide, 4.92, 4.82, 4.92, 6.85, TEAL, 2.1); add_text(slide, 5.02, 6.72, 0.6, 0.3, "mg", 12, TEXT)
    add_line(slide, 4.92, 4.82, 6.55, 4.10, ORANGE, 1.8); add_text(slide, 6.60, 4.00, 0.45, 0.3, "f", 13, ORANGE, True)
    add_text(slide, 1.35, 6.95, 5.1, 0.3, "斜面方向建立动力学方程，反推出动摩擦系数", 12, MUTED)
    add_text(slide, 9.05, 1.35, 2, 0.34, "理论公式", 16, TEAL, True)
    card(slide, 9.05, 1.95, 5.05, 1.0, "a = g(sinθ - μcosθ)")
    add_text(slide, 11.28, 3.08, 0.6, 0.3, "↓", 27, RGBColor(145,165,172), True, PP_ALIGN.CENTER)
    card(slide, 9.05, 3.55, 5.05, 1.20, "μ = (sinθ - a/g) / cosθ")
    add_text(slide, 9.10, 5.10, 4.4, 1.5, "a —— 沿斜面方向的加速度\ng —— 重力加速度\nθ —— 斜面倾角\nμ —— 动摩擦系数", 13, TEXT)


def page05(slide, a):
    base(slide, 5, "装置设计")
    add_text(slide, 1.15, 1.18, 3.0, 0.34, "斜面冰面模块", 16, TEAL, True)
    slide.shapes.add_picture(str(a["p05_left"]), Inches(0.95), Inches(1.52), width=Inches(7.55), height=Inches(5.10))
    add_text(slide, 8.80, 1.18, 4.0, 0.34, "垂直冰面模块（压力加载）", 16, TEAL, True)
    slide.shapes.add_picture(str(a["p05_right"]), Inches(8.65), Inches(1.52), width=Inches(6.25), height=Inches(5.10))
    labels = ["模块化设计", "快速更换", "低温环境模拟", "高精度测量"]
    for i, lab in enumerate(labels):
        icon_circle(slide, 2.0 + i * 3.5, 7.30, "◎", lab)


def page06(slide):
    base(slide, 6, "变量设计", "六大影响变量与水平设置")
    items = [
        ("温度", "-5 ℃\n-10 ℃\n-15 ℃"), ("压强", "0.05 MPa\n0.10 MPa\n0.20 MPa"),
        ("材质", "铝合金\n不锈钢\n钛合金"), ("粗糙度", "Ra 0.2 μm\nRa 1.0 μm\nRa 3.0 μm"),
        ("面积", "1 cm²\n4 cm²\n9 cm²"), ("成分", "淡水冰\n海冰\n人工掺盐冰"),
    ]
    for i, (t, b) in enumerate(items):
        x = 1.35 + i * 2.35
        icon_circle(slide, x, 3.0, ["❄","压","层","Ra","□","分"][i])
        add_text(slide, x - 0.65, 3.90, 1.3, 0.28, t, 16, TEAL, True, PP_ALIGN.CENTER)
        add_text(slide, x - 0.80, 4.48, 1.6, 1.2, b, 12, TEXT, False, PP_ALIGN.CENTER)


def page07(slide):
    base(slide, 7, "问题改进")
    rows = [
        ("冰面制备不平整", "表面粗糙和沉积过大\n影响测量重复性", "定制制冰机 + 控温体系\n实现光滑透明平整冰面"),
        ("低温环境稳定性差", "温度波动导致数据漂移\n影响实验结果", "双层保温 + PID温控\n温度波动 ≤ ±0.2 ℃"),
        ("传感器可靠性不足", "低温漂移与噪声干扰大\n影响信号准确性", "低温型传感器 + 屏蔽布线\n多点校准与滤波处理"),
    ]
    for i, (p, pb, s) in enumerate(rows):
        y = 1.42 + i * 2.10
        card(slide, 1.10, y, 5.45, 1.42, p, pb)
        add_text(slide, 6.78, y + 0.46, 0.45, 0.3, "→", 26, RGBColor(129,157,164), True)
        card(slide, 7.38, y, 6.45, 1.42, "改进方案", s)
        add_text(slide, 14.05, y + 0.44, 0.35, 0.3, "✓", 25, RGBColor(45,135,93), True)


def page08(slide):
    base(slide, 8, "数据处理")
    steps = [
        ("1\n数据采集", "力信号\n位移信号\n温度/压力\n采样频率 100 Hz"),
        ("2\n加速度计算", "位移二次微分\n滤波处理\n得到 a(t)"),
        ("3\n摩擦系数计算", "μ = (sinθ - a/g) / cosθ\n代入角度与加速度\n计算 μ"),
        ("4\n不确定度评估", "A类不确定度\nB类不确定度\n合成标准不确定度\n扩展不确定度 U"),
    ]
    for i, (t, b) in enumerate(steps):
        x = 0.95 + i * 3.78
        card(slide, x, 1.48, 3.00, 6.05, t, b)
        if i < 3:
            add_text(slide, x + 3.03, 4.15, 0.4, 0.3, "→", 22, RGBColor(129,157,164), True)
        if i == 0:
            add_line(slide, x+0.55, 3.55, x+1.85, 3.05, TEAL, 1.3); add_line(slide, x+1.85, 3.05, x+2.10, 3.25, TEAL, 1.3)
        if i == 1:
            add_line(slide, x+0.60, 3.75, x+2.10, 2.85, TEAL, 1.3)
        if i == 3:
            c = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ARC, Inches(x+0.92), Inches(2.75), Inches(1.0), Inches(0.8)); c.line.color.rgb = TEAL; c.line.width = Pt(2)


def page09(slide):
    base(slide, 9, "结果分析")
    add_text(slide, 2.95, 1.28, 4.3, 0.35, "摩擦系数对比趋势（示例）", 16, TEAL, True, PP_ALIGN.CENTER)
    # Chart axes and approximate data.
    x0, y0, cw, ch = 1.35, 1.95, 7.35, 4.95
    add_line(slide, x0, y0+ch, x0+cw, y0+ch, RGBColor(70,80,84), 1.2)
    add_line(slide, x0, y0, x0, y0+ch, RGBColor(70,80,84), 1.2)
    series = [
        (RGBColor(39, 132, 197), [0.34,0.30,0.26,0.20,0.15]),
        (ORANGE, [0.78,0.70,0.62,0.52,0.42]),
        (RGBColor(110, 164, 64), [0.56,0.50,0.43,0.34,0.25]),
    ]
    for col, vals in series:
        pts = []
        for j, v in enumerate(vals):
            pts.append((x0+0.6+j*(cw-1.2)/4, y0+ch - v*ch))
        for p1, p2 in zip(pts, pts[1:]):
            add_line(slide, p1[0], p1[1], p2[0], p2[1], col, 1.6)
        for px, py in pts:
            o = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(px-0.04), Inches(py-0.04), Inches(0.08), Inches(0.08))
            o.fill.solid(); o.fill.fore_color.rgb = col; o.line.color.rgb = col
    add_text(slide, 3.9, 7.12, 1.4, 0.25, "温度 / ℃", 12, TEXT, False, PP_ALIGN.CENTER)
    add_text(slide, 0.72, 3.90, 0.3, 1.2, "摩擦系数 μ", 12, TEXT)
    add_text(slide, 9.55, 1.70, 1.8, 0.34, "数据解读", 16, TEAL, True)
    for k, line in enumerate(["温度升高，μ整体下降", "粗糙度增大，μ上升", "压强增大，先升后降", "材质差异显著，钛合金综合表现最优"]):
        bullet(slide, 9.55, 2.35+k*0.62, line, 13)
    card(slide, 9.55, 5.55, 3.95, 1.25, "结果与文献趋势一致，验证方法的有效性")


def page10(slide, a):
    base(slide, 10, "总结展望")
    add_text(slide, 3.0, 1.08, 1.2, 0.28, "创新点", 15, TEAL, True)
    points = ["模块化多工况装置设计", "低温稳定控制与高精度测量", "多因素系统研究与量化分析", "不确定度评估与可靠性提升"]
    for i, p in enumerate(points):
        card(slide, 1.10, 1.45+i*1.32, 5.55, 0.92, p)
    add_text(slide, 9.8, 1.08, 2.3, 0.28, "工程应用前景", 15, TEAL, True)
    imgs = [a["p10_ship1"], a["p10_iceblock"], a["p10_ship2"]]
    labels = ["极地船舶设计优化\n降低航行阻力", "表面工程与减阻开发\n提升抗冰性能", "极地装备与结构安全\n提供理论依据"]
    for i in range(3):
        y = 1.45 + i*1.65
        slide.shapes.add_picture(str(imgs[i]), Inches(8.30), Inches(y), width=Inches(3.05), height=Inches(1.15))
        card(slide, 11.55, y, 3.45, 1.15, labels[i].split("\n")[0], labels[i].split("\n")[1])
    card(slide, 1.10, 7.25, 13.90, 0.82, "未来将拓展更多材料与工况，建立摩擦数据库，服务极地工程与安全航行。")


def build():
    a = prep_assets()
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    for i in range(10):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if i == 0: build_page01(slide, a)
        elif i == 1: page02(slide, a)
        elif i == 2: page03(slide)
        elif i == 3: page04(slide)
        elif i == 4: page05(slide, a)
        elif i == 5: page06(slide)
        elif i == 6: page07(slide)
        elif i == 7: page08(slide)
        elif i == 8: page09(slide)
        elif i == 9: page10(slide, a)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
