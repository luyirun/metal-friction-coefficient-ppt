from pathlib import Path
from math import cos, sin, pi

from PIL import Image, ImageDraw, ImageEnhance, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
PROJECT = ROOT / "projects" / "slide09_whitepaper_static_v5"
ASSET_DIR = PROJECT / "assets"
OUT_DIR = PROJECT / "exports"
VALIDATION_DIR = PROJECT / "validation"
OUT = OUT_DIR / "slide09_whitepaper_static_v5.pptx"

W, H = 16, 9
FONT = "Microsoft YaHei"

PAPER = RGBColor(247, 250, 250)
WHITE = RGBColor(255, 255, 255)
DEEP = RGBColor(0, 72, 84)
DEEP_2 = RGBColor(0, 54, 68)
ICE = RGBColor(225, 243, 247)
ICE_2 = RGBColor(205, 233, 240)
ICE_LINE = RGBColor(177, 215, 225)
BLUE = RGBColor(75, 172, 198)
DARK = RGBColor(28, 44, 49)
MUTED = RGBColor(88, 108, 114)
LINE = RGBColor(211, 225, 229)
SOFT = RGBColor(230, 239, 242)
PANEL = RGBColor(242, 247, 248)

FACTORS = [
    ("F01", "温度", "Temperature", "℃", "-15 / -10 / -5 / 0 ℃"),
    ("F02", "压强", "Pressure", "P", "0.05 / 0.10 / 0.20 MPa"),
    ("F03", "金属材质", "Metal Material", "▰", "Al / Steel / Ti"),
    ("F04", "表面粗糙度", "Surface Roughness", "≈", "Ra 0.2 / 1.0 / 3.0 μm"),
    ("F05", "接触面积", "Contact Area", "□", "1 / 4 / 9 cm²"),
    ("F06", "冰的成分", "Ice Composition", "⌬", "淡水冰 / 海冰 / 盐冰"),
]


def make_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    bg = Image.new("RGB", (2048, 1152), (247, 250, 250))
    d = ImageDraw.Draw(bg, "RGBA")
    for x in range(-220, 2260, 260):
        d.line([(x, 132), (x + 120, 1060)], fill=(140, 178, 188, 14), width=2)
    for y in range(180, 1040, 190):
        d.line([(100, y), (1980, y)], fill=(140, 178, 188, 11), width=2)
    for box, a in [((500, -330, 1560, 730), 14), ((630, -210, 1430, 590), 11)]:
        d.ellipse(box, outline=(75, 172, 198, a), width=4)
    bg = bg.filter(ImageFilter.GaussianBlur(0.45))
    texture = ASSET_DIR / "whitepaper_background_texture_v4.png"
    bg.save(texture, quality=96)

    ref = ROOT / "outputs" / "whitepaper_native_assets" / "p05_inclined_clean.png"
    ref_img = Image.open(ref).convert("RGB")
    crop = ref_img.crop((135, 45, 710, 575))
    scene = Image.new("RGB", (1500, 820), (244, 248, 250))
    d = ImageDraw.Draw(scene, "RGBA")
    for y in range(820):
        shade = 250 - int(y * 5 / 820)
        d.line([(0, y), (1500, y)], fill=(shade, shade + 1, 252, 255))
    scale = max(1180 / crop.width, 710 / crop.height)
    crop = crop.resize((int(crop.width * scale), int(crop.height * scale)), Image.Resampling.LANCZOS)
    crop = ImageEnhance.Color(crop).enhance(0.72)
    crop = ImageEnhance.Contrast(crop).enhance(1.05)
    crop = ImageEnhance.Sharpness(crop).enhance(1.08)
    scene.paste(crop, ((1500 - crop.width) // 2 + 40, (820 - crop.height) // 2 + 15))
    overlay = Image.new("RGBA", scene.size, (0, 0, 0, 0))
    od = ImageDraw.Draw(overlay, "RGBA")
    od.rectangle((0, 0, 1499, 819), outline=(177, 215, 225, 125), width=3)
    od.line((50, 58, 205, 58), fill=(0, 72, 84, 120), width=4)
    od.line((50, 58, 50, 158), fill=(0, 72, 84, 56), width=2)
    scene = Image.alpha_composite(scene.convert("RGBA"), overlay).convert("RGB")
    scene = scene.filter(ImageFilter.UnsharpMask(radius=1.0, percent=120, threshold=3))
    scene_path = ASSET_DIR / "metal_ice_contact_reference_clean_v5.png"
    scene.save(scene_path, quality=97)
    return texture, scene_path


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=False, transparency=0, name=None):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    if name:
        shp.name = name
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.fill.transparency = transparency
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    return shp


def line(slide, x1, y1, x2, y2, color=LINE, width=0.75, name=None, dash=False):
    shp = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    if name:
        shp.name = name
    shp.line.color.rgb = color
    shp.line.width = Pt(width)
    if dash:
        shp.line.dash_style = 4
    return shp


def text(slide, x, y, w, h, value, size, color=DARK, bold=False, align=PP_ALIGN.LEFT, name=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for idx, raw in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = raw
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def tab(slide, x, y, w, h, fill=DEEP):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = fill
    return shp


def ellipse(slide, x, y, w, h, outline=ICE_LINE, width=0.75, name=None):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shp.name = name
    shp.fill.background()
    shp.line.color.rgb = outline
    shp.line.width = Pt(width)
    return shp


def add_header(slide):
    tab(slide, -0.10, 0.08, 1.35, 0.66)
    text(slide, 0.18, 0.20, 0.58, 0.32, "09", 22, WHITE, True, name="slide09_title_index")
    text(slide, 1.42, 0.22, 3.1, 0.36, "结果分析", 23.5, DEEP, True, name="slide09_title_main")
    line(slide, 0.72, 0.86, 1.28, 0.86, DEEP, 3)
    line(slide, 1.40, 0.86, 2.62, 0.86, DEEP, 1.4)
    text(slide, 1.42, 1.00, 7.75, 0.20, "控制变量法下的摩擦系数变化规律 / Controlled Variable Analysis", 10.4, MUTED, name="slide09_title_sub")
    tab(slide, 15.12, 8.50, 0.80, 0.36)
    text(slide, 15.36, 8.58, 0.27, 0.12, "09", 10, WHITE, True, PP_ALIGN.CENTER)
    text(slide, 0.92, 8.48, 8.4, 0.16, "注：当前数据为版式示例，后续可替换为真实实验数据。", 7.5, MUTED, name="slide09_data_source_note")


def add_factor_node(slide, x, y, fid, cn, en, glyph, levels):
    rect(slide, x, y, 1.78, 0.66, WHITE, ICE_LINE, False, name=f"factor_node_{fid}")
    rect(slide, x, y, 0.08, 0.66, DEEP, None, False)
    rect(slide, x + 0.18, y + 0.15, 0.34, 0.34, ICE, ICE_LINE, False)
    text(slide, x + 0.25, y + 0.255, 0.20, 0.08, glyph, 7.1, DEEP, True, PP_ALIGN.CENTER)
    text(slide, x + 0.64, y + 0.11, 0.34, 0.08, fid, 6.5, BLUE, True)
    text(slide, x + 0.64, y + 0.27, 0.90, 0.12, cn, 8.4, DEEP, True)
    text(slide, x + 0.64, y + 0.45, 0.98, 0.08, en, 5.2, MUTED)
    text(slide, x + 0.16, y + 0.72, 1.36, 0.08, levels, 5.1, MUTED, False, PP_ALIGN.CENTER)


def add_model(slide):
    text(slide, 0.92, 1.42, 2.38, 0.22, "六因素控制变量模型", 13.2, DEEP, True)
    text(slide, 5.60, 1.44, 1.12, 0.14, "response: μ", 7.2, BLUE, True, PP_ALIGN.RIGHT)
    rect(slide, 0.88, 1.78, 6.45, 5.76, WHITE, LINE, False, name="slide09_factor_matrix")
    rect(slide, 1.04, 1.98, 6.13, 5.30, PANEL, SOFT, False, transparency=18000)
    rect(slide, 1.04, 1.98, 6.13, 0.10, DEEP, None, False)
    cx, cy = 4.08, 4.67
    ellipse(slide, cx - 2.08, cy - 1.52, 4.16, 3.04, ICE_LINE, 0.65)
    ellipse(slide, cx - 1.48, cy - 1.04, 2.96, 2.08, SOFT, 0.55)
    for deg in [0, 45, 90, 135, 180, 225, 270, 315]:
        a = pi * deg / 180
        line(slide, cx, cy, cx + 2.02 * cos(a), cy + 1.46 * sin(a), SOFT, 0.45)
    line(slide, cx - 2.25, cy, cx + 2.25, cy, ICE_LINE, 0.55)
    line(slide, cx, cy - 1.62, cx, cy + 1.62, ICE_LINE, 0.55)

    rect(slide, cx - 0.70, cy - 0.42, 1.40, 0.84, DEEP, None, False, name="model_response_core")
    rect(slide, cx - 0.54, cy - 0.29, 1.08, 0.58, ICE, ICE_LINE, False)
    text(slide, cx - 0.38, cy - 0.22, 0.76, 0.20, "μ", 22, DEEP, True, PP_ALIGN.CENTER, name="model_response_mu")
    text(slide, cx - 0.48, cy + 0.12, 0.96, 0.08, "response", 5.6, MUTED, False, PP_ALIGN.CENTER)

    positions = [
        (2.02, 2.66, FACTORS[0]), (4.38, 2.66, FACTORS[1]),
        (1.32, 4.34, FACTORS[5]), (5.04, 4.34, FACTORS[2]),
        (2.02, 5.98, FACTORS[4]), (4.38, 5.98, FACTORS[3]),
    ]
    anchors = [(2.92, 2.99), (5.28, 2.99), (2.22, 4.67), (5.94, 4.67), (2.92, 6.31), (5.28, 6.31)]
    for ax, ay in anchors:
        line(slide, cx, cy, ax, ay, ICE_LINE, 0.45)
    for x, y, factor in positions:
        add_factor_node(slide, x, y, *factor)
    text(slide, 1.22, 7.18, 5.70, 0.15, "模型图用于定位控制变量与响应量关系；交互阶段再在该结构上叠加分析态。", 7.7, MUTED)


def add_right(slide, scene_path):
    text(slide, 8.14, 1.42, 2.9, 0.22, "实验手段选择与模型构建", 13.2, DEEP, True)
    rect(slide, 7.86, 1.78, 7.06, 5.76, WHITE, LINE, False, name="slide09_experiment_model_overview")
    rect(slide, 8.02, 1.98, 6.74, 5.30, PANEL, SOFT, False, transparency=17000)
    rect(slide, 8.02, 1.98, 6.74, 0.10, DEEP, None, False)
    pic = slide.shapes.add_picture(str(scene_path), Inches(8.28), Inches(2.42), width=Inches(3.32), height=Inches(1.82))
    pic.name = "slide09_metal_ice_inset_image"
    text(slide, 8.56, 4.38, 2.76, 0.14, "金属-冰接触实验场景（局部视觉资产）", 7.2, MUTED, False, PP_ALIGN.CENTER)
    line(slide, 11.86, 2.46, 11.86, 4.52, ICE_LINE, 0.65)
    text(slide, 12.18, 2.48, 2.18, 1.02, "在同一测量框架下改变单一因素，保持其余条件稳定，提取 μ 的变化趋势并回到物理机理解释。", 9.7, DARK)
    rect(slide, 12.18, 3.86, 1.95, 0.44, WHITE, ICE_LINE, False)
    text(slide, 12.38, 4.00, 1.56, 0.10, "controlled-variable protocol", 5.9, BLUE, True, PP_ALIGN.CENTER)

    y = 5.46
    line(slide, 8.42, y + 0.28, 14.02, y + 0.28, ICE_LINE, 0.75)
    flow = ["控制变量法", "实验水平设计", "μ 数据采集", "趋势建模", "结果分析"]
    x = 8.22
    widths = [0.98, 1.08, 1.04, 0.98, 0.98]
    for i, (label, w) in enumerate(zip(flow, widths)):
        fill = DEEP if i == 4 else WHITE
        fg = WHITE if i == 4 else DEEP
        rect(slide, x, y, w, 0.58, fill, ICE_LINE, False, name=f"overview_step_{i+1}")
        text(slide, x + 0.10, y + 0.10, 0.26, 0.08, f"{i+1:02d}", 5.5, ICE if i == 4 else BLUE, True)
        text(slide, x + 0.12, y + 0.34, w - 0.24, 0.10, label, 6.8, fg, True, PP_ALIGN.CENTER)
        if i < 4:
            text(slide, x + w + 0.08, y + 0.22, 0.14, 0.10, "→", 7.2, MUTED, True, PP_ALIGN.CENTER)
        x += w + 0.28
    rect(slide, 8.26, 6.58, 5.98, 0.42, WHITE, LINE, False)
    text(slide, 8.72, 6.72, 5.08, 0.12, "静态阶段只呈现分析入口：模型总览、实验方法与建模流程。", 8.0, MUTED)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    VALIDATION_DIR.mkdir(parents=True, exist_ok=True)
    texture, scene = make_assets()
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(texture), Inches(0), Inches(0), width=Inches(W), height=Inches(H)).name = "slide09_background_texture"
    add_header(slide)
    add_model(slide)
    add_right(slide, scene)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
