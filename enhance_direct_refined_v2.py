from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
SRC = ROOT / "outputs" / "metal_ice_direct_editable_refined.pptx"
OUT = ROOT / "outputs" / "metal_ice_whitepaper_direct_refined_v2.pptx"
FONT = "Microsoft YaHei"

TEAL = RGBColor(0, 72, 84)
DARK = RGBColor(28, 44, 49)
MUTED = RGBColor(78, 98, 104)
LINE = RGBColor(211, 225, 229)
ICE = RGBColor(226, 243, 247)
ORANGE = RGBColor(226, 126, 39)
GREEN = RGBColor(61, 137, 78)
WHITE = RGBColor(255, 255, 255)


def text(slide, x, y, w, h, value, size, color=DARK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, raw in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = raw
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=True, transparency=0):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.fill.transparency = transparency
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    return shp


def line(slide, x1, y1, x2, y2, color=LINE, width=1):
    shp = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shp.line.color.rgb = color
    shp.line.width = Pt(width)
    return shp


def bullet(slide, x, y, value, size=10.5):
    rect(slide, x, y + 0.10, 0.055, 0.055, TEAL, radius=True)
    text(slide, x + 0.16, y, 3.4, 0.20, value, size, DARK)


def add_polyline(slide, pts, color=TEAL, width=1.4):
    for a, b in zip(pts, pts[1:]):
        line(slide, a[0], a[1], b[0], b[1], color, width)


def enhance_slide4(slide):
    # Restore the visual richness of the original force model without replacing editable text.
    text(slide, 1.35, 1.82, 1.35, 0.20, "受力分解", 12, TEAL, True)
    text(slide, 2.10, 5.82, 0.28, 0.16, "θ", 11, DARK, True)
    text(slide, 5.76, 4.36, 0.34, 0.16, "Σf", 11, DARK, True)
    text(slide, 3.12, 4.18, 0.30, 0.16, "R", 11, DARK, True)
    add_polyline(slide, [(2.10, 6.38), (2.50, 5.80), (2.90, 6.38)], MUTED, 0.8)
    line(slide, 4.55, 4.92, 5.85, 4.28, ORANGE, 1.4)
    line(slide, 4.55, 4.92, 3.56, 5.42, ORANGE, 1.4)
    rect(slide, 9.00, 6.30, 4.45, 0.62, RGBColor(244, 248, 248), LINE, True)
    text(slide, 9.24, 6.48, 3.90, 0.18, "通过位移-时间曲线求 a，再代入公式得到 μ", 10.5, TEAL, True, PP_ALIGN.CENTER)


def enhance_slide6(slide):
    # Add the lower factor-matrix strip that the original draft had, keeping it editable.
    labels = ["低温区间", "加载范围", "材料组合", "表面状态", "接触尺度", "冰体来源"]
    values = ["-5~-15 ℃", "0.05~0.20 MPa", "铝/钢/钛", "Ra 梯度", "1~9 cm²", "淡水/海冰"]
    for i, (label, value) in enumerate(zip(labels, values)):
        x = 0.96 + i * 2.38
        rect(slide, x, 6.58, 1.70, 0.72, RGBColor(248, 251, 251), LINE, True)
        text(slide, x + 0.18, 6.76, 1.30, 0.16, label, 9.8, TEAL, True, PP_ALIGN.CENTER)
        text(slide, x + 0.18, 7.04, 1.30, 0.15, value, 8.7, MUTED, False, PP_ALIGN.CENTER)


def enhance_slide8(slide):
    # Fill the process cards with visible, domain-specific editable micro-diagrams.
    # Card 1: signal trace
    add_polyline(slide, [(1.38, 3.70), (1.68, 3.05), (1.96, 4.12), (2.24, 3.32), (2.52, 3.62), (2.80, 3.10)], TEAL, 1.2)
    text(slide, 1.25, 4.32, 1.90, 0.18, "x / t / θ / F", 11, TEAL, True, PP_ALIGN.CENTER)

    # Card 2: fitted acceleration curve
    line(slide, 5.05, 4.12, 6.72, 4.12, LINE, 0.9)
    line(slide, 5.05, 4.12, 5.05, 3.00, LINE, 0.9)
    add_polyline(slide, [(5.10, 3.92), (5.45, 3.72), (5.80, 3.88), (6.15, 3.35), (6.55, 3.08)], TEAL, 1.2)
    text(slide, 5.08, 4.36, 1.62, 0.18, "二次微分 / 滤波", 10.5, TEAL, True, PP_ALIGN.CENTER)

    # Card 3: formula emphasis
    rect(slide, 8.72, 3.28, 1.95, 0.74, RGBColor(246, 249, 249), LINE, True)
    text(slide, 8.88, 3.52, 1.65, 0.18, "μ = (sinθ - a/g) / cosθ", 11.2, DARK, False, PP_ALIGN.CENTER)
    text(slide, 8.86, 4.40, 1.72, 0.18, "代入 θ 与 a", 10.5, TEAL, True, PP_ALIGN.CENTER)

    # Card 4: uncertainty bell curve
    add_polyline(slide, [(12.35, 4.10), (12.65, 3.70), (12.95, 3.20), (13.25, 3.70), (13.55, 4.10)], TEAL, 1.2)
    line(slide, 12.20, 4.10, 13.72, 4.10, LINE, 0.9)
    line(slide, 12.95, 3.18, 12.95, 4.18, LINE, 0.8)
    text(slide, 12.30, 4.38, 1.45, 0.18, "U(μ)", 11, TEAL, True, PP_ALIGN.CENTER)


def enhance_slide9(slide):
    # Restore chart legend density and result interpretation rhythm.
    legends = [(7.70, 2.10, RGBColor(41, 137, 190), "铝合金"), (7.70, 2.46, ORANGE, "不锈钢"), (7.70, 2.82, RGBColor(120, 168, 42), "钛合金")]
    for x, y, col, label in legends:
        line(slide, x, y + 0.08, x + 0.32, y + 0.08, col, 1.5)
        text(slide, x + 0.42, y, 0.72, 0.15, label, 9.2, DARK)
    rect(slide, 9.10, 6.28, 3.90, 0.55, RGBColor(248, 251, 251), LINE, True)
    text(slide, 9.36, 6.45, 3.28, 0.16, "趋势一致性用于验证实验方案可行性", 9.8, TEAL, True, PP_ALIGN.CENTER)


def enhance_slide10(slide):
    # Add closing emphasis line to make the ending slide less sparse.
    rect(slide, 1.10, 6.82, 13.70, 0.38, RGBColor(248, 251, 251), LINE, True)
    text(slide, 1.55, 6.94, 12.7, 0.14, "面向极地装备摩擦数据库建设：从实验装置、变量控制到工程应用形成闭环。", 9.5, TEAL, True, PP_ALIGN.CENTER)


def main():
    prs = Presentation(SRC)
    enhance_slide4(prs.slides[3])
    enhance_slide6(prs.slides[5])
    enhance_slide8(prs.slides[7])
    enhance_slide9(prs.slides[8])
    enhance_slide10(prs.slides[9])
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
