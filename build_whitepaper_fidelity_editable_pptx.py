from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "outputs"
SRC_DIR = OUT_DIR / "single_pages"
OUT = OUT_DIR / "metal_ice_whitepaper_fidelity_editable.pptx"

W, H = 16, 9
FONT = "Microsoft YaHei"

PAPER = RGBColor(248, 250, 250)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(0, 72, 84)
TEAL_2 = RGBColor(0, 96, 112)
DARK = RGBColor(25, 43, 48)
MUTED = RGBColor(80, 100, 106)
LINE = RGBColor(218, 229, 232)
ICE = RGBColor(236, 247, 250)
ORANGE = RGBColor(226, 126, 39)
GREEN = RGBColor(58, 138, 78)


def rgb(hex_value):
    hex_value = hex_value.lstrip("#")
    return RGBColor(int(hex_value[0:2], 16), int(hex_value[2:4], 16), int(hex_value[4:6], 16))


def add_bg(slide, page):
    path = SRC_DIR / f"page_{page:02d}_industrial_whitepaper.png"
    slide.shapes.add_picture(str(path), Inches(0), Inches(0), width=Inches(W), height=Inches(H))


def rect(slide, x, y, w, h, fill=WHITE, line=None, transparency=0, radius=False):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.fill.transparency = transparency
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    return shp


def line(slide, x1, y1, x2, y2, color=TEAL, width=2):
    shp = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shp.line.color.rgb = color
    shp.line.width = Pt(width)
    return shp


def txt(slide, x, y, w, h, value, size, color=DARK, bold=False, align=PP_ALIGN.LEFT, name=None, spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for idx, raw_line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = raw_line
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def bullet_lines(slide, x, y, lines, size=14, gap=0.42, color=DARK, dot=TEAL):
    for i, value in enumerate(lines):
        yy = y + i * gap
        rect(slide, x, yy + 0.11, 0.055, 0.055, dot, radius=True)
        txt(slide, x + 0.18, yy, 4.2, 0.24, value, size, color)


def page_tag(slide, page):
    rect(slide, 0.0, 0.0, 1.52, 0.72, TEAL)
    txt(slide, 0.23, 0.14, 0.72, 0.34, f"{page:02d}", 25, WHITE, True)
    rect(slide, 15.08, 8.52, 0.82, 0.34, TEAL)
    txt(slide, 15.33, 8.59, 0.28, 0.12, f"{page:02d}", 10.5, WHITE, True, PP_ALIGN.CENTER)


def header(slide, page, title, subtitle=None):
    rect(slide, 0, 0, W, 0.98, PAPER)
    page_tag(slide, page)
    txt(slide, 1.42, 0.14, 7.0, 0.34, title, 24, TEAL, True)
    line(slide, 0.72, 0.82, 1.20, 0.82, TEAL, 3)
    line(slide, 1.32, 0.82, 2.55, 0.82, TEAL_2, 1.4)
    if subtitle:
        rect(slide, 1.42, 0.62, 8.0, 0.34, WHITE)
        txt(slide, 1.42, 0.65, 8.0, 0.22, subtitle, 12.8, MUTED, False)


def cover(slide, boxes):
    for item in boxes:
        if len(item) == 4:
            rect(slide, *item, WHITE)
        else:
            x, y, w, h, color = item
            rect(slide, x, y, w, h, color)


def slide01(slide):
    add_bg(slide, 1)
    cover(slide, [
        (0.0, 0.0, 2.0, 0.78, PAPER),
        (0.45, 1.02, 7.20, 2.80, rgb("#f7fbfb")),
        (0.62, 3.28, 4.20, 0.62, rgb("#f7fbfb")),
        (0.72, 4.05, 2.60, 1.68, rgb("#f7fbfb")),
        (0.80, 7.94, 3.72, 0.28, TEAL),
        (15.08, 8.52, 0.82, 0.34, TEAL),
    ])
    page_tag(slide, 1)
    txt(slide, 0.78, 1.38, 5.30, 1.58, "金属与冰\n摩擦系数的测量", 35, TEAL, True)
    txt(slide, 0.85, 3.34, 3.2, 0.32, "物理实验创新竞赛答辩", 17, DARK, True)
    line(slide, 0.85, 3.78, 1.38, 3.78, TEAL, 2.3)
    line(slide, 1.48, 3.78, 2.65, 3.78, LINE, 1.3)
    txt(slide, 1.20, 4.36, 2.0, 0.24, "团队：祝冰间隙者", 12.5, DARK)
    txt(slide, 1.20, 5.23, 2.0, 0.24, "日期：2024.05.20", 12.5, DARK)
    txt(slide, 1.08, 8.00, 3.05, 0.18, "精准测量  ·  理解机理  ·  服务工程", 10.5, WHITE, True)
    txt(slide, 15.33, 8.59, 0.28, 0.12, "01", 10.5, WHITE, True, PP_ALIGN.CENTER)


def slide02(slide):
    add_bg(slide, 2)
    cover(slide, [
        (0.0, 0.0, 4.85, 0.92, PAPER),
        (8.80, 1.28, 5.1, 2.95, WHITE),
        (1.05, 5.28, 5.70, 1.72, WHITE),
        (15.08, 8.52, 0.82, 0.34, TEAL),
    ])
    header(slide, 2, "选题背景")
    txt(slide, 9.04, 1.44, 4.5, 0.28, "北极航运与冰阻力问题", 17, TEAL, True)
    txt(slide, 9.05, 2.02, 4.0, 0.20, "航行阻力中，冰障碍阻力占比", 12.2, MUTED)
    txt(slide, 9.05, 2.34, 2.2, 0.34, "25%-55%", 21, ORANGE, True)
    txt(slide, 9.05, 3.42, 4.3, 0.20, "冰面相关事故中，摩擦失控问题占比", 12.2, MUTED)
    txt(slide, 9.05, 3.74, 2.2, 0.34, "13%-15%", 21, ORANGE, True)
    rect(slide, 1.12, 5.38, 5.55, 1.48, WHITE, LINE, radius=True)
    txt(slide, 1.38, 5.56, 1.4, 0.22, "研究意义", 12.8, TEAL, True)
    bullet_lines(slide, 1.22, 5.95, [
        "揭示船舶破冰与航行效率的关键因素之一",
        "为船舶材料选型、表面处理与结构设计提供依据",
        "推动极地装备与工程安全的进步",
    ], 10.8, 0.34)


def slide03(slide):
    add_bg(slide, 3)
    cover(slide, [
        (0.0, 0.0, 4.50, 0.92, PAPER),
        (1.08, 1.12, 13.8, 5.95, WHITE),
        (15.08, 8.52, 0.82, 0.34, TEAL),
    ])
    header(slide, 3, "题目解读")
    cols = [
        ("机理分析", "冰的黏附性变弱\n表面微观水膜\n水膜润滑作用\n温度与压力影响"),
        ("装置搭建", "斜面摩擦实验平台\n力学测量与采集\n低温环境控制\n稳定与校准"),
        ("结果与不确定度", "摩擦系数计算\n不确定度评估\n数据性分析\n结果讨论与验证"),
    ]
    for i, (title, body) in enumerate(cols):
        x = 1.35 + i * 4.55
        rect(slide, x, 1.22, 4.05, 5.60, WHITE, LINE)
        txt(slide, x + 0.45, 1.56, 3.15, 0.28, title, 16.5, TEAL, True, PP_ALIGN.CENTER)
        txt(slide, x + 1.52, 2.60, 1.0, 0.62, ["□", "⚒", "↗"][i], 33, TEAL_2, False, PP_ALIGN.CENTER)
        bullet_lines(slide, x + 0.62, 4.20, body.split("\n"), 12.5, 0.46)


def slide04(slide):
    add_bg(slide, 4)
    cover(slide, [
        (0.0, 0.0, 6.8, 0.95, PAPER),
        (1.08, 1.18, 1.7, 0.28, WHITE),
        (8.42, 1.04, 5.35, 5.85, WHITE),
        (15.08, 8.52, 0.82, 0.34, TEAL),
    ])
    header(slide, 4, "方法选择 —— 动态斜面法")
    txt(slide, 1.18, 1.18, 1.5, 0.22, "力学模型", 13.8, TEAL, True)
    txt(slide, 8.86, 1.18, 1.7, 0.22, "理论公式", 13.8, TEAL, True)
    rect(slide, 8.82, 1.54, 4.45, 0.72, rgb("#f2f4f4"), line=None, radius=True)
    txt(slide, 9.08, 1.76, 3.9, 0.18, "a = g(sinθ - μcosθ)", 17, DARK, False, PP_ALIGN.CENTER)
    txt(slide, 10.95, 2.55, 0.32, 0.22, "↓", 18, TEAL_2, True, PP_ALIGN.CENTER)
    rect(slide, 8.82, 2.92, 4.45, 0.88, rgb("#f2f4f4"), line=None, radius=True)
    txt(slide, 9.02, 3.18, 4.05, 0.22, "μ = (sinθ - a/g) / cosθ", 16.5, DARK, False, PP_ALIGN.CENTER)
    txt(slide, 9.18, 4.18, 3.4, 1.20, "a —— 沿斜面方向的加速度\ng —— 重力加速度\nθ —— 斜面倾角\nμ —— 动摩擦系数", 12.4, DARK)


def slide05(slide):
    add_bg(slide, 5)
    cover(slide, [
        (0.0, 0.0, 4.25, 0.92, PAPER),
        (1.05, 0.98, 3.2, 0.35, WHITE),
        (8.45, 0.98, 4.7, 0.35, WHITE),
        (1.02, 1.70, 2.60, 3.05, WHITE),
        (10.48, 1.70, 2.95, 3.20, WHITE),
        (0.95, 7.42, 14.2, 0.65, WHITE),
        (15.08, 8.52, 0.82, 0.34, TEAL),
    ])
    header(slide, 5, "装置设计")
    txt(slide, 1.16, 1.04, 2.6, 0.24, "斜面冰面模块", 13.8, TEAL, True)
    txt(slide, 8.62, 1.04, 3.6, 0.24, "垂直冰面模块（压力加载）", 13.8, TEAL, True)
    for x, y, label in [
        (1.20, 1.72, "导轨与滑块"), (1.15, 2.23, "力传感器"), (1.05, 2.72, "位移传感器"),
        (1.10, 3.20, "冰面板"), (1.12, 3.70, "角度调节机构"), (1.05, 4.20, "支撑框架"),
        (10.75, 1.75, "压力加载装置"), (10.72, 2.50, "力传感器"), (10.72, 3.17, "冰面板"),
        (10.75, 3.85, "温控仓体"), (10.75, 4.55, "底盘"),
    ]:
        txt(slide, x, y, 1.25, 0.18, label, 9.6, MUTED)
    labels = ["模块化设计", "快速更换", "低温环境模拟", "高精度测量"]
    for i, label in enumerate(labels):
        txt(slide, 1.48 + i * 3.28, 7.68, 1.55, 0.18, label, 11.2, DARK, True, PP_ALIGN.CENTER)


def slide06(slide):
    add_bg(slide, 6)
    cover(slide, [
        (0.0, 0.0, 4.20, 0.92, PAPER),
        (0.70, 0.90, 7.5, 0.35, WHITE),
        (1.02, 4.00, 14.0, 2.90, WHITE),
        (15.08, 8.52, 0.82, 0.34, TEAL),
    ])
    header(slide, 6, "变量设计", "六大影响变量与水平设置")
    items = [
        ("温度", "-5 ℃\n-10 ℃\n-15 ℃"),
        ("压强", "0.05 MPa\n0.10 MPa\n0.20 MPa"),
        ("材质", "铝合金\n不锈钢\n钛合金"),
        ("粗糙度", "Ra 0.2 μm\nRa 1.0 μm\nRa 3.0 μm"),
        ("面积", "1 cm²\n4 cm²\n9 cm²"),
        ("成分", "淡水冰\n海冰\n人工掺盐冰"),
    ]
    for i, (title, body) in enumerate(items):
        x = 1.30 + i * 2.35
        txt(slide, x - 0.40, 4.05, 0.8, 0.20, title, 13.5, TEAL, True, PP_ALIGN.CENTER)
        txt(slide, x - 0.55, 4.72, 1.10, 0.92, body, 11.0, DARK, False, PP_ALIGN.CENTER)


def slide07(slide):
    add_bg(slide, 7)
    cover(slide, [
        (0.0, 0.0, 4.20, 0.92, PAPER),
        (1.42, 1.35, 4.10, 5.35, rgb("#f8faf9")),
        (7.12, 1.35, 4.45, 5.35, rgb("#f8faf9")),
        (15.08, 8.52, 0.82, 0.34, TEAL),
    ])
    header(slide, 7, "问题改进")
    rows = [
        ("冰面制备不平整", "表面粗糙和气泡大\n影响测量重复性", "改进方案", "定制制冰机 + 控温体系\n实现光滑透明平整冰面"),
        ("低温环境稳定性差", "温度波动导致数据漂移\n影响实验结果", "改进方案", "双层保温 + PID温控\n温度波动 ≤ ±0.2 ℃"),
        ("传感器可靠性不足", "低温漂移与噪声干扰大\n影响信号准确性", "改进方案", "低温型传感器 + 屏蔽布线\n多点校准与滤波处理"),
    ]
    for i, (p, desc, sol_title, sol) in enumerate(rows):
        y = 1.55 + i * 1.77
        rect(slide, 0.98, y - 0.20, 4.75, 1.20, rgb("#f7faf9"), LINE, radius=True)
        rect(slide, 6.86, y - 0.20, 4.90, 1.20, rgb("#f7faf9"), LINE, radius=True)
        txt(slide, 1.62, y, 2.55, 0.22, p, 13.0, DARK, True)
        txt(slide, 1.62, y + 0.34, 3.0, 0.48, desc, 10.5, MUTED)
        txt(slide, 7.42, y, 1.25, 0.22, sol_title, 13.0, GREEN, True)
        txt(slide, 7.42, y + 0.34, 3.15, 0.48, sol, 10.5, DARK)


def slide08(slide):
    add_bg(slide, 8)
    cover(slide, [
        (0.0, 0.0, 4.20, 0.92, PAPER),
        (1.10, 1.32, 14.10, 6.30, WHITE),
        (15.08, 8.52, 0.82, 0.34, TEAL),
    ])
    header(slide, 8, "数据处理")
    steps = [
        ("1", "数据采集", "力信号\n位移信号\n温度/压力\n采样频率 100 Hz"),
        ("2", "加速度计算", "位移二次微分\n滤波处理\n得到 a(t)"),
        ("3", "摩擦系数计算", "μ = (sinθ - a/g) / cosθ\n代入角度与加速度\n计算 μ"),
        ("4", "不确定度评估", "A类不确定度\nB类不确定度\n合成标准不确定度\n扩展不确定度 U"),
    ]
    for i, (num, title, body) in enumerate(steps):
        x = 1.32 + i * 3.55
        rect(slide, x, 1.58, 2.55, 5.55, WHITE, LINE, radius=True)
        oval = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 1.08), Inches(1.70), Inches(0.38), Inches(0.38))
        oval.fill.solid()
        oval.fill.fore_color.rgb = TEAL
        oval.line.fill.background()
        txt(slide, x + 1.15, 1.78, 0.24, 0.16, num, 9.5, WHITE, True, PP_ALIGN.CENTER)
        txt(slide, x + 0.35, 2.24, 1.85, 0.22, title, 12.5, TEAL, True, PP_ALIGN.CENTER)
        if i == 2:
            txt(slide, x + 0.25, 4.08, 2.05, 0.52, "μ = (sinθ - a/g)\n      / cosθ", 15.0, DARK, False, PP_ALIGN.CENTER)
            txt(slide, x + 0.40, 5.58, 1.75, 0.48, "代入角度与加速度\n计算 μ", 10.5, DARK)
        else:
            bullet_lines(slide, x + 0.38, 5.18, body.split("\n"), 9.6, 0.28)


def slide09(slide):
    add_bg(slide, 9)
    cover(slide, [
        (0.0, 0.0, 4.20, 0.92, PAPER),
        (2.15, 1.05, 4.3, 0.28, WHITE),
        (8.60, 1.28, 5.20, 3.30, WHITE),
        (8.86, 5.30, 3.95, 0.95, WHITE),
        (15.08, 8.52, 0.82, 0.34, TEAL),
    ])
    header(slide, 9, "结果分析")
    txt(slide, 2.50, 1.05, 3.1, 0.24, "摩擦系数对比趋势（示例）", 13.2, TEAL, True, PP_ALIGN.CENTER)
    txt(slide, 8.88, 1.34, 1.3, 0.22, "数据解读", 13.5, TEAL, True)
    bullet_lines(slide, 8.90, 1.88, [
        "温度升高，μ整体下降",
        "粗糙度增大，μ上升",
        "压强增大，先升后降",
        "材质差异显著，钛合金综合表现最优",
    ], 11.2, 0.48)
    rect(slide, 9.00, 5.35, 3.25, 0.72, rgb("#f4f8f8"), LINE, radius=True)
    txt(slide, 9.45, 5.60, 2.35, 0.32, "结果与文献趋势一致，\n验证方法的有效性", 11.5, TEAL, True, PP_ALIGN.CENTER)


def slide10(slide):
    add_bg(slide, 10)
    cover(slide, [
        (0.0, 0.0, 4.30, 0.92, PAPER),
        (1.05, 1.00, 5.75, 5.05, WHITE),
        (10.76, 1.20, 3.55, 4.65, WHITE),
        (1.20, 7.26, 13.30, 0.58, WHITE),
        (15.08, 8.52, 0.82, 0.34, TEAL),
    ])
    header(slide, 10, "总结展望")
    txt(slide, 2.35, 1.05, 1.0, 0.20, "创新点", 12.5, TEAL, True)
    points = [
        "模块化多工况装置设计",
        "低温稳定控制与高精度测量",
        "多因素系统研究与量化分析",
        "不确定度评估与可靠性提升",
    ]
    for i, value in enumerate(points):
        txt(slide, 2.05, 1.60 + i * 0.93, 3.2, 0.22, value, 12.2, DARK)
    txt(slide, 11.00, 1.05, 2.0, 0.20, "工程应用前景", 12.5, TEAL, True)
    apps = [
        ("极地船舶设计优化", "降低航行阻力"),
        ("表面工程与减阻开发", "提升抗冰性能"),
        ("极地装备与结构安全", "提供理论依据"),
    ]
    for i, (title, body) in enumerate(apps):
        y = 1.60 + i * 1.22
        txt(slide, 11.05, y, 2.2, 0.22, title, 11.8, TEAL, True)
        txt(slide, 11.05, y + 0.32, 1.6, 0.18, body, 9.5, MUTED)
    txt(slide, 1.78, 7.47, 11.8, 0.20, "未来将拓展更多材料与工况，建立摩擦数据库，服务极地工程与安全航行。", 10.8, DARK)


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    builders = [slide01, slide02, slide03, slide04, slide05, slide06, slide07, slide08, slide09, slide10]
    for builder in builders:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        builder(slide)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
