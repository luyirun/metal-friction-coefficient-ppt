from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
PREVIEW_DIR = ROOT / "outputs" / "whitepaper_preview_rebuild"
OUT_DIR = ROOT / "outputs"
OUT = OUT_DIR / "metal_ice_industrial_whitepaper_rebuild_v1.pptx"

W, H = 16, 9
FONT = "Microsoft YaHei"

PAPER = RGBColor(247, 250, 250)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(0, 72, 84)
TEAL_2 = RGBColor(0, 95, 111)
DARK = RGBColor(25, 42, 47)
MUTED = RGBColor(78, 98, 104)
LINE = RGBColor(215, 229, 233)
ICE = RGBColor(237, 247, 250)
ORANGE = RGBColor(226, 126, 39)
GREEN = RGBColor(55, 139, 82)
GRID = RGBColor(229, 237, 239)


def output_path() -> Path:
    if not OUT.exists():
        return OUT
    stem = OUT.stem
    for i in range(2, 100):
        candidate = OUT.with_name(f"{stem[:-2]}v{i}.pptx")
        if not candidate.exists():
            return candidate
    raise RuntimeError("Could not choose versioned output path")


def add_bg(slide, page: int) -> None:
    path = PREVIEW_DIR / f"page_{page:02d}_preview.png"
    slide.shapes.add_picture(str(path), Inches(0), Inches(0), width=Inches(W), height=Inches(H))


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=False, transparency=0, name=None):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    if name:
        shp.name = name
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.fill.transparency = transparency
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    return shp


def line(slide, x1, y1, x2, y2, color=TEAL, width=1.3, name=None):
    shp = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    if name:
        shp.name = name
    shp.line.color.rgb = color
    shp.line.width = Pt(width)
    return shp


def text(slide, x, y, w, h, value, size, color=DARK, bold=False, align=PP_ALIGN.LEFT, name=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, raw in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = raw
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def tag(slide, page: int):
    rect(slide, 0.00, 0.00, 1.48, 0.78, TEAL, name=f"S{page:02d}_page_tag_bg")
    text(slide, 0.23, 0.15, 0.74, 0.34, f"{page:02d}", 24, WHITE, True, name=f"S{page:02d}_page_tag")
    rect(slide, 15.08, 8.51, 0.82, 0.36, TEAL, name=f"S{page:02d}_footer_bg")
    text(slide, 15.34, 8.59, 0.28, 0.12, f"{page:02d}", 10, WHITE, True, PP_ALIGN.CENTER, name=f"S{page:02d}_footer")


def header(slide, page: int, title_value: str, subtitle: str | None = None):
    rect(slide, 0, 0, 9.2, 1.12, PAPER, name=f"S{page:02d}_header_clean")
    tag(slide, page)
    text(slide, 1.42, 0.18, 7.0, 0.34, title_value, 22.5, TEAL, True, name=f"S{page:02d}_title")
    line(slide, 0.72, 0.83, 1.22, 0.83, TEAL, 3)
    line(slide, 1.34, 0.83, 2.56, 0.83, TEAL_2, 1.4)
    if subtitle:
        text(slide, 1.42, 0.96, 6.5, 0.20, subtitle, 12.0, MUTED, name=f"S{page:02d}_subtitle")


def card(slide, x, y, w, h, fill=WHITE, border=LINE, name=None):
    return rect(slide, x, y, w, h, fill, border, True, 0, name)


def bullet(slide, x, y, value, size=10.5, width=4.2, name=None):
    rect(slide, x, y + 0.10, 0.055, 0.055, TEAL, radius=True)
    text(slide, x + 0.18, y, width, 0.24, value, size, DARK, name=name)


def slide01(slide):
    add_bg(slide, 1)
    rect(slide, 0, 0, 1.55, 0.78, PAPER)
    rect(slide, 0.62, 1.35, 5.95, 2.28, PAPER)
    rect(slide, 0.70, 3.72, 3.95, 0.42, PAPER)
    rect(slide, 0.62, 4.70, 3.85, 1.35, PAPER)
    rect(slide, 0.70, 7.95, 5.05, 0.38, TEAL)
    tag(slide, 1)
    text(slide, 0.78, 1.50, 5.45, 1.34, "金属与冰\n摩擦系数的测量", 34, TEAL, True, name="S01_main_title")
    line(slide, 0.80, 3.46, 1.36, 3.46, ORANGE, 2.4)
    line(slide, 1.48, 3.46, 2.62, 3.46, LINE, 1.3)
    text(slide, 0.78, 3.72, 3.7, 0.32, "物理实验创新竞赛答辩", 16.5, DARK, True, name="S01_subtitle")
    text(slide, 1.18, 4.88, 2.15, 0.20, "团队：极冰钥匙者", 11.5, DARK, name="S01_team")
    text(slide, 1.18, 5.58, 2.15, 0.20, "日期：2024.05.20", 11.5, DARK, name="S01_date")
    text(slide, 1.02, 8.04, 4.25, 0.16, "物理实验    |    摩擦学研究    |    精密测量    |    工程创新", 9.6, WHITE, True, name="S01_footer_band")


def slide02(slide):
    add_bg(slide, 2)
    header(slide, 2, "选题背景")
    rect(slide, 8.78, 1.18, 5.75, 3.85, WHITE)
    rect(slide, 0.96, 5.38, 6.58, 1.42, WHITE)
    text(slide, 9.00, 1.36, 4.60, 0.26, "北极航运与冰阻力问题", 16.5, TEAL, True, name="S02_right_title")
    text(slide, 9.00, 2.06, 4.2, 0.20, "航行阻力中，冰障碍阻力占比", 11.0, MUTED, name="S02_desc1")
    text(slide, 9.00, 2.38, 2.10, 0.32, "25%-55%", 21, ORANGE, True, name="S02_number1")
    text(slide, 9.00, 3.40, 4.2, 0.20, "冰面相关事故中，摩擦失控因素占比", 11.0, MUTED, name="S02_desc2")
    text(slide, 9.00, 3.72, 2.10, 0.32, "13%-15%", 21, ORANGE, True, name="S02_number2")
    text(slide, 1.28, 5.52, 1.2, 0.20, "研究意义", 12.5, TEAL, True, name="S02_research")
    for i, item in enumerate([
        "揭示船舶破冰与航行效率的关键摩擦机理之一",
        "为船体材料选型、表面处理与结构设计提供依据",
        "提升极地装备在低温与冰面接触下的安全性",
    ]):
        bullet(slide, 1.18, 5.90 + i * 0.30, item, 9.8, 5.0, f"S02_bullet{i+1}")


def slide03(slide):
    add_bg(slide, 3)
    header(slide, 3, "题目解读")
    rect(slide, 0.95, 1.24, 13.75, 5.92, WHITE)
    columns = [
        ("机理分析", ["冰的黏附性变形", "表面微观水膜", "水膜润滑作用", "温度与压力影响"]),
        ("装置搭建", ["斜面摩擦实验平台", "力学测量与采集系统", "低温环境控制", "标定与校准"]),
        ("结果与不确定度", ["摩擦系数计算", "不确定度评估", "敏感性分析", "结果讨论与验证"]),
    ]
    for i, (head, lines) in enumerate(columns):
        x = 1.35 + i * 4.55
        text(slide, x + 0.55, 1.62, 3.0, 0.26, head, 15.2, TEAL, True, PP_ALIGN.CENTER, name=f"S03_head{i+1}")
        for j, item in enumerate(lines):
            bullet(slide, x + 0.62, 4.18 + j * 0.42, item, 10.8, 2.9, f"S03_b{i+1}_{j+1}")
        if i < 2:
            line(slide, x + 4.35, 1.62, x + 4.35, 6.70, LINE, 0.8)


def slide04(slide):
    add_bg(slide, 4)
    header(slide, 4, "方法选择 —— 动态斜面法")
    rect(slide, 0.95, 1.15, 2.10, 0.38, WHITE)
    rect(slide, 8.48, 1.10, 5.65, 5.70, WHITE)
    text(slide, 1.12, 1.23, 1.4, 0.20, "力学模型", 12.8, TEAL, True, name="S04_model_title")
    text(slide, 8.88, 1.24, 1.4, 0.20, "理论公式", 12.8, TEAL, True, name="S04_formula_title")
    card(slide, 8.98, 1.72, 4.55, 0.64, RGBColor(247, 248, 248), None)
    text(slide, 9.32, 1.92, 3.90, 0.20, "a = g(sinθ − μcosθ)", 15.8, DARK, False, PP_ALIGN.CENTER, "S04_formula_1")
    text(slide, 11.10, 2.58, 0.28, 0.20, "↓", 18, TEAL, True, PP_ALIGN.CENTER)
    card(slide, 8.98, 3.00, 4.55, 0.78, RGBColor(247, 248, 248), None)
    text(slide, 9.25, 3.22, 4.05, 0.24, "μ = (sinθ − a/g) / cosθ", 15.8, DARK, False, PP_ALIGN.CENTER, "S04_formula_2")
    text(slide, 9.16, 4.18, 3.7, 1.00, "a —— 沿斜面方向的加速度\ng —— 重力加速度\nθ —— 斜面倾角\nμ —— 动摩擦系数", 11.2, DARK, name="S04_symbols")


def slide05(slide):
    add_bg(slide, 5)
    header(slide, 5, "装置设计")
    rect(slide, 0.80, 1.00, 3.35, 0.32, WHITE)
    rect(slide, 8.25, 1.00, 4.80, 0.32, WHITE)
    rect(slide, 0.55, 1.62, 1.70, 4.20, WHITE)
    rect(slide, 12.85, 1.62, 1.70, 4.10, WHITE)
    rect(slide, 0.75, 7.40, 14.2, 0.58, WHITE)
    text(slide, 1.00, 1.08, 2.20, 0.18, "斜面冰面模块", 12.6, TEAL, True, name="S05_left_title")
    text(slide, 8.42, 1.08, 3.85, 0.18, "垂直冰面模块（压力加载）", 12.6, TEAL, True, name="S05_right_title")
    for i, item in enumerate(["导轨与滑块", "力传感器", "位移传感器", "冰面板", "角度调节机构", "支撑框架"]):
        text(slide, 0.72, 1.82 + i * 0.47, 1.18, 0.15, item, 8.8, MUTED, name=f"S05_l{i+1}")
    for i, item in enumerate(["压力加载装置", "力传感器", "冰面板", "温控舱体", "底盘"]):
        text(slide, 13.10, 1.84 + i * 0.68, 1.24, 0.15, item, 8.8, MUTED, name=f"S05_r{i+1}")
    for i, item in enumerate(["模块化设计", "快速更换", "低温环境模拟", "高精度测量"]):
        text(slide, 1.15 + i * 3.33, 7.68, 1.32, 0.16, item, 9.6, TEAL, True, PP_ALIGN.CENTER, f"S05_feature{i+1}")


def slide06(slide):
    add_bg(slide, 6)
    header(slide, 6, "变量设计", "六大影响变量与水平设置")
    rect(slide, 0.86, 3.28, 14.35, 2.15, WHITE)
    items = [
        ("温度", "-5 °C\n-10 °C\n-15 °C"),
        ("压强", "0.05 MPa\n0.10 MPa\n0.20 MPa"),
        ("材质", "铝合金\n不锈钢\n钛合金"),
        ("粗糙度", "Ra 0.2 μm\nRa 1.0 μm\nRa 3.0 μm"),
        ("面积", "1 cm²\n4 cm²\n9 cm²"),
        ("成分", "淡水冰\n海冰\n人工掺盐冰"),
    ]
    for i, (head, vals) in enumerate(items):
        x = 1.20 + i * 2.42
        text(slide, x - 0.48, 3.35, 0.96, 0.18, head, 11.6, TEAL, True, PP_ALIGN.CENTER, f"S06_head{i+1}")
        text(slide, x - 0.62, 4.02, 1.24, 0.72, vals, 9.8, DARK, False, PP_ALIGN.CENTER, f"S06_vals{i+1}")


def slide07(slide):
    add_bg(slide, 7)
    header(slide, 7, "问题改进")
    rect(slide, 1.35, 1.42, 4.10, 5.35, RGBColor(247, 250, 249))
    rect(slide, 7.18, 1.42, 4.50, 5.35, RGBColor(247, 250, 249))
    rows = [
        ("冰面制备不平整", "表面粗糙和气泡较大\n影响测量重复性", "改进方案", "定制制冰模具 + 控温体系\n实现均匀透明平整冰面"),
        ("低温环境稳定性差", "温度波动导致数据漂移\n影响测量精度", "改进方案", "双层保温 + PID温控\n温度波动 ≤ ±0.2 °C"),
        ("传感器可靠性不足", "低温漂移与噪声干扰大\n影响数据信号", "改进方案", "低温型传感器 + 屏蔽布线\n多点校准与滤波处理"),
    ]
    for i, (p, pd, st, sv) in enumerate(rows):
        y = 1.66 + i * 1.72
        text(slide, 1.66, y, 2.36, 0.20, p, 11.2, DARK, True, name=f"S07_problem{i+1}")
        text(slide, 1.66, y + 0.30, 2.60, 0.36, pd, 9.5, MUTED, name=f"S07_problem_desc{i+1}")
        text(slide, 7.52, y, 1.10, 0.20, st, 11.2, GREEN, True, name=f"S07_solution_tag{i+1}")
        text(slide, 7.52, y + 0.30, 2.90, 0.36, sv, 9.5, DARK, name=f"S07_solution{i+1}")


def slide08(slide):
    add_bg(slide, 8)
    header(slide, 8, "数据处理")
    rect(slide, 1.06, 1.45, 14.15, 5.95, WHITE)
    steps = [
        ("数据采集", ["力信号", "位移信号", "温度/压力", "采样频率 100 Hz"]),
        ("加速度计算", ["位移二次微分", "滤波处理", "得到 a(t)"]),
        ("摩擦系数计算", ["μ = (sinθ − a/g) / cosθ", "代入角度与加速度", "计算 μ"]),
        ("不确定度评估", ["A类不确定度", "B类不确定度", "合成标准不确定度", "扩展不确定度 U"]),
    ]
    for i, (head, vals) in enumerate(steps):
        x = 1.30 + i * 3.48
        text(slide, x + 0.40, 2.10, 1.72, 0.20, head, 11.4, TEAL, True, PP_ALIGN.CENTER, f"S08_head{i+1}")
        if i == 2:
            text(slide, x + 0.10, 3.88, 2.32, 0.45, "μ = (sinθ − a/g)\n      / cosθ", 13.0, DARK, False, PP_ALIGN.CENTER, "S08_formula")
            for j, item in enumerate(vals[1:]):
                bullet(slide, x + 0.48, 5.20 + j * 0.30, item, 8.8, 1.8, f"S08_b{i+1}_{j+1}")
        else:
            for j, item in enumerate(vals):
                bullet(slide, x + 0.42, 5.02 + j * 0.30, item, 8.8, 1.8, f"S08_b{i+1}_{j+1}")


def chart_point(x0, y0, cw, ch, x_index, y_value):
    px = x0 + 0.55 + x_index * (cw - 1.10) / 3
    py = y0 + ch - (y_value / 0.15) * ch
    return px, py


def slide09(slide):
    add_bg(slide, 9)
    header(slide, 9, "结果分析")
    rect(slide, 1.05, 1.14, 7.20, 6.28, WHITE)
    rect(slide, 8.68, 1.24, 5.05, 3.60, WHITE)
    rect(slide, 8.90, 5.38, 3.90, 0.94, RGBColor(242, 248, 248), LINE, True)
    text(slide, 2.32, 1.22, 3.9, 0.20, "摩擦系数对比趋势（示例）", 12.4, TEAL, True, PP_ALIGN.CENTER, "S09_chart_title")
    x0, y0, cw, ch = 1.45, 1.90, 5.85, 4.25
    for k in range(6):
        y = y0 + k * ch / 5
        line(slide, x0, y, x0 + cw, y, GRID, 0.6)
        text(slide, x0 - 0.44, y - 0.06, 0.32, 0.12, f"{0.15 - k * 0.03:.2f}", 7.2, DARK, False, PP_ALIGN.RIGHT)
    line(slide, x0, y0, x0, y0 + ch, DARK, 0.8)
    line(slide, x0, y0 + ch, x0 + cw, y0 + ch, DARK, 0.8)
    for i, val in enumerate([-15, -10, -5, 0]):
        x = chart_point(x0, y0, cw, ch, i, 0)[0]
        text(slide, x - 0.20, y0 + ch + 0.14, 0.40, 0.12, str(val), 7.4, DARK, False, PP_ALIGN.CENTER)
    data = [
        ("铝合金", RGBColor(41, 137, 190), [0.047, 0.040, 0.030, 0.020]),
        ("不锈钢", ORANGE, [0.116, 0.105, 0.087, 0.064]),
        ("钛合金", RGBColor(112, 164, 52), [0.082, 0.070, 0.058, 0.039]),
    ]
    for s_idx, (name, color, vals) in enumerate(data):
        pts = [chart_point(x0, y0, cw, ch, i, val) for i, val in enumerate(vals)]
        for p1, p2 in zip(pts, pts[1:]):
            line(slide, p1[0], p1[1], p2[0], p2[1], color, 1.4, f"S09_chart_{name}_line")
        for p_idx, (px, py) in enumerate(pts):
            dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(px - 0.035), Inches(py - 0.035), Inches(0.07), Inches(0.07))
            dot.name = f"S09_chart_{name}_point_{p_idx + 1}"
            dot.fill.solid()
            dot.fill.fore_color.rgb = color
            dot.line.color.rgb = color
        text(slide, 7.55, 2.10 + s_idx * 0.28, 0.90, 0.14, name, 7.4, DARK)
    text(slide, 3.35, 6.42, 1.45, 0.16, "温度 / °C", 8.8, DARK, False, PP_ALIGN.CENTER, "S09_x_label")
    text(slide, 0.76, 3.30, 0.42, 0.90, "摩擦系数 μ", 8.8, DARK, name="S09_y_label")
    text(slide, 9.02, 1.42, 1.25, 0.20, "数据解读", 12.6, TEAL, True, name="S09_insight_title")
    for i, item in enumerate(["温度升高，μ整体下降", "粗糙度增大，μ上升", "压强增大，μ先升后降", "材质差异显著，钛合金综合表现更优"]):
        bullet(slide, 9.02, 2.02 + i * 0.45, item, 9.8, 3.6, f"S09_insight{i+1}")
    text(slide, 9.35, 5.72, 2.85, 0.34, "结果与文献趋势一致，\n验证方法的有效性。", 10.4, TEAL, True, PP_ALIGN.CENTER, "S09_conclusion")


def slide10(slide):
    add_bg(slide, 10)
    header(slide, 10, "总结展望")
    rect(slide, 1.70, 1.02, 4.80, 5.05, WHITE)
    rect(slide, 10.72, 1.02, 3.70, 5.10, WHITE)
    rect(slide, 1.72, 7.22, 12.25, 0.58, WHITE)
    text(slide, 2.28, 1.12, 1.10, 0.18, "创新点", 11.8, TEAL, True, name="S10_innovation_title")
    for i, item in enumerate(["模块化多工况装置设计", "低温稳定控制与高精度测量", "多因素系统研究与量化分析", "不确定度评估与可靠性提升"]):
        text(slide, 2.06, 1.76 + i * 0.86, 3.7, 0.18, item, 10.4, DARK, name=f"S10_innovation{i+1}")
    text(slide, 10.90, 1.12, 1.90, 0.18, "工程应用前景", 11.8, TEAL, True, name="S10_application_title")
    apps = [("极地船舶设计优化", "降低航行阻力"), ("表面工程与减阻开发", "提升抗冰性能"), ("极地装备与结构安全", "提供理论依据")]
    for i, (head, body) in enumerate(apps):
        y = 1.78 + i * 1.22
        text(slide, 11.12, y, 2.10, 0.18, head, 10.4, TEAL, True, name=f"S10_app{i+1}_title")
        text(slide, 11.12, y + 0.30, 1.60, 0.16, body, 8.8, MUTED, name=f"S10_app{i+1}_body")
    text(slide, 1.96, 7.41, 10.8, 0.18, "未来将拓展更多材料与工况，建立摩擦数据库，服务极地工程与安全航行。", 9.0, DARK, name="S10_future")


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    for idx, builder in enumerate([slide01, slide02, slide03, slide04, slide05, slide06, slide07, slide08, slide09, slide10], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        builder(slide)
    out = output_path()
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
