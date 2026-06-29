from __future__ import annotations

import json
import math
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt


SLIDE_W_PT = 960
SLIDE_H_PT = 640
EMU_PER_PT = 12700

ROOT = Path(r"F:\桌面\Keynote缩略图\因素矩阵")
REF_DIR = ROOT / "static_reference_states_20260627"
ASSET_DIR = ROOT / "editable_static_overview"
OUT_DIR = ROOT / "gorden_static_single_pages"
BG_PATH = ASSET_DIR / "ice_blue_clean_background.png"

BLUE = "003B73"
MID_BLUE = "005A99"
LINE_BLUE = "477EB8"
ICE_LINE = "A8C7DB"
ACCENT = "FF7A1A"
TEXT_MUTED = "315C7A"


def ep(pt: float) -> Emu:
    return Emu(int(round(pt * EMU_PER_PT)))


def rgb(hex_value: str) -> RGBColor:
    hex_value = hex_value.strip().lstrip("#")
    return RGBColor(int(hex_value[0:2], 16), int(hex_value[2:4], 16), int(hex_value[4:6], 16))


def set_no_fill(shape) -> None:
    shape.fill.background()


def set_line(shape, color: str = LINE_BLUE, width: float = 0.85) -> None:
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: float,
    color: str = BLUE,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    name: str | None = None,
):
    box = slide.shapes.add_textbox(ep(x), ep(y), ep(w), ep(h))
    if name:
        box.name = name
    set_no_fill(box)
    box.line.fill.background()
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_panel(slide, x: float, y: float, w: float, h: float, name: str, line: str = LINE_BLUE, fill_alpha: float = 0.24):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, ep(x), ep(y), ep(w), ep(h))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb("FFFFFF")
    shape.fill.transparency = fill_alpha
    set_line(shape, line, 0.85)
    return shape


def add_line(slide, x1: float, y1: float, x2: float, y2: float, color: str = MID_BLUE, width: float = 1.0, name: str | None = None):
    line = slide.shapes.add_connector(1, ep(x1), ep(y1), ep(x2), ep(y2))
    if name:
        line.name = name
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def add_circle(slide, x: float, y: float, d: float, name: str, fill: str | None = None, line: str = MID_BLUE, width: float = 1.2):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, ep(x), ep(y), ep(d), ep(d))
    shape.name = name
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    else:
        set_no_fill(shape)
    set_line(shape, line, width)
    return shape


def add_factor_icon(slide, key: str, x: float, y: float, scale: float = 1.0):
    def sx(v): return x + v * scale
    def sy(v): return y + v * scale
    c = MID_BLUE
    if key == "temperature":
        add_line(slide, sx(18), sy(2), sx(18), sy(36), c, 1.9)
        add_circle(slide, sx(9), sy(32), 18 * scale, f"icon_{key}_dot", None, c, 1.9)
        add_text(slide, "*", sx(31), sy(12), 16 * scale, 16 * scale, 15 * scale, c, True)
    elif key == "pressure":
        add_circle(slide, sx(4), sy(5), 44 * scale, f"icon_{key}_dial", None, c, 1.6)
        add_line(slide, sx(25), sy(28), sx(39), sy(16), c, 1.7)
        add_line(slide, sx(15), sy(50), sx(39), sy(50), c, 1.3)
    elif key == "roughness":
        pts = [(0, 42), (5, 18), (11, 43), (16, 16), (22, 43), (28, 19), (34, 43), (40, 17), (46, 42)]
        for (a, b), (c1, d) in zip(pts, pts[1:]):
            add_line(slide, sx(a), sy(b), sx(c1), sy(d), c, 1.2)
        add_line(slide, sx(0), sy(52), sx(46), sy(52), c, 1.2)
    elif key == "area":
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, ep(sx(3)), ep(sy(5)), ep(45 * scale), ep(45 * scale))
        set_no_fill(box)
        set_line(box, c, 1.4)
        for r in range(4):
            for col in range(4):
                add_circle(slide, sx(13 + col * 8), sy(15 + r * 8), 2.5 * scale, f"icon_area_{r}_{col}", c, c, 0.6)
    else:
        lines = [
            (10, 38, 27, 49), (27, 49, 45, 38), (10, 38, 27, 27),
            (27, 27, 45, 38), (27, 27, 27, 49)
        ]
        if key == "ice":
            lines += [(10, 38, 10, 50), (45, 38, 45, 50), (10, 50, 27, 60), (27, 60, 45, 50)]
        for i, (x1, y1, x2, y2) in enumerate(lines):
            add_line(slide, sx(x1), sy(y1), sx(x2), sy(y2), c, 1.3, f"icon_{key}_{i}")


FACTORS = [
    {"key": "temperature", "no": "F01", "name": "温度", "en": "Temperature", "x": 42, "y": 218},
    {"key": "pressure", "no": "F02", "name": "压强", "en": "Pressure", "x": 194, "y": 218},
    {"key": "material", "no": "F03", "name": "金属材质", "en": "Metal Material", "x": 346, "y": 218},
    {"key": "roughness", "no": "F04", "name": "表面粗糙度", "en": "Surface Roughness", "x": 42, "y": 376},
    {"key": "area", "no": "F05", "name": "接触面积", "en": "Contact Area", "x": 194, "y": 376},
    {"key": "ice", "no": "F06", "name": "冰的成分", "en": "Ice Composition", "x": 346, "y": 376},
]


DETAILS = {
    "temperature": {
        "slug": "01_temperature_f01",
        "title": "温度变量分析",
        "en": "Temperature",
        "no": "F01",
        "big": "F01",
        "active": "temperature",
        "table_head": ["水平", "变量取值", "滑（摩擦系数）"],
        "rows": [["T1", "-15°C", "0.118"], ["T2", "-10°C", "0.103"], ["T3", "-5°C", "0.089"], ["T4", "0°C", "0.076"]],
        "x_label": "温度水平（°C）",
        "conclusion": "温度升高时，冰表面液膜增强，摩擦系数 μ 整体下降。",
    },
    "pressure": {
        "slug": "02_pressure_f02",
        "title": "压强变量分析",
        "en": "Pressure",
        "no": "F02",
        "big": "F02",
        "active": "pressure",
        "table_head": ["水平", "变量取值", "滑（摩擦系数）"],
        "rows": [["P1", "0.05 MPa", "0.082"], ["P2", "0.10 MPa", "0.096"], ["P3", "0.20 MPa", "0.091"]],
        "x_label": "压强水平",
        "conclusion": "压强增大会改变局部接触状态与压力融化程度，摩擦系数可能呈现先升后降趋势。",
    },
    "material": {
        "slug": "03_material_f03",
        "title": "金属材质变量分析",
        "en": "Metal Material",
        "no": "F03",
        "big": "F03",
        "active": "material",
        "table_head": ["水平", "变量取值", "滑（摩擦系数）"],
        "rows": [["M1", "铝合金", "0.118"], ["M2", "不锈钢", "0.103"], ["M3", "铜", "0.089"], ["M4", "钛合金", "0.076"]],
        "x_label": "金属材质水平",
        "conclusion": "金属材质对摩擦系数影响显著：铝合金 > 不锈钢 > 铜 > 钛合金。",
    },
    "roughness": {
        "slug": "04_roughness_f04",
        "title": "表面粗糙度变量分析",
        "en": "Surface Roughness",
        "no": "F04",
        "big": "F04",
        "active": "roughness",
        "table_head": ["水平", "表面粗糙度 Ra", "滑（摩擦系数）"],
        "rows": [["R1", "Ra 0.2 μm", "0.074"], ["R2", "Ra 1.0 μm", "0.096"], ["R3", "Ra 3.0 μm", "0.124"]],
        "x_label": "表面粗糙度水平",
        "conclusion": "表面粗糙度增大时，微观嵌入和犁削效应增强，摩擦系数 μ 上升。",
    },
    "area": {
        "slug": "05_contact_area_f05",
        "title": "接触面积变量分析",
        "en": "Contact Area",
        "no": "F05",
        "big": "F05",
        "active": "area",
        "table_head": ["水平", "变量取值", "滑（摩擦系数）"],
        "rows": [["A1", "1 cm²", "0.112"], ["A2", "4 cm²", "0.096"], ["A3", "9 cm²", "0.087"]],
        "x_label": "接触面积水平",
        "conclusion": "接触面积变化会改变单位压强和局部液膜状态，使摩擦系数发生变化。",
    },
    "ice": {
        "slug": "06_ice_composition_f06",
        "title": "冰的成分变量分析",
        "en": "Ice Composition",
        "no": "F06",
        "big": "F06",
        "active": "ice",
        "table_head": ["水平", "变量取值", "滑（摩擦系数）"],
        "rows": [["I1", "纯水冰", "0.104"], ["I2", "海冰", "0.089"], ["I3", "人工盐冰", "0.078"]],
        "x_label": "冰的成分水平",
        "conclusion": "冰中盐分或杂质会改变冰点与表面液膜状态，从而影响金属-冰界面的摩擦行为。",
    },
}


def new_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = ep(SLIDE_W_PT)
    prs.slide_height = ep(SLIDE_H_PT)
    return prs


def add_common(slide):
    slide.shapes.add_picture(str(BG_PATH), ep(0), ep(0), ep(SLIDE_W_PT), ep(SLIDE_H_PT))
    add_text(slide, "09", 44, 38, 86, 58, 40, BLUE, True, name="slide_index")
    add_line(slide, 46, 94, 116, 94, ACCENT, 3, "orange_rule")
    add_text(slide, "因素矩阵", 142, 43, 238, 38, 28, BLUE, True, name="main_title")
    add_text(slide, "控制变量法下的摩擦系数影响因素总览", 144, 86, 520, 20, 12, BLUE, True, name="main_subtitle")
    add_text(slide, "点击任一因素查看实验水平、趋势图与物理结论", 44, 128, 480, 18, 9, "476378", False, name="main_hint")
    add_line(slide, 42, 592, 918, 592, "8DB2C7", 0.75, "footer_rule")
    add_text(slide, "*", 44, 604, 22, 20, 20, "0077A3", True, name="footer_icon")
    add_text(slide, "金属与冰摩擦系数测量", 72, 608, 260, 18, 10, BLUE, True, name="footer_title")
    add_text(slide, "09/10", 842, 604, 72, 20, 14, BLUE, True, name="footer_page")


def add_button(slide, x: float, y: float, w: float = 68, h: float = 23):
    shape = add_panel(slide, x, y, w, h, "button", MID_BLUE, 0.35)
    add_text(slide, "点击查看", x + 11, y + 6, 44, 10, 8.4, "00447A", True, PP_ALIGN.CENTER)
    add_text(slide, ">", x + w - 17, y + 4, 10, 12, 10, MID_BLUE, True)
    return shape


def add_card(slide, factor: dict, active: str | None = None):
    x, y = factor["x"], factor["y"]
    is_active = factor["key"] == active
    card = add_panel(slide, x, y, 142, 136, f"card_{factor['key']}", "13A8FF" if is_active else LINE_BLUE, 0.24)
    if is_active:
        card.line.width = Pt(1.4)
    add_text(slide, factor["no"], x + 14, y + 16, 36, 16, 8.7, "00447A", True)
    add_text(slide, factor["name"], x + 14, y + 44, 96, 24, 15.4, BLUE, True)
    add_text(slide, factor["en"], x + 14, y + 73, 96, 14, 7.8, TEXT_MUTED)
    add_factor_icon(slide, factor["key"], x + 92, y + 45, 1.0)
    add_button(slide, x + 60, y + 102)


def add_cards(slide, active: str | None = None):
    for f in FACTORS:
        add_card(slide, f, active)


def add_overview(slide):
    add_panel(slide, 502, 118, 430, 448, "overview_panel", "70A8D0", 0.24)
    add_text(slide, "因素总览", 524, 142, 146, 26, 18, BLUE, True)
    add_text(slide, "OVERVIEW", 526, 170, 90, 14, 8, TEXT_MUTED)
    add_text(slide, "OVERVIEW", 746, 128, 168, 42, 30, "99BDD8", True)
    add_panel(slide, 518, 198, 398, 104, "flow_box", "B2CCE0", 0.55)
    flow = [
        ("=", "控制变量", "固定其他因素"),
        ("^", "单因素变化", "改变一个因素水平"),
        ("u", "摩擦系数 μ", "测量响应变化"),
        ("[]", "对比结论", "分析影响规律"),
    ]
    for i, (icon, title, sub) in enumerate(flow):
        cx = 548 + i * 102
        add_circle(slide, cx, 212, 42, f"flow_{i}", None, "2A78C4", 1.0)
        add_text(slide, icon, cx + 8, 220, 26, 16, 16, MID_BLUE, True, PP_ALIGN.CENTER)
        add_text(slide, title, cx - 12, 260, 66, 14, 8.1, BLUE, True, PP_ALIGN.CENTER)
        add_text(slide, sub, cx - 13, 280, 72, 12, 7.2, "456579", False, PP_ALIGN.CENTER)
        if i < 3:
            add_text(slide, ">", cx + 68, 226, 22, 18, 18, "8FB9D6", True)

    add_text(slide, "因素分组概览", 524, 324, 110, 20, 11.5, BLUE, True)
    add_text(slide, "六大因素关系概览", 704, 324, 150, 20, 11.5, BLUE, True)
    groups = [
        ("temperature", "环境条件", "温度、冰的成分", 350),
        ("pressure", "接触状态", "压强、接触面积", 412),
        ("material", "材料表面", "表面材质、表面粗糙度", 474),
    ]
    for key, title, desc, y in groups:
        add_panel(slide, 524, y, 148, 54, f"group_{key}", "98BEDA", 0.65)
        add_factor_icon(slide, key, 536, y + 10, 0.7)
        add_text(slide, title, 578, y + 13, 70, 14, 8.7, BLUE, True)
        add_text(slide, desc, 578, y + 32, 86, 12, 7.0, TEXT_MUTED)
    add_panel(slide, 686, 350, 230, 178, "radar_box", "B2CCE0", 0.55)
    add_radar(slide, 800, 438, 58)
    add_panel(slide, 518, 536, 398, 34, "conclusion_box", "B2CCE0", 0.5)
    add_text(slide, "!", 534, 545, 14, 16, 12, MID_BLUE, True)
    add_text(slide, "结论：摩擦系数由材料属性、冰面状态与接触条件共同决定，需逐项控制变量分析。", 556, 544, 344, 18, 8.3, BLUE, True)


def add_radar(slide, cx: float, cy: float, r: float):
    points = []
    for i in range(6):
        ang = math.radians(-90 + i * 60)
        x = cx + math.cos(ang) * r
        y = cy + math.sin(ang) * r
        points.append((x, y))
        add_line(slide, cx, cy, x, y, ICE_LINE, 0.7)
    value = []
    for x, y in points:
        value.append((cx + (x - cx) * 0.62, cy + (y - cy) * 0.62))
    for pts, color, width in [(points, ICE_LINE, 0.7), (value, "0077A3", 1.5)]:
        for i, (x, y) in enumerate(pts):
            x2, y2 = pts[(i + 1) % len(pts)]
            add_line(slide, x, y, x2, y2, color, width)
    for i, (x, y) in enumerate(value):
        add_circle(slide, x - 3, y - 3, 6, f"radar_dot_{i}", "EAF6FF", "0077A3", 1.2)
    labels = [
        ("温度\nF01", cx - 16, cy - r - 32, 50),
        ("压强\nF02", cx + r + 8, cy - 13, 48),
        ("金属材质\nF03", cx + 48, cy + 48, 64),
        ("表面粗糙度\nF04", cx - 25, cy + r + 6, 72),
        ("接触面积\nF05", cx - r - 62, cy + 45, 62),
        ("冰的成分\nF06", cx - r - 64, cy - 12, 62),
    ]
    for text, x, y, w in labels:
        add_text(slide, text, x, y, w, 28, 8, BLUE, True, PP_ALIGN.CENTER)


def add_active_arrow(slide, active: str):
    factor = next(f for f in FACTORS if f["key"] == active)
    x = factor["x"] + 136
    y = factor["y"] + 72
    arrow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, ep(x), ep(y), ep(44), ep(26))
    arrow.name = f"active_arrow_{active}"
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = rgb("29A8FF")
    arrow.fill.transparency = 0.10
    arrow.line.fill.background()


def add_table(slide, x: float, y: float, w: float, h: float, headers: list[str], rows: list[list[str]]):
    add_panel(slide, x, y, w, h, "data_table", "B5D1E5", 0.72)
    cols = len(headers)
    total_rows = len(rows) + 1
    row_h = h / total_rows
    col_w = [0.22 * w, 0.39 * w, 0.39 * w]
    xs = [x, x + col_w[0], x + col_w[0] + col_w[1], x + w]
    for i in range(1, total_rows):
        add_line(slide, x, y + i * row_h, x + w, y + i * row_h, "B7D4E8", 0.55)
    for vx in xs[1:-1]:
        add_line(slide, vx, y, vx, y + h, "B7D4E8", 0.55)
    for c, head in enumerate(headers):
        add_text(slide, head, xs[c] + 5, y + 10, xs[c + 1] - xs[c] - 10, 12, 7.7, BLUE, True, PP_ALIGN.CENTER)
    for r_i, row in enumerate(rows):
        for c, value in enumerate(row):
            add_text(slide, value, xs[c] + 5, y + (r_i + 1) * row_h + row_h * 0.32, xs[c + 1] - xs[c] - 10, 12, 8.0, BLUE, False, PP_ALIGN.CENTER)


def add_chart(slide, x: float, y: float, w: float, h: float, labels: list[str], values: list[float], x_label: str):
    y_min = 0.0
    y_max = 0.16 if max(values) > 0.12 else 0.14
    px0, py0 = x + 42, y + h - 38
    px1, py1 = x + w - 16, y + 28
    add_line(slide, px0, py0, px1, py0, "7EAAD0", 0.8)
    add_line(slide, px0, py0, px0, py1, "7EAAD0", 0.8)
    tick_count = 4
    for i in range(tick_count + 1):
        tv = y_min + (y_max - y_min) * i / tick_count
        ty = py0 - (py0 - py1) * i / tick_count
        add_line(slide, px0 - 3, ty, px0, ty, "7EAAD0", 0.55)
        add_text(slide, f"{tv:.2f}", x + 4, ty - 5, 32, 10, 6.9, BLUE, False, PP_ALIGN.RIGHT)
        if i > 0:
            add_line(slide, px0, ty, px1, ty, "D0E2EF", 0.35)
    pts = []
    for i, value in enumerate(values):
        denom = max(1, len(values) - 1)
        px = px0 + (px1 - px0) * i / denom
        py = py0 - (py0 - py1) * ((value - y_min) / (y_max - y_min))
        pts.append((px, py))
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        add_line(slide, x1, y1, x2, y2, "0077D9", 1.8)
    for (px, py), label, value in zip(pts, labels, values):
        add_circle(slide, px - 3.5, py - 3.5, 7, f"chart_dot_{label}", "EAF6FF", "0077D9", 1.5)
        add_text(slide, f"{value:.3f}", px - 16, py - 22, 34, 10, 6.8, BLUE, False, PP_ALIGN.CENTER)
        add_text(slide, label, px - 12, py0 + 8, 24, 10, 7.2, BLUE, False, PP_ALIGN.CENTER)
    add_text(slide, x_label, px0 + 35, py0 + 26, 160, 12, 7.4, BLUE, False, PP_ALIGN.CENTER)
    add_text(slide, "滑（摩擦系数）", x + 24, y + 10, 82, 12, 7.2, BLUE, False)
    add_line(slide, x + w - 72, y + 10, x + w - 58, y + 10, "0077D9", 1.6)
    add_circle(slide, x + w - 67, y + 7.8, 4, "chart_legend_dot", "EAF6FF", "0077D9", 1)
    add_text(slide, "摩擦系数（滑）", x + w - 55, y + 5, 72, 12, 6.8, BLUE)


def add_detail_panel(slide, detail: dict):
    add_panel(slide, 502, 118, 430, 448, "detail_panel", "70A8D0", 0.24)
    add_text(slide, detail["title"], 524, 142, 178, 26, 17, BLUE, True)
    add_panel(slide, 654, 146, 28, 18, "factor_badge", "70A8D0", 0.62)
    add_text(slide, detail["no"], 659, 149, 20, 10, 8, BLUE, True, PP_ALIGN.CENTER)
    add_text(slide, detail["en"], 526, 171, 110, 14, 8.5, BLUE)
    add_text(slide, detail["big"], 812, 130, 88, 42, 30, "99BDD8", True)
    add_text(slide, "X", 905, 134, 18, 18, 14, BLUE, True, PP_ALIGN.CENTER)
    add_panel(slide, 518, 198, 398, 332, "content_box", "B2CCE0", 0.62)
    add_text(slide, "实验水平表", 556, 222, 106, 20, 12.5, BLUE, True)
    add_text(slide, "趋势图", 716, 222, 80, 20, 12.5, BLUE, True)
    add_text(slide, "△", 534, 222, 16, 18, 15, MID_BLUE, True)
    add_text(slide, "▥", 696, 222, 16, 18, 13, MID_BLUE, True)
    add_table(slide, 524, 258, 172, 166, detail["table_head"], detail["rows"])
    labels = [r[0] for r in detail["rows"]]
    values = [float(r[2]) for r in detail["rows"]]
    add_chart(slide, 704, 254, 198, 174, labels, values, detail["x_label"])
    add_panel(slide, 518, 454, 398, 64, "detail_conclusion", "B2CCE0", 0.5)
    add_text(slide, "!", 536, 466, 16, 16, 13, MID_BLUE, True)
    add_text(slide, "结论", 556, 465, 60, 18, 13, BLUE, True)
    add_text(slide, detail["conclusion"], 536, 491, 356, 18, 8.5, BLUE, True)
    add_text(slide, "注：当前数据为示例，后续可替换为真实实验数据。", 526, 544, 250, 14, 7.4, "476378")
    add_panel(slide, 804, 532, 106, 26, "back_button", MID_BLUE, 0.35)
    add_text(slide, "<  返回矩阵", 816, 540, 82, 10, 9, BLUE, True, PP_ALIGN.CENTER)


def save_single(slug: str, builder) -> dict:
    prs = new_prs()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_common(slide)
    builder(slide)
    pptx = OUT_DIR / f"factor_matrix_static_{slug}.pptx"
    prs.save(pptx)
    return {"slug": slug, "pptx": str(pptx)}


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(BG_PATH, OUT_DIR / BG_PATH.name)

    outputs = []
    outputs.append(save_single("00_overview", lambda s: (add_cards(s), add_overview(s))))
    for key, detail in DETAILS.items():
        def builder(slide, d=detail):
            add_cards(slide, d["active"])
            add_active_arrow(slide, d["active"])
            add_detail_panel(slide, d)
        outputs.append(save_single(detail["slug"], builder))

    manifest = {
        "workflow": "GordenImage2PPTX-inspired layered static reconstruction",
        "slide_size_pt": [SLIDE_W_PT, SLIDE_H_PT],
        "source_reference_dir": str(REF_DIR),
        "output_dir": str(OUT_DIR),
        "layers": {
            "background": "image layer, generated clean ice-blue texture",
            "framework": "editable PowerPoint panels/lines/tables/charts reconstructed from reference screenshots for editability",
            "icons_decor": "editable PowerPoint line icons and chart marks",
            "text": "editable PowerPoint text boxes",
        },
        "outputs": outputs,
        "notes": [
            "No animation or one-page trigger logic is added in this static batch.",
            "F05 detail badge and x-axis labels were corrected to F05/A1-A3 instead of copying visible reference typos.",
        ],
    }
    (OUT_DIR / "static_pages_layer_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
